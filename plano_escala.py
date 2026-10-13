"""
MÓDULO: Plano de escala — Sistema Nala   [PLANO ESCALA]
======================================================
v1, Etapa 1 (12/10/2026), SÓ MERCADO LIVRE. Regra do [ANALISTA ESCALA]
(08/10, validada pelo Thiago); método da planilha "Plano de escala" do Drive
(1R12xedUR4VqLCYNY8UwstDwmaOESLssws0ksvEYiFdM). Onde as duas divergem vale a
ESPECIFICAÇÃO (ver tests/test_plano_escala.py).

O QUE FAZ
  - Cadastro (seção 2): um plano por anúncio, 1 ativo por (loja, anúncio).
    ÚNICA porta de entrada: o .xlsx ou a tabela editável da tela, que passam
    pela MESMA previa() + gravar(). Subir de novo SUBSTITUI o plano da chave;
    plano já começado exige motivo (vai para plano_escala_decisao) e tem o
    previsto recalculado inteiro, mantendo as vendas reais.
  - Previsto (seções 3 e 4), calculado no cadastro e gravado por rodada:
      meta mensal = 20% da venda/mês do concorrente nº 1 (ou a meta própria);
      vendas por rodada = 0,5× a meta, 2× a meta, depois 1× a meta;
      preço = concorrente × 0,85, depois ×1,05, depois ×1,03 por rodada; o
      último degrau é o PRÓPRIO preço meta (o degrau que passaria dele é
      travado nele); cada rodada aceita preço digitado à mão;
      ACOS-alvo 10%, 10%, depois −1 ponto por rodada até 5%;
      parcela da venda vinda do ads 75%, −5 pontos por rodada até 40%;
      TACOS = ACOS × parcela;
      resultado da rodada = margem unitária NO PREÇO DA RODADA × vendas −
        receita × TACOS (decisão do Thiago, 12/10: a margem sai da conta da
        Tabela de Preço — comissão e imposto em %, frete pela barreira de
        R$ 79 —, não de um preço de margem zero fixo; o preço de margem zero
        aparece só como referência);
      TETO = o menor resultado acumulado (buraco máximo); ponto de virada = a
        venda em que o acumulado volta a zero.
    Tarifa do Full FORA (v2). Margem = antes de ads e Full, como margem_total.
  - Real (seções 5, 6, 8, 10, 13): vendas do anúncio desde a data de início
    até o último dia FECHADO (sinais_dia.DIAS_ATRASO_API); a rodada fecha no
    dia em que as vendas acumuladas atingem as previstas (a sobra daquele dia
    fica nela); margem real = margem_total − gasto de ads do anúncio
    (fact_ads_performance, ML) — ads subtraído UMA vez; parcela do ads real =
    unidades de ads (diretas + indiretas) ÷ unidades vendidas; ritmo e
    "lenta"; situação ativo / pede revisão (acumulado real abaixo do TETO
    antes da virada) / concluído / encerrado.

  O sistema AVISA e RECOMENDA: não escreve nada no marketplace.

INSUMOS DA MARGEM (de onde vem cada um aparece na tela)
  custo      dim_produtos.preco_a_ser_considerado (custo de hoje)
  imposto    dim_lojas.imposto da loja
  comissão   informada no cadastro → vendas do anúncio (90 dias) → vendas do
             mesmo SKU em outros anúncios da loja → média da categoria na
             loja. Só vendas da API em FULL/ENVIOS: no FLEX a comissão vem com
             a tarifa do Flex e o imposto zerado.
  frete      abaixo de R$ 79: vendas do anúncio → do SKU na loja → taxa fixa
             da Tabela de Preço; acima de R$ 79: informado no cadastro →
             vendas do anúncio → do SKU na loja (o sistema não tem peso
             cadastrado, então a tabela de frete por peso não serve).
  extra      R$ 6 na ML-YanniSP (prep center), como a Tabela de Preço.
"""

import io
import math
import re
from datetime import date, datetime, timedelta

import pandas as pd

from sinais_dia import DIAS_ATRASO_API, BRT, MARKETPLACE
from tabela_preco import calcular_margem, taxa_fixa_dba

# ============================================================
# REGRA (calibrar aqui)
# ============================================================

META_PCT_CONCORRENTE = 0.20          # meta mensal = 20% da venda do concorrente nº 1
FATORES_VENDAS = (0.5, 2.0, 1.0)     # rodada 1, rodada 2, da 3ª em diante (× meta)
FATOR_PRECO_R1 = 0.85                # preço da rodada 1 = concorrente × 0,85
FATOR_PRECO_R2 = 1.05                # rodada 2 = rodada 1 × 1,05
FATOR_PRECO_DEPOIS = 1.03            # da 3ª em diante: × 1,03 por rodada
ACOS_INICIAL = 0.10                  # rodadas 1 e 2
ACOS_PASSO = 0.01                    # −1 ponto por rodada a partir da 3ª
ACOS_PISO = 0.05
PARCELA_INICIAL = 0.75               # parcela da venda vinda do ads, rodada 1
PARCELA_PASSO = 0.05
PARCELA_PISO = 0.40
DIAS_RODADA = (15, 30, 30)           # opção A do [MESTRE ANÁLISES]: rodada 1, 2, 3ª+
LENTA_CORTE = 0.70                   # ritmo < 70% do esperado
LENTA_DIAS_MIN = 7                   # ... depois de 7 dias de rodada
LENTA_SEGUIDOS_VOLTAR = 7            # lenta 7 dias seguidos em rodada negativa → voltar
ALERTA_PARCELA_PP = 0.10             # parcela do ads real ≥ prevista + 10 pontos
ALERTA_PARCELA_RODADAS = 2           # ... em 2 rodadas (fechadas) seguidas
MAX_RODADAS = 40                     # guarda contra preço meta muito acima do concorrente

BARREIRA_ML = 79.0                   # abaixo: taxa fixa; a partir de: frete do vendedor
JANELA_INSUMOS = 90                  # dias de venda para comissão e frete observados
MIN_UNID_INSUMO = 3                  # unidades mínimas para usar o observado
LOGISTICAS_INSUMO = ('FULL', 'ENVIOS')
EXTRA_POR_LOJA = {'ML-YanniSP': 6.0}  # igual a tabela_preco.render_tab_ml

CHAR_NOME = 45                       # nome resumido do item

# ============================================================
# PREVISTO (sem banco)
# ============================================================


def _arred(x):
    """Arredonda meio para cima (o round do Python arredonda 0,5 para o par)."""
    return int(math.floor(x + 0.5))


def meta_mensal(venda_concorrente, meta_propria):
    """A meta própria, se informada; senão 20% da venda do concorrente nº 1."""
    if meta_propria:
        return float(meta_propria)
    return float(venda_concorrente or 0) * META_PCT_CONCORRENTE


def acos_da_rodada(n):
    return max(ACOS_PISO, ACOS_INICIAL - ACOS_PASSO * max(0, n - 2))


def parcela_da_rodada(n):
    return max(PARCELA_PISO, PARCELA_INICIAL - PARCELA_PASSO * (n - 1))


def dias_da_rodada(n):
    return DIAS_RODADA[min(n, len(DIAS_RODADA)) - 1]


def vendas_da_rodada(meta, n):
    return max(1, _arred(meta * FATORES_VENDAS[min(n, len(FATORES_VENDAS)) - 1]))


def montar_rodadas(meta, preco_concorrente, preco_meta, precos_mao=None):
    """As rodadas até o preço meta. precos_mao: {rodada: preço} digitado à
    mão; o degrau seguinte parte dele. A rodada cujo preço chega ao preço meta
    (ou passaria dele) é a última, travada no preço meta."""
    precos_mao = precos_mao or {}
    rodadas, preco = [], None
    for n in range(1, MAX_RODADAS + 1):
        if n in precos_mao:
            preco, a_mao = float(precos_mao[n]), True
        else:
            if n == 1:
                preco = preco_concorrente * FATOR_PRECO_R1
            else:
                preco = preco * (FATOR_PRECO_R2 if n == 2 else FATOR_PRECO_DEPOIS)
            a_mao = False
        ultima = preco >= preco_meta - 1e-9
        if ultima and not a_mao:
            preco = float(preco_meta)
        acos, parcela = acos_da_rodada(n), parcela_da_rodada(n)
        rodadas.append({'rodada': n, 'vendas': vendas_da_rodada(meta, n), 'preco': preco,
                        'preco_a_mao': a_mao, 'acos': acos, 'parcela_ads': parcela,
                        'tacos': acos * parcela, 'dias_esperados': dias_da_rodada(n)})
        if ultima:
            return rodadas
    raise ValueError(f'O preço não chega ao preço meta em {MAX_RODADAS} rodadas.')


def prever(rodadas, margem_unit):
    """Completa cada rodada com margem unitária (no preço da rodada), receita,
    gasto de ads previsto, resultado e acumulado. margem_unit(preço) -> R$."""
    acum = 0.0
    for r in rodadas:
        m = margem_unit(r['preco'])
        if m is None:
            raise ValueError(f"Sem frete para o preço da rodada {r['rodada']} "
                             f"(R$ {r['preco']:.2f}).")
        r['margem_unit'] = m
        r['receita'] = r['vendas'] * r['preco']
        r['gasto_ads'] = r['receita'] * r['tacos']
        r['resultado'] = m * r['vendas'] - r['gasto_ads']
        acum += r['resultado']
        r['acumulado'] = acum
    return rodadas


def teto_e_virada(rodadas):
    """(TETO, ponto de virada). TETO = menor acumulado (≤ 0; dentro da rodada
    o acumulado é uma reta, então o mínimo cai no fim de uma rodada). Virada =
    a venda (unidades desde o início) em que o acumulado volta a zero depois
    do TETO; 0 se nunca fica negativo; None se não volta dentro do plano."""
    teto, i_min = 0.0, -1
    for i, r in enumerate(rodadas):
        if r['acumulado'] < teto:
            teto, i_min = r['acumulado'], i
    if i_min < 0:
        return 0.0, 0
    unid = sum(r['vendas'] for r in rodadas[:i_min + 1])
    acum = teto
    for r in rodadas[i_min + 1:]:
        if acum + r['resultado'] >= 0 and r['resultado'] > 0:
            por_unid = r['resultado'] / r['vendas']
            return teto, math.ceil(unid + (-acum) / por_unid - 1e-9)
        acum += r['resultado']
        unid += r['vendas']
    return teto, None


# ---- Margem do sistema no preço (conta da Tabela de Preço) ---------------

def margem_ml(custo, comissao, imposto, frete_abaixo=None, frete_acima=None, extra=0.0,
              preco_ref_abaixo=None):
    """margem_unit(preço) do ML com a conta de tabela_preco.calcular_margem:
    comissão e imposto em % do preço; abaixo da barreira de R$ 79 o frete é a
    taxa fixa (observada, ou a da Tabela de Preço); a partir dela, o frete do
    vendedor (None = desconhecido).
    O frete observado abaixo de R$ 79 é a média das vendas ao preço de hoje
    (preco_ref_abaixo); a rodada em outra faixa da taxa fixa (R$ 4,50 até
    R$ 30, R$ 6,50 até R$ 50, R$ 6,75 até R$ 79) leva a diferença entre as
    faixas (ressalva 1 do auditor, 12/10)."""
    def f(preco):
        if preco < BARREIRA_ML:
            if frete_abaixo is None:
                frete = taxa_fixa_dba(preco)
            elif preco_ref_abaixo:
                frete = frete_abaixo + taxa_fixa_dba(preco) - taxa_fixa_dba(preco_ref_abaixo)
            else:
                frete = frete_abaixo
        else:
            if frete_acima is None:
                return None
            frete = frete_acima
        return calcular_margem(preco, custo, comissao, frete, imposto, extra)[0]
    return f


def preco_margem_zero(margem_unit):
    """Menor preço com margem ≥ 0 (referência na tela). Procura abaixo da
    barreira de R$ 79 e, se não houver, acima. None se não achar."""
    for lo, hi in ((0.01, BARREIRA_ML - 0.01), (BARREIRA_ML, 100000.0)):
        m_hi = margem_unit(hi)
        if m_hi is None or m_hi < 0:
            continue
        m_lo = margem_unit(lo)
        if m_lo is not None and m_lo >= 0:
            return round(lo, 2)
        for _ in range(60):
            meio = (lo + hi) / 2
            m = margem_unit(meio)
            if m is not None and m >= 0:
                hi = meio
            else:
                lo = meio
        return math.ceil(hi * 100 - 1e-6) / 100
    return None


# ---- Insumos: de onde vêm comissão e frete --------------------------------

def _obs_vazio():
    return {'q': 0, 'v': 0.0, 'c': 0.0, 'f': 0.0}


def _somar(dst, q, v, c, f):
    dst['q'] += q
    dst['v'] += v
    dst['c'] += c
    dst['f'] += f


def agregar_observado(linhas):
    """linhas: (loja, anuncio, sku, categoria, acima, qtd, valor, comissao,
    frete) da SQL_OBSERVADO -> {'anuncio'|'sku'|'categoria': {chave: {False:
    obs abaixo de R$ 79, True: obs acima}}}."""
    niveis = {'anuncio': {}, 'sku': {}, 'categoria': {}}
    for loja, anuncio, sku, cat, acima, q, v, c, f in linhas:
        q, v, c, f = int(q or 0), float(v or 0), float(c or 0), float(f or 0)
        for nivel, chave in (('anuncio', (loja, anuncio)), ('sku', (loja, sku)),
                             ('categoria', (loja, cat))):
            if chave[1] is None:
                continue
            d = niveis[nivel].setdefault(chave, {False: _obs_vazio(), True: _obs_vazio()})
            _somar(d[bool(acima)], q, v, c, f)
    return niveis


def _junto(d):
    t = _obs_vazio()
    for faixa in (False, True):
        _somar(t, d[faixa]['q'], d[faixa]['v'], d[faixa]['c'], d[faixa]['f'])
    return t


ROTULO_NIVEL = {
    'anuncio': 'vendas do anúncio',
    'sku': 'vendas do SKU em outros anúncios da loja',
    'categoria': 'média da categoria na loja',
}


def resolver_insumos(loja, anuncio, sku, categoria, observado,
                     comissao_informada=None, frete_acima_informado=None):
    """{'comissao', 'comissao_origem', 'frete_abaixo', 'frete_abaixo_origem',
    'frete_acima', 'frete_acima_origem'}; valor None = não achou."""
    chaves = {'anuncio': (loja, anuncio), 'sku': (loja, sku), 'categoria': (loja, categoria)}
    jan = f'{JANELA_INSUMOS} dias'
    out = {}
    if comissao_informada is not None:
        out['comissao'], out['comissao_origem'] = comissao_informada, 'informada no cadastro'
    else:
        out['comissao'], out['comissao_origem'] = None, 'não achada: informe a Comissão %'
        for nivel in ('anuncio', 'sku', 'categoria'):
            d = observado.get(nivel, {}).get(chaves[nivel])
            t = _junto(d) if d else None
            if t and t['q'] >= MIN_UNID_INSUMO and t['v'] > 0:
                out['comissao'] = t['c'] / t['v']
                out['comissao_origem'] = f"{ROTULO_NIVEL[nivel]} ({jan}, {t['q']} un.)"
                break
    for faixa, campo, padrao in ((False, 'frete_abaixo', 'taxa fixa da Tabela de Preço'),
                                 (True, 'frete_acima', None)):
        if faixa and frete_acima_informado is not None:
            out[campo], out[campo + '_origem'] = frete_acima_informado, 'informado no cadastro'
            continue
        out[campo] = None
        out[campo + '_origem'] = padrao or 'não achado: informe o Frete acima de R$ 79'
        if not faixa:
            out['preco_ref_abaixo'] = None
        for nivel in ('anuncio', 'sku'):
            d = observado.get(nivel, {}).get(chaves[nivel])
            o = d[faixa] if d else None
            if o and o['q'] >= MIN_UNID_INSUMO:
                out[campo] = o['f'] / o['q']
                if not faixa:
                    out['preco_ref_abaixo'] = o['v'] / o['q']
                out[campo + '_origem'] = f"{ROTULO_NIVEL[nivel]} ({jan}, {o['q']} un.)"
                break
    return out


def calcular_plano(cad, produto, imposto, insumos):
    """O previsto completo de uma linha de cadastro validada. produto: {'custo',
    'categoria'}; imposto em fração. Devolve (rodadas, resumo) ou levanta
    ValueError com a frase para a gestora."""
    if insumos['comissao'] is None:
        raise ValueError('Comissão não achada nas vendas: informe a Comissão %.')
    custo = produto.get('custo')
    if not custo or custo <= 0:
        raise ValueError('SKU sem custo cadastrado.')
    extra = EXTRA_POR_LOJA.get(cad['loja'], 0.0)
    f = margem_ml(custo, insumos['comissao'], imposto, insumos['frete_abaixo'],
                  insumos['frete_acima'], extra, insumos.get('preco_ref_abaixo'))
    meta = meta_mensal(cad['venda_concorrente'], cad['meta_propria'])
    rodadas = prever(montar_rodadas(meta, cad['preco_concorrente'], cad['preco_meta'],
                                    cad['precos_mao']), f)
    teto, virada = teto_e_virada(rodadas)
    avisos = []
    sobra = sorted(k for k in cad['precos_mao'] if k < 1 or k > len(rodadas))
    if sobra:
        avisos.append('Preço à mão ignorado na(s) rodada(s) ' + ', '.join(map(str, sobra)) +
                      f': o plano tem {len(rodadas)} rodada(s).')
    insumos = dict(insumos)
    if all(r['preco'] < BARREIRA_ML for r in rodadas) and insumos['frete_acima'] is None:
        insumos['frete_acima_origem'] = 'não usado: todas as rodadas abaixo de R$ 79'
    if all(r['preco'] >= BARREIRA_ML for r in rodadas) and insumos['frete_abaixo'] is None:
        insumos['frete_abaixo_origem'] = 'não usado: todas as rodadas a partir de R$ 79'
    resumo = {
        'meta_mensal': meta, 'teto': teto, 'virada': virada,
        'resultado_total': rodadas[-1]['acumulado'],
        'vendas_total': sum(r['vendas'] for r in rodadas),
        'preco_margem_zero': preco_margem_zero(f),
        'avisos': avisos,
        'insumos': {'custo': custo, 'imposto': imposto, 'extra': extra, **insumos},
    }
    return rodadas, resumo


# ============================================================
# REAL × PREVISTO (sem banco)
# ============================================================

def fim_fechado(hoje):
    """Último dia FECHADO de venda (o mesmo fim das janelas do Sinais do Dia)."""
    return hoje - timedelta(days=DIAS_ATRASO_API + 1)


def serie_diaria(inicio, fim, vendas, ads):
    """Um dia por linha de inicio a fim (dia sem venda = 0). vendas: {data:
    (unid, receita, margem)}; ads: {data: (gasto, unid_ads)}."""
    dias, d = [], inicio
    while d <= fim:
        u, rec, mg = vendas.get(d, (0, 0.0, 0.0))
        g, ua = ads.get(d, (0.0, 0))
        dias.append({'data': d, 'unid': int(u), 'receita': float(rec), 'margem': float(mg),
                     'gasto': float(g), 'unid_ads': int(ua)})
        d += timedelta(days=1)
    return dias


def alocar_rodadas(previstas, dias):
    """Distribui os dias pelas rodadas, em ordem: o dia em que as vendas da
    rodada atingem as previstas a fecha (a sobra do dia fica nela) e a
    seguinte começa no dia seguinte. Dias depois da última rodada fechada não
    contam. Devolve uma lista (uma por rodada alcançada) com o real."""
    reais, k, atual = [], 0, None
    for dia in dias:
        if k >= len(previstas):
            break
        if atual is None:
            atual = {'rodada': previstas[k]['rodada'], 'inicio': dia['data'], 'fim': None,
                     'dias': [], 'unid': 0, 'receita': 0.0, 'margem': 0.0, 'gasto': 0.0,
                     'unid_ads': 0, 'fechada': False}
            reais.append(atual)
        atual['dias'].append(dia)
        for c in ('unid', 'receita', 'margem', 'gasto', 'unid_ads'):
            atual[c] += dia[c]
        if atual['unid'] >= previstas[k]['vendas']:
            atual['fechada'], atual['fim'] = True, dia['data']
            k, atual = k + 1, None
    for r in reais:
        r['margem_real'] = r['margem'] - r['gasto']
        r['parcela_ads'] = r['unid_ads'] / r['unid'] if r['unid'] else None
    return reais


def ritmo(prevista, real):
    """Ritmo da rodada aberta: real/dia × esperado/dia, lenta (abaixo de 70%
    depois de 7 dias) e quantos dias seguidos (até o último fechado) ela está
    lenta."""
    esperado = prevista['vendas'] / prevista['dias_esperados']
    acum, seguidos = 0, 0
    for i, dia in enumerate(real['dias'], start=1):
        acum += dia['unid']
        lenta_i = i >= LENTA_DIAS_MIN and (acum / i) < LENTA_CORTE * esperado
        seguidos = seguidos + 1 if lenta_i else 0
    n = len(real['dias'])
    por_dia = real['unid'] / n if n else 0.0
    return {'dias': n, 'por_dia': por_dia, 'esperado': esperado,
            'razao': por_dia / esperado if esperado else None,
            'lenta': n >= LENTA_DIAS_MIN and por_dia < LENTA_CORTE * esperado,
            'lenta_seguidos': seguidos}


def alerta_parcela(previstas, reais):
    """Parcela do ads real ≥ prevista + 10 pontos em 2 rodadas FECHADAS seguidas."""
    seguidas = 0
    for p, r in zip(previstas, reais):
        if not r['fechada'] or r['parcela_ads'] is None:
            seguidas = 0
            continue
        seguidas = seguidas + 1 if r['parcela_ads'] - p['parcela_ads'] >= ALERTA_PARCELA_PP - 1e-9 else 0
        if seguidas >= ALERTA_PARCELA_RODADAS:
            return True
    return False


def acompanhar(plano, previstas, dias, hoje):
    """O real do plano: rodadas, rodada atual, ritmo, situação e
    recomendação. plano: {'situacao', 'data_inicio', 'teto', 'como_encerrou'}."""
    fim = fim_fechado(hoje)
    reais = alocar_rodadas(previstas, dias)
    acum, virou, foi_negativo, acum_dia = 0.0, False, False, []
    for r in reais:
        for d in r['dias']:
            acum += d['margem'] - d['gasto']
            acum_dia.append(acum)
            if acum < 0:
                foi_negativo = True
            elif foi_negativo:
                virou = True
    concluido = len(reais) == len(previstas) and bool(reais) and reais[-1]['fechada']
    aberta = next((r for r in reais if not r['fechada']), None)
    if aberta is None and not concluido:
        n_atual = len(reais) + 1 if len(reais) < len(previstas) else len(previstas)
    else:
        n_atual = aberta['rodada'] if aberta else len(previstas)
    prev_atual = previstas[n_atual - 1]
    rit = ritmo(prev_atual, aberta) if aberta else None
    teto = plano['teto']
    if plano['situacao'] == 'encerrado':
        situacao = 'encerrado'
    elif plano['data_inicio'] > fim:
        situacao = 'aguardando início'
    elif concluido:
        situacao = 'concluído'
    elif acum < teto - 1e-9 and not virou:
        situacao = 'pede revisão'
    else:
        situacao = 'ativo'
    recomendacao = ''
    if situacao == 'pede revisão':
        recomendacao = 'Acumulado real abaixo do TETO antes da virada: rever o plano.'
    elif rit and prev_atual['resultado'] < 0 and rit['lenta_seguidos'] >= LENTA_SEGUIDOS_VOLTAR:
        recomendacao = ('Rodada negativa lenta há 7 dias seguidos: voltar ao preço de margem '
                        'zero e rever anúncio e ads.')
    elif rit and rit['lenta']:
        recomendacao = 'Rodada lenta: ritmo abaixo de 70% do esperado.'
    elif situacao == 'concluído':
        recomendacao = 'Plano concluído: uma rodada inteira rodou no preço meta.'
    return {'reais': reais, 'rodada_atual': n_atual, 'prevista_atual': prev_atual,
            'ritmo': rit, 'acumulado_real': acum, 'virou': virou, 'situacao': situacao,
            'alerta_parcela': alerta_parcela(previstas, reais), 'recomendacao': recomendacao,
            'unid_total': sum(r['unid'] for r in reais), 'fim': fim}


# ============================================================
# CADASTRO: layout do .xlsx / tabela, validação, prévia (sem banco)
# ============================================================

COLUNAS = [
    ('loja', 'Loja'),
    ('anuncio', 'Anúncio (MLB)'),
    ('sku', 'SKU'),
    ('produto', 'Produto (não editar)'),
    ('venda_concorrente', 'Venda/mês do concorrente nº 1'),
    ('meta_propria', 'Meta própria (vendas/mês)'),
    ('preco_concorrente', 'Preço do concorrente'),
    ('preco_meta', 'Preço meta (~15% de margem)'),
    ('data_inicio', 'Data de início'),
    ('gestora', 'Gestora responsável'),
    ('comissao', 'Comissão % (opcional)'),
    ('frete_acima', 'Frete acima de R$ 79 (opcional)'),
    ('precos_mao', 'Preços à mão (rodada=preço; ...)'),
    ('motivo', 'Motivo (obrigatório para revisar plano já começado)'),
]
CABECALHO = [c[1] for c in COLUNAS]
CAMPO_DO_CABECALHO = {c[1]: c[0] for c in COLUNAS}
CAMPOS_PLANO = ('sku', 'venda_concorrente', 'meta_propria', 'preco_concorrente', 'preco_meta',
                'data_inicio', 'gestora', 'comissao', 'frete_acima', 'precos_mao')

RE_MLB = re.compile(r'^MLB\d{6,}$')


def _vazio(v):
    return v is None or (isinstance(v, float) and math.isnan(v)) or \
        (isinstance(v, str) and not v.strip()) or v is pd.NaT


def _num(v, campo):
    if _vazio(v):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace('R$', '').replace('%', '').replace(' ', '')
    if ',' in s:
        s = s.replace('.', '').replace(',', '.')
    try:
        return float(s)
    except ValueError:
        raise ValueError(f'{campo}: "{v}" não é número.')


def _data(v):
    if _vazio(v):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, pd.Timestamp):
        return v.date()
    s = str(v).strip()
    for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d/%m/%y'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f'Data de início: "{v}" não é data (use dd/mm/aaaa).')


def ler_precos_mao(v):
    """'1=35,90; 3=37,90' -> {1: 35.9, 3: 37.9}."""
    if _vazio(v):
        return {}
    out = {}
    for parte in re.split(r'[;\n]+', str(v)):
        parte = parte.strip()
        if not parte:
            continue
        m = re.match(r'^(?:R(?:odada)?\s*)?(\d+)\s*[=:]\s*(?:R\$)?\s*([\d.,]+)$', parte, re.I)
        if not m:
            raise ValueError(f'Preços à mão: "{parte}" fora do formato rodada=preço.')
        out[int(m.group(1))] = _num(m.group(2), 'Preços à mão')
    return out


def texto_precos_mao(d):
    return '; '.join(f'{k}={v:.2f}'.replace('.', ',') for k, v in sorted((d or {}).items()))


def normalizar(bruta):
    """Linha crua (dict campo->valor) -> cadastro normalizado. Levanta
    ValueError com TODAS as falhas da linha numa frase."""
    erros, cad = [], {}
    cad['loja'] = str(bruta.get('loja') or '').strip()
    cad['anuncio'] = str(bruta.get('anuncio') or '').strip().upper().replace('-', '')
    cad['sku'] = str(bruta.get('sku') or '').strip().upper()
    cad['gestora'] = str(bruta.get('gestora') or '').strip()
    cad['motivo'] = '' if _vazio(bruta.get('motivo')) else str(bruta.get('motivo')).strip()
    for campo, rotulo in (('venda_concorrente', 'Venda do concorrente'),
                          ('meta_propria', 'Meta própria'),
                          ('preco_concorrente', 'Preço do concorrente'),
                          ('preco_meta', 'Preço meta'), ('comissao', 'Comissão %'),
                          ('frete_acima', 'Frete acima de R$ 79')):
        try:
            cad[campo] = _num(bruta.get(campo), rotulo)
        except ValueError as e:
            erros.append(str(e))
            cad[campo] = None
    try:
        cad['data_inicio'] = _data(bruta.get('data_inicio'))
    except ValueError as e:
        erros.append(str(e))
        cad['data_inicio'] = None
    try:
        cad['precos_mao'] = ler_precos_mao(bruta.get('precos_mao'))
    except ValueError as e:
        erros.append(str(e))
        cad['precos_mao'] = {}
    for campo in ('venda_concorrente', 'meta_propria'):
        if cad[campo] is not None:
            cad[campo] = int(round(cad[campo]))
    if cad['comissao'] is not None and cad['comissao'] > 1:
        cad['comissao'] = cad['comissao'] / 100          # "11,5" -> 0,115
    if not cad['loja']:
        erros.append('Loja vazia.')
    if not RE_MLB.match(cad['anuncio']):
        erros.append(f'Anúncio "{cad["anuncio"]}" não é um MLB.')
    if not cad['sku']:
        erros.append('SKU vazio.')
    if not cad['gestora']:
        erros.append('Gestora vazia.')
    if not (cad['venda_concorrente'] or 0) > 0 and not (cad['meta_propria'] or 0) > 0:
        erros.append('Informe a venda/mês do concorrente nº 1 OU a meta própria.')
    if not (cad['preco_concorrente'] or 0) > 0:
        erros.append('Preço do concorrente vazio.')
    if not (cad['preco_meta'] or 0) > 0:
        erros.append('Preço meta vazio.')
    if cad['data_inicio'] is None and not any('Data' in e for e in erros):
        erros.append('Data de início vazia.')
    if cad['comissao'] is not None and not 0 < cad['comissao'] < 0.5:
        erros.append('Comissão % fora de 0–50%.')
    if any(p <= 0 for p in cad['precos_mao'].values()):
        erros.append('Preço à mão precisa ser maior que zero.')
    if erros:
        raise ValueError(' '.join(erros))
    return cad


def _igual(cad, ativo):
    """O cadastro não muda nada no plano ativo (subir o mesmo arquivo = nada)."""
    for c in CAMPOS_PLANO:
        a, b = cad.get(c), ativo.get(c)
        if isinstance(a, float) or isinstance(b, float):
            if (a is None) != (b is None) or (a is not None and abs(float(a) - float(b)) > 0.005):
                return False
        elif c == 'precos_mao':
            if {int(k): round(float(v), 2) for k, v in (a or {}).items()} != \
               {int(k): round(float(v), 2) for k, v in (b or {}).items()}:
                return False
        elif a != b:
            return False
    return True


def previa(linhas_brutas, ativos, lojas_permitidas, hoje, calcular):
    """Uma entrada por linha: {'n', 'cad', 'acao', 'erro', 'rodadas',
    'resumo', 'ativo'}. acao: 'novo' | 'revisao' | 'cadastro' (plano ativo
    ainda não começado, reenviado) | 'sem mudança' | 'recusado'.
    ativos: {(loja, anuncio): plano ativo (com 'concluido')}.
    calcular(cad) -> (rodadas, resumo) ou ValueError."""
    fim = fim_fechado(hoje)
    vistos, out = set(), []
    for n, bruta in enumerate(linhas_brutas, start=2):
        if all(_vazio(v) for k, v in bruta.items() if k != 'produto'):
            continue
        item = {'n': n, 'cad': None, 'acao': 'recusado', 'erro': '', 'aviso': '',
                'rodadas': None, 'resumo': None, 'ativo': None}
        out.append(item)
        try:
            cad = normalizar(bruta)
        except ValueError as e:
            item['erro'] = str(e)
            continue
        item['cad'] = cad
        chave = (cad['loja'], cad['anuncio'])
        if cad['loja'] not in lojas_permitidas:
            item['erro'] = f'Loja {cad["loja"]} fora das suas lojas do Mercado Livre.'
            continue
        if chave in vistos:
            item['erro'] = 'Anúncio repetido no arquivo (1 plano por anúncio).'
            continue
        vistos.add(chave)
        ativo = ativos.get(chave)
        item['ativo'] = ativo
        if ativo and _igual(cad, ativo):
            item['acao'] = 'sem mudança'
            continue
        try:
            item['rodadas'], item['resumo'] = calcular(cad)
        except ValueError as e:
            item['erro'] = str(e)
            continue
        item['aviso'] = ' '.join(item['resumo'].get('avisos') or [])
        if not ativo or ativo.get('concluido'):
            item['acao'] = 'novo'
        elif ativo['data_inicio'] <= fim:
            if not cad['motivo']:
                item['erro'] = 'Plano já começado: escreva o motivo da revisão.'
                continue
            item['acao'] = 'revisao'
        else:
            item['acao'] = 'cadastro'
    return out


ROTULO_ACAO = {'novo': '🆕 novo plano', 'revisao': '✏️ revisão (recalcula o previsto)',
               'cadastro': '✏️ substitui (ainda não começou)', 'sem mudança': '— sem mudança',
               'recusado': '⛔ recusado'}


# ============================================================
# EXCEL (ida e volta no mesmo layout)
# ============================================================

def linha_do_plano(p, nomes):
    par = p.get('parametros') or {}
    inf = par.get('informado') or {}
    return {
        'loja': p['loja'], 'anuncio': p['anuncio'], 'sku': p['sku'],
        'produto': nome_resumido(nomes.get(p['sku'])),
        'venda_concorrente': p.get('venda_concorrente'), 'meta_propria': p.get('meta_propria'),
        'preco_concorrente': p.get('preco_concorrente'), 'preco_meta': p.get('preco_meta'),
        'data_inicio': p.get('data_inicio'), 'gestora': p.get('gestora'),
        'comissao': (inf['comissao'] * 100) if inf.get('comissao') is not None else None,
        'frete_acima': inf.get('frete_acima'),
        'precos_mao': texto_precos_mao({int(k): v for k, v in (par.get('precos_mao') or {}).items()}),
        'motivo': None,
    }


def excel_cadastro(linhas):
    """O .xlsx de cadastro: aba Planos (o layout que volta) + aba Como
    preencher."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = 'Planos'
    ws.append(CABECALHO)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='DDEBF7')
    for linha in linhas:
        ws.append([linha.get(c) for c, _ in COLUNAS])
    larg = {'produto': 40, 'precos_mao': 30, 'motivo': 45}
    for i, (c, rot) in enumerate(COLUNAS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = larg.get(c, max(14, len(rot) * 0.8))
        if c == 'data_inicio':
            for cel in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                for x in cel:
                    x.number_format = 'DD/MM/YYYY'
    ws.freeze_panes = 'A2'
    ajuda = wb.create_sheet('Como preencher')
    for t in (
        'Uma linha por anúncio (MLB). Subir de novo SUBSTITUI o plano do anúncio; não soma.',
        'Venda/mês do concorrente nº 1 OU Meta própria: a meta mensal é 20% da venda do concorrente, ou a meta própria.',
        'Preço meta: o preço com ~15% de margem bruta; é o último degrau.',
        'Comissão % e Frete acima de R$ 79: só se quiser forçar; vazio = o sistema tira das vendas e mostra de onde.',
        'Preços à mão: troca o preço de uma rodada, ex.: 1=35,90; 3=37,90. Os degraus seguintes partem dele.',
        'Motivo: obrigatório para mudar um plano que já começou (fica no registro de decisões).',
        'Produto: só para conferência; é ignorado na volta.',
    ):
        ajuda.append([t])
    ajuda.column_dimensions['A'].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def ler_excel(conteudo):
    """.xlsx -> [linha crua {campo: valor}]. Recusa cabeçalho diferente."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(conteudo), data_only=True)
    ws = wb['Planos'] if 'Planos' in wb.sheetnames else wb.active
    linhas = list(ws.iter_rows(values_only=True))
    if not linhas:
        raise ValueError('Planilha vazia.')
    cab = [str(c).strip() if c is not None else '' for c in linhas[0]]
    faltam = [c for c in CABECALHO if c not in cab]
    if faltam:
        raise ValueError('Faltam colunas do layout: ' + ', '.join(faltam) +
                         '. Baixe o modelo da tela e preencha nele.')
    idx = {CAMPO_DO_CABECALHO[c]: cab.index(c) for c in CABECALHO}
    return [{campo: (row[i] if i < len(row) else None) for campo, i in idx.items()}
            for row in linhas[1:]]


def linhas_do_editor(df):
    """DataFrame da tabela editável (colunas = CABECALHO) -> linhas cruas."""
    return [{CAMPO_DO_CABECALHO[c]: r.get(c) for c in CABECALHO}
            for r in df.to_dict('records')]


# ============================================================
# SQL (psycopg2, parâmetros ligados; sem schema: o teste usa TEMP)
# ============================================================

SQL_EXISTE = "SELECT to_regclass('plano_escala') IS NOT NULL"

SQL_LOJAS_ML = "SELECT loja, imposto FROM dim_lojas WHERE marketplace = %(marketplace)s ORDER BY loja"

SQL_ATIVOS = """
    SELECT id, loja, anuncio, sku, venda_concorrente, meta_propria, preco_concorrente,
           preco_meta, data_inicio, gestora, parametros, situacao, como_encerrou,
           criado_por, criado_em, atualizado_por, atualizado_em
      FROM plano_escala
     WHERE marketplace = %(marketplace)s AND loja = ANY(%(lojas)s) AND situacao = 'ativo'
     ORDER BY loja, anuncio
"""

SQL_RODADAS = """
    SELECT plano_id, rodada, vendas, preco, preco_a_mao, acos, parcela_ads, tacos,
           dias_esperados, margem_unit, receita, gasto_ads, resultado, acumulado
      FROM plano_escala_rodada
     WHERE plano_id = ANY(%(ids)s)
     ORDER BY plano_id, rodada
"""

SQL_DECISOES = """
    SELECT plano_id, tipo, rodada, motivo, criado_por, criado_em
      FROM plano_escala_decisao
     WHERE plano_id = ANY(%(ids)s)
     ORDER BY plano_id, criado_em
"""

SQL_PRODUTOS = """
    SELECT sku, nome, categoria, preco_a_ser_considerado
      FROM dim_produtos WHERE sku = ANY(%(skus)s)
"""

# Comissão e frete observados: só vendas da API em FULL/ENVIOS (o FLEX traz a
# tarifa do Flex na comissão e o imposto zerado; o upload antigo não separa).
SQL_OBSERVADO = """
    SELECT f.loja_origem, f.codigo_anuncio, f.sku, p.categoria,
           (f.valor_venda_efetivo / f.quantidade) >= %(barreira)s AS acima,
           SUM(f.quantidade), SUM(f.valor_venda_efetivo), SUM(f.comissao), SUM(f.frete)
      FROM fact_vendas_snapshot f
      LEFT JOIN dim_produtos p ON p.sku = f.sku
     WHERE f.marketplace_origem = %(marketplace)s AND f.loja_origem = ANY(%(lojas)s)
       AND f.logistica = ANY(%(logisticas)s)
       AND f.data_venda >= %(ini)s AND f.data_venda <= %(fim)s
     GROUP BY 1, 2, 3, 4, 5
"""

SQL_REAL_VENDAS = """
    SELECT loja_origem, codigo_anuncio, data_venda,
           SUM(quantidade), SUM(valor_venda_efetivo), SUM(margem_total)
      FROM fact_vendas_snapshot
     WHERE marketplace_origem = %(marketplace)s AND loja_origem = ANY(%(lojas)s)
       AND codigo_anuncio = ANY(%(anuncios)s)
       AND data_venda >= %(ini)s AND data_venda <= %(fim)s
     GROUP BY 1, 2, 3
"""

# Parcela do ads em UNIDADES: no ML o coletor grava vendas_diretas/indiretas =
# direct/indirect_units_quantity e vendas = units_quantity (nala-coletor-ml,
# coletar_ads.py); o denominador é SUM(quantidade) do snapshot, também unidade.
SQL_REAL_ADS = """
    SELECT loja, codigo_anuncio, data, SUM(gasto_ads),
           SUM(COALESCE(vendas_diretas + vendas_indiretas, vendas))
      FROM fact_ads_performance
     WHERE marketplace = %(marketplace)s AND loja = ANY(%(lojas)s)
       AND codigo_anuncio = ANY(%(anuncios)s)
       AND data >= %(ini)s AND data <= %(fim)s
     GROUP BY 1, 2, 3
"""

SQL_INSERIR_PLANO = """
    INSERT INTO plano_escala
           (marketplace, loja, anuncio, sku, venda_concorrente, meta_propria,
            preco_concorrente, preco_meta, data_inicio, gestora, parametros, criado_por)
    VALUES (%(marketplace)s, %(loja)s, %(anuncio)s, %(sku)s, %(venda_concorrente)s,
            %(meta_propria)s, %(preco_concorrente)s, %(preco_meta)s, %(data_inicio)s,
            %(gestora)s, %(parametros)s, %(usuario)s)
    RETURNING id
"""

SQL_ATUALIZAR_PLANO = """
    UPDATE plano_escala
       SET sku = %(sku)s, venda_concorrente = %(venda_concorrente)s,
           meta_propria = %(meta_propria)s, preco_concorrente = %(preco_concorrente)s,
           preco_meta = %(preco_meta)s, data_inicio = %(data_inicio)s, gestora = %(gestora)s,
           parametros = %(parametros)s, atualizado_por = %(usuario)s,
           atualizado_em = (now() AT TIME ZONE 'America/Sao_Paulo')
     WHERE id = %(id)s AND marketplace = %(marketplace)s AND loja = ANY(%(lojas)s)
       AND situacao = 'ativo'
"""

SQL_ENCERRAR = """
    UPDATE plano_escala
       SET situacao = 'encerrado', como_encerrou = %(como)s, encerrado_por = %(usuario)s,
           encerrado_em = (now() AT TIME ZONE 'America/Sao_Paulo')
     WHERE id = %(id)s AND marketplace = %(marketplace)s AND loja = ANY(%(lojas)s)
       AND situacao = 'ativo'
"""

SQL_APAGAR_RODADAS = "DELETE FROM plano_escala_rodada WHERE plano_id = %(id)s"

SQL_INSERIR_RODADA = """
    INSERT INTO plano_escala_rodada
           (plano_id, rodada, vendas, preco, preco_a_mao, acos, parcela_ads, tacos,
            dias_esperados, margem_unit, receita, gasto_ads, resultado, acumulado)
    VALUES (%(plano_id)s, %(rodada)s, %(vendas)s, %(preco)s, %(preco_a_mao)s, %(acos)s,
            %(parcela_ads)s, %(tacos)s, %(dias_esperados)s, %(margem_unit)s, %(receita)s,
            %(gasto_ads)s, %(resultado)s, %(acumulado)s)
"""

SQL_INSERIR_DECISAO = """
    INSERT INTO plano_escala_decisao
           (plano_id, marketplace, loja, anuncio, tipo, rodada, motivo, criado_por)
    VALUES (%(plano_id)s, %(marketplace)s, %(loja)s, %(anuncio)s, %(tipo)s, %(rodada)s,
            %(motivo)s, %(usuario)s)
"""

USUARIO_SISTEMA = 'sistema'


def _ler(conn, sql, p):
    cur = conn.cursor()
    try:
        cur.execute(sql, p)
        return cur.fetchall()
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()


def ler_observado(conn, lojas, hoje):
    fim = fim_fechado(hoje)
    return agregar_observado(_ler(conn, SQL_OBSERVADO, {
        'marketplace': MARKETPLACE, 'lojas': list(lojas), 'barreira': BARREIRA_ML,
        'logisticas': list(LOGISTICAS_INSUMO), 'ini': fim - timedelta(days=JANELA_INSUMOS - 1),
        'fim': fim}))


def ler_produtos(conn, skus):
    if not skus:
        return {}
    return {s: {'nome': n, 'categoria': c, 'custo': float(p) if p is not None else None}
            for s, n, c, p in _ler(conn, SQL_PRODUTOS, {'skus': sorted(skus)})}


def ler_ativos(conn, lojas):
    cols = ('id', 'loja', 'anuncio', 'sku', 'venda_concorrente', 'meta_propria',
            'preco_concorrente', 'preco_meta', 'data_inicio', 'gestora', 'parametros',
            'situacao', 'como_encerrou', 'criado_por', 'criado_em', 'atualizado_por',
            'atualizado_em')
    planos = [dict(zip(cols, r)) for r in _ler(conn, SQL_ATIVOS,
                                                {'marketplace': MARKETPLACE, 'lojas': list(lojas)})]
    for p in planos:
        for c in ('preco_concorrente', 'preco_meta'):
            p[c] = float(p[c])
        par = p['parametros'] or {}
        inf = par.get('informado') or {}
        p['comissao'] = inf.get('comissao')
        p['frete_acima'] = inf.get('frete_acima')
        p['precos_mao'] = {int(k): float(v) for k, v in (par.get('precos_mao') or {}).items()}
        p['teto'] = float(par.get('teto') or 0)
    ids = [p['id'] for p in planos]
    rod, dec = {}, {}
    if ids:
        cols_r = ('plano_id', 'rodada', 'vendas', 'preco', 'preco_a_mao', 'acos', 'parcela_ads',
                  'tacos', 'dias_esperados', 'margem_unit', 'receita', 'gasto_ads', 'resultado',
                  'acumulado')
        for r in _ler(conn, SQL_RODADAS, {'ids': ids}):
            d = dict(zip(cols_r, r))
            for c in cols_r[3:]:
                if c != 'preco_a_mao':
                    d[c] = float(d[c])
            rod.setdefault(d['plano_id'], []).append(d)
        for r in _ler(conn, SQL_DECISOES, {'ids': ids}):
            dec.setdefault(r[0], []).append(dict(zip(
                ('plano_id', 'tipo', 'rodada', 'motivo', 'criado_por', 'criado_em'), r)))
    for p in planos:
        p['rodadas'] = rod.get(p['id'], [])
        p['decisoes'] = dec.get(p['id'], [])
    return planos


def ler_real(conn, planos, hoje):
    """{plano_id: série diária} de cada plano (até o último dia fechado)."""
    fim = fim_fechado(hoje)
    vivos = [p for p in planos if p['data_inicio'] <= fim]
    if not vivos:
        return {}
    par = {'marketplace': MARKETPLACE, 'lojas': sorted({p['loja'] for p in vivos}),
           'anuncios': sorted({p['anuncio'] for p in vivos}),
           'ini': min(p['data_inicio'] for p in vivos), 'fim': fim}
    vendas, ads = {}, {}
    for loja, an, d, q, rec, mg in _ler(conn, SQL_REAL_VENDAS, par):
        vendas.setdefault((loja, an), {})[d] = (int(q or 0), float(rec or 0), float(mg or 0))
    for loja, an, d, g, ua in _ler(conn, SQL_REAL_ADS, par):
        ads.setdefault((loja, an), {})[d] = (float(g or 0), int(ua or 0))
    return {p['id']: serie_diaria(p['data_inicio'], fim, vendas.get((p['loja'], p['anuncio']), {}),
                                  ads.get((p['loja'], p['anuncio']), {}))
            for p in vivos}


def calculadora(produtos, impostos, observado):
    """calcular(cad) para previa(), com os insumos lidos do banco."""
    def calcular(cad):
        prod = produtos.get(cad['sku'])
        if not prod:
            raise ValueError(f'SKU {cad["sku"]} não existe no cadastro de produtos.')
        if cad['loja'] not in impostos:
            raise ValueError(f'Loja {cad["loja"]} não é do Mercado Livre.')
        ins = resolver_insumos(cad['loja'], cad['anuncio'], cad['sku'], prod.get('categoria'),
                               observado, cad['comissao'], cad['frete_acima'])
        return calcular_plano(cad, prod, impostos[cad['loja']], ins)
    return calcular


def _parametros(item):
    cad, res = item['cad'], item['resumo']
    return {
        'informado': {'comissao': cad['comissao'], 'frete_acima': cad['frete_acima']},
        'precos_mao': {str(k): v for k, v in cad['precos_mao'].items()},
        'meta_mensal': res['meta_mensal'], 'teto': round(res['teto'], 2),
        'virada': res['virada'], 'resultado_total': round(res['resultado_total'], 2),
        'vendas_total': res['vendas_total'], 'preco_margem_zero': res['preco_margem_zero'],
        'insumos': res['insumos'], 'versao_regra': 'v1-2026-10-12',
    }


def gravar(conn, itens, usuario, lojas_permitidas):
    """Grava a prévia numa transação só (tudo ou nada). Devolve {acao: n}."""
    import json
    from psycopg2.extras import Json
    cont = {}
    cur = conn.cursor()
    try:
        for it in itens:
            if it['acao'] not in ('novo', 'revisao', 'cadastro'):
                continue
            cad, ativo = it['cad'], it['ativo']
            base = {'marketplace': MARKETPLACE, 'loja': cad['loja'], 'anuncio': cad['anuncio'],
                    'usuario': usuario, 'lojas': list(lojas_permitidas)}
            dados = {**base, **{c: cad[c] for c in ('sku', 'venda_concorrente', 'meta_propria',
                                                     'preco_concorrente', 'preco_meta',
                                                     'data_inicio', 'gestora')},
                     'parametros': Json(json.loads(json.dumps(_parametros(it), default=str)))}
            if it['acao'] == 'novo':
                if ativo and ativo.get('concluido'):
                    cur.execute(SQL_ENCERRAR, {**base, 'id': ativo['id'], 'como': 'concluido',
                                               'usuario': USUARIO_SISTEMA})
                    if cur.rowcount != 1:
                        raise RuntimeError('O plano concluído mudou enquanto você conferia.')
                    cur.execute(SQL_INSERIR_DECISAO, {**base, 'plano_id': ativo['id'],
                                                      'tipo': 'concluido', 'rodada': None,
                                                      'motivo': None, 'usuario': USUARIO_SISTEMA})
                cur.execute(SQL_INSERIR_PLANO, dados)
                plano_id = cur.fetchone()[0]
                tipo = 'cadastro'
            else:
                plano_id = ativo['id']
                cur.execute(SQL_ATUALIZAR_PLANO, {**dados, 'id': plano_id})
                if cur.rowcount != 1:
                    raise RuntimeError(f'O plano de {cad["loja"]} {cad["anuncio"]} mudou '
                                       'enquanto você conferia. Confira de novo.')
                cur.execute(SQL_APAGAR_RODADAS, {'id': plano_id})
                tipo = it['acao']
            for r in it['rodadas']:
                cur.execute(SQL_INSERIR_RODADA, {'plano_id': plano_id, **{
                    c: (round(r[c], 6) if isinstance(r[c], float) else r[c])
                    for c in ('rodada', 'vendas', 'preco', 'preco_a_mao', 'acos', 'parcela_ads',
                              'tacos', 'dias_esperados', 'margem_unit', 'receita', 'gasto_ads',
                              'resultado', 'acumulado')}})
            cur.execute(SQL_INSERIR_DECISAO, {**base, 'plano_id': plano_id, 'tipo': tipo,
                                              'rodada': None, 'motivo': cad['motivo'] or None})
            cont[it['acao']] = cont.get(it['acao'], 0) + 1
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
    return cont


def encerrar(conn, plano, motivo, usuario, lojas_permitidas, rodada=None):
    """A gestora encerra antes da hora, com motivo (seção 13)."""
    if not (motivo or '').strip():
        raise ValueError('Escreva o motivo do encerramento.')
    base = {'marketplace': MARKETPLACE, 'loja': plano['loja'], 'anuncio': plano['anuncio'],
            'lojas': list(lojas_permitidas)}
    cur = conn.cursor()
    try:
        cur.execute(SQL_ENCERRAR, {**base, 'id': plano['id'], 'como': 'gestora',
                                   'usuario': usuario})
        if cur.rowcount != 1:
            raise ValueError('Plano não está mais ativo (ou não é de uma loja sua).')
        cur.execute(SQL_INSERIR_DECISAO, {**base, 'plano_id': plano['id'], 'tipo': 'encerramento',
                                          'rodada': rodada, 'motivo': motivo.strip(),
                                          'usuario': usuario})
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()


# ============================================================
# TELA
# ============================================================

def nome_resumido(nome):
    nome = (nome or '(sem cadastro)').strip()
    return nome if len(nome) <= CHAR_NOME else nome[:CHAR_NOME - 1].rstrip() + '…'


def _brl(v):
    if v is None:
        return '—'
    s = f'{abs(v):,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    return f'−R$ {s}' if v < 0 else f'R$ {s}'


def _pct(v, casas=1):
    return '—' if v is None else f'{v * 100:.{casas}f}%'.replace('.', ',')


ICONE_SITUACAO = {'ativo': '🟢 ativo', 'pede revisão': '🔴 pede revisão',
                  'aguardando início': '⏳ aguardando início', 'concluído': '🏁 concluído',
                  'encerrado': '⚫ encerrado'}


def tabela_resumo(planos, acomp, nomes):
    linhas = []
    for p in planos:
        a = acomp[p['id']]
        pa, rit = a['prevista_atual'], a['ritmo']
        real_atual = next((r for r in a['reais'] if r['rodada'] == pa['rodada']), None)
        par = p['parametros'] or {}
        linhas.append({
            'Loja': p['loja'], 'Anúncio': p['anuncio'], 'SKU': p['sku'],
            'Produto': nome_resumido(nomes.get(p['sku'])), 'Gestora': p['gestora'],
            'Início': f"{p['data_inicio']:%d/%m/%Y}",
            'Situação': ICONE_SITUACAO.get(a['situacao'], a['situacao']),
            'Rodada': f"{pa['rodada']} de {len(p['rodadas'])}",
            'Preço da rodada': _brl(pa['preco']), 'ACOS-alvo': _pct(pa['acos'], 0),
            'Vendas da rodada': f"{real_atual['unid'] if real_atual else 0} de {pa['vendas']}",
            'Ritmo (real × esperado/dia)': (
                f"{rit['por_dia']:.1f} × {rit['esperado']:.1f}".replace('.', ',')
                + (' 🐢 lenta' if rit['lenta'] else '')) if rit else '—',
            'Acumulado real (margem − ads)': _brl(a['acumulado_real']),
            'TETO (buraco máx.)': _brl(par.get('teto')),
            'Virada (venda nº)': str(par['virada']) if par.get('virada') is not None else 'não vira',
            'Vendas desde o início': a['unid_total'],
            'Parcela do ads': ('⚠️ acima do previsto' if a['alerta_parcela'] else ''),
            'Recomendação': a['recomendacao'],
        })
    return pd.DataFrame(linhas)


def tabela_rodadas(p, a):
    reais = {r['rodada']: r for r in a['reais']}
    linhas = []
    for r in p['rodadas']:
        x = reais.get(r['rodada'])
        linhas.append({
            'Rodada': r['rodada'],
            'Preço': _brl(r['preco']) + (' ✋' if r['preco_a_mao'] else ''),
            'Vendas prev.': r['vendas'], 'Vendas real': str(x['unid']) if x else '',
            'Dias esp.': r['dias_esperados'],
            'Período real': (f"{x['inicio']:%d/%m}–" + (f"{x['fim']:%d/%m}" if x['fim'] else 'aberta'))
            if x else '',
            'ACOS-alvo': _pct(r['acos'], 0), 'Parcela ads prev.': _pct(r['parcela_ads'], 0),
            'Parcela ads real': _pct(x['parcela_ads'], 0) if x else '',
            'TACOS prev.': _pct(r['tacos'], 2),
            'Margem unit. prev.': _brl(r['margem_unit']),
            'Resultado prev.': _brl(r['resultado']),
            'Resultado real (margem − ads)': _brl(x['margem_real']) if x else '',
            'Gasto ads prev.': _brl(r['gasto_ads']), 'Gasto ads real': _brl(x['gasto']) if x else '',
            'Acumulado prev.': _brl(r['acumulado']),
        })
    return pd.DataFrame(linhas)


def texto_insumos(par):
    ins = (par or {}).get('insumos') or {}
    if not ins:
        return ''
    partes = [
        f"Custo {_brl(ins.get('custo'))} (cadastro de produtos)",
        f"Imposto {_pct(ins.get('imposto'))} (loja)",
        f"Comissão {_pct(ins.get('comissao'), 2)} — {ins.get('comissao_origem')}",
        f"Frete abaixo de R$ 79: "
        + (_brl(ins.get('frete_abaixo')) if ins.get('frete_abaixo') is not None else 'taxa fixa')
        + f" — {ins.get('frete_abaixo_origem')}",
        f"Frete acima de R$ 79: {_brl(ins.get('frete_acima'))} — {ins.get('frete_acima_origem')}",
    ]
    if ins.get('extra'):
        partes.append(f"Extra da loja {_brl(ins['extra'])}")
    partes.append(f"Preço de margem zero (referência): {_brl((par or {}).get('preco_margem_zero'))}")
    partes.append('Tarifa do Full fora da conta (v2).')
    return ' · '.join(partes)


def tabela_previa(itens, nomes):
    linhas = []
    for it in itens:
        cad, res = it['cad'] or {}, it['resumo'] or {}
        ins = res.get('insumos') or {}
        linhas.append({
            'Linha': str(it['n']), 'Ação': ROTULO_ACAO[it['acao']],
            'Loja': cad.get('loja', ''), 'Anúncio': cad.get('anuncio', ''),
            'SKU': cad.get('sku', ''), 'Produto': nome_resumido(nomes.get(cad.get('sku'))),
            'Rodadas': str(len(it['rodadas'])) if it['rodadas'] else '',
            'TETO': _brl(res.get('teto')) if res else '',
            'Virada (venda nº)': (str(res['virada']) if res.get('virada') is not None else 'não vira')
            if res else '',
            'Resultado do plano': _brl(res.get('resultado_total')) if res else '',
            'Margem zero (ref.)': _brl(res.get('preco_margem_zero')) if res else '',
            'Comissão usada': (f"{_pct(ins.get('comissao'), 2)} — {ins.get('comissao_origem')}"
                               if ins else ''),
            'Por que recusou': it['erro'],
            'Aviso': it.get('aviso', ''),
        })
    return pd.DataFrame(linhas)


def _usuario():
    import streamlit as st
    return (st.session_state.get('usuario') or {}).get('username', '')


def render(engine):
    import streamlit as st
    from permissoes import get_nivel_acesso
    from sinais_dia import restricao_de_lojas, lojas_visiveis
    import estoque_peca as ep

    st.title('🚀 Plano de escala')
    st.caption('Mercado Livre. Um plano por anúncio: rodadas de preço e volume até o preço meta, '
               'previsto × real, ritmo e TETO. O sistema avisa e recomenda; quem decide é a '
               'gestora. Margem antes da tarifa do Full.')
    hoje = datetime.now(BRT).date()
    pode_gravar = get_nivel_acesso('plano_escala') in ('completo', 'parcial')
    conn = engine.raw_connection()
    try:
        if not _ler(conn, SQL_EXISTE, {})[0][0]:
            st.warning('As tabelas do plano de escala ainda não foram criadas no banco '
                       '(sql/plano_escala.sql).')
            return
        lojas_imp = {l: float(i or 0) / 100 for l, i in _ler(conn, SQL_LOJAS_ML,
                                                              {'marketplace': MARKETPLACE})}
        lojas = lojas_visiveis(sorted(lojas_imp), restricao_de_lojas(engine))
        if not lojas:
            st.info('Nenhuma loja do Mercado Livre liberada para o seu usuário.')
            return
        planos = ler_ativos(conn, lojas)
        series = ler_real(conn, planos, hoje)
        nomes = ep.ler_nomes(conn, {p['sku'] for p in planos})
        acomp = {p['id']: acompanhar(p, p['rodadas'], series.get(p['id'], []), hoje)
                 for p in planos if p['rodadas']}
        planos = [p for p in planos if p['id'] in acomp]
    finally:
        conn.close()

    st.caption(f'Real até {fim_fechado(hoje):%d/%m/%Y} (último dia fechado).')
    st.subheader('Planos ativos')
    if planos:
        st.dataframe(tabela_resumo(planos, acomp, nomes), hide_index=True, use_container_width=True)
        for p in planos:
            a = acomp[p['id']]
            with st.expander(f"{p['loja']} · {p['anuncio']} · {p['sku']} "
                             f"{nome_resumido(nomes.get(p['sku']))} — "
                             f"{ICONE_SITUACAO.get(a['situacao'], a['situacao'])}"):
                st.dataframe(tabela_rodadas(p, a), hide_index=True, use_container_width=True)
                st.caption(texto_insumos(p['parametros']))
                if p['decisoes']:
                    st.markdown('**Decisões**')
                    st.dataframe(pd.DataFrame([{
                        'Quando': f"{pd.Timestamp(d['criado_em']):%d/%m/%Y %H:%M}",
                        'Quem': d['criado_por'], 'Tipo': d['tipo'], 'Motivo': d['motivo'] or ''}
                        for d in p['decisoes']]), hide_index=True, use_container_width=True)
    else:
        st.info('Nenhum plano ativo nas suas lojas.')

    if not pode_gravar:
        return

    st.subheader('Cadastro')
    st.caption('Uma porta só: baixe, edite e suba o .xlsx, OU edite a tabela abaixo; os dois '
               'passam pela mesma conferência. Subir de novo substitui o plano do anúncio. '
               'Plano já começado: motivo obrigatório e o previsto é recalculado inteiro.')
    atuais = [linha_do_plano(p, nomes) for p in planos]
    st.download_button('⬇️ Baixar o cadastro (.xlsx)', excel_cadastro(atuais),
                       file_name=f'plano_escala_{hoje:%Y%m%d}.xlsx',
                       mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    arq = st.file_uploader('Subir o cadastro preenchido (.xlsx)', type=['xlsx'],
                           key='plano_escala_upload')
    df_edit = pd.DataFrame(atuais, columns=[c for c, _ in COLUNAS]).rename(columns=dict(COLUNAS))
    editado = st.data_editor(df_edit, num_rows='dynamic', hide_index=True,
                             use_container_width=True, key='plano_escala_editor',
                             disabled=['Produto (não editar)'])
    c1, c2 = st.columns(2)
    origem = None
    if c1.button('Conferir o .xlsx', disabled=arq is None):
        origem = 'xlsx'
    if c2.button('Conferir a tabela'):
        origem = 'tabela'
    if origem:
        try:
            brutas = ler_excel(arq.getvalue()) if origem == 'xlsx' else linhas_do_editor(editado)
        except ValueError as e:
            st.error(str(e))
            brutas = None
        if brutas is not None:
            conn = engine.raw_connection()
            try:
                skus = {str(b.get('sku') or '').strip().upper() for b in brutas} - {''}
                produtos = ler_produtos(conn, skus)
                observado = ler_observado(conn, lojas, hoje)
            finally:
                conn.close()
            ativos = {(p['loja'], p['anuncio']): {**p, 'concluido': acomp[p['id']]['situacao']
                                                  == 'concluído'} for p in planos}
            st.session_state['plano_escala_previa'] = previa(
                brutas, ativos, set(lojas), hoje, calculadora(produtos, lojas_imp, observado))
            st.session_state['plano_escala_nomes'] = {s: v['nome'] for s, v in produtos.items()}
    itens = st.session_state.get('plano_escala_previa')
    if itens:
        st.dataframe(tabela_previa(itens, st.session_state.get('plano_escala_nomes', {})),
                     hide_index=True, use_container_width=True)
        a_gravar = [i for i in itens if i['acao'] in ('novo', 'revisao', 'cadastro')]
        recusados = sum(1 for i in itens if i['acao'] == 'recusado')
        if recusados:
            st.warning(f'{recusados} linha(s) recusada(s): corrija e confira de novo. '
                       'As outras podem ser gravadas.')
        if a_gravar and st.button(f'Gravar {len(a_gravar)} plano(s)', type='primary'):
            conn = engine.raw_connection()
            try:
                cont = gravar(conn, a_gravar, _usuario(), lojas)
                st.success('Gravado: ' + ', '.join(f'{ROTULO_ACAO[k]}: {v}' for k, v in cont.items()))
                st.session_state.pop('plano_escala_previa', None)
            except Exception as e:  # noqa: BLE001
                st.error(f'Nada foi gravado: {e}')
            finally:
                conn.close()

    if planos:
        st.subheader('Encerrar um plano antes da hora')
        rot = {f"{p['loja']} · {p['anuncio']} · {p['sku']}": p for p in planos}
        esc = st.selectbox('Plano', list(rot), key='plano_escala_encerrar')
        motivo = st.text_input('Motivo do encerramento (obrigatório)', key='plano_escala_motivo')
        if st.button('Encerrar plano'):
            p = rot[esc]
            conn = engine.raw_connection()
            try:
                encerrar(conn, p, motivo, _usuario(), lojas,
                         rodada=acomp[p['id']]['rodada_atual'])
                st.success('Plano encerrado; a decisão ficou registrada.')
            except ValueError as e:
                st.error(str(e))
            finally:
                conn.close()
