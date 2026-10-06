"""
SINAIS DO DIA — painel diário das gestoras (frente [SINAIS DO DIA], v1, 05/10/2026)

Menu "📡 Sinais do Dia". Uma aba por marketplace (v1 = Mercado Livre); dentro,
as lojas uma embaixo da outra. Por loja:
  (a) Resumo: venda de ontem × média do MESMO dia da semana nas 4 semanas
      anteriores; mês até ontem × meta (dim_metas_loja, aba Performance).
  (b) Sinais: até MAX_SINAIS_POR_LOJA por loja, ordenados por R$ em jogo.

Regras em SQL/Python determinísticos (sem IA), parâmetros ligados. A única
escrita é o "Ciente" (v1.1): tabela sinal_ciente, gravada pelo app. Plano aprovado pelo Thiago em 05/10/2026, com a calibragem
do [MESTRE ANÁLISES] (rodada manual na ML-LPT, 05/10):

  - A venda da API ainda entra depois: D-1 e D-2 ficam FORA das comparações
    (DIAS_ATRASO_API). Semana avaliada = D-9..D-3; base = as 4 semanas
    anteriores (D-37..D-10), comparada como MÉDIA SEMANAL (o sábado vende ~25%
    menos; semanas inteiras anulam o dia da semana; ciclo do mês não é modelado).
  - Dias de Full = (Full disponível + em transferência) ÷ venda/dia, venda/dia =
    max(7 dias ÷ 7, 30 dias ÷ 30), as duas janelas terminando em D-3: ritmo
    recente com piso na média de 30 (caso K-L-0421-A, que zerou antes do
    previsto).
  - Full POR estoque_id, NUNCA somado por SKU: o K-L-0421-A tem 2 anúncios na
    ML-LPT, um com 6 un. (o que vende) e outro com 89 (quase parado); somar
    esconde a ruptura. Anúncios que DIVIDEM o mesmo estoque_id somam a venda
    contra aquele Full (o estoque é o mesmo). A venda casa com o estoque por
    (loja, anúncio, SKU): unidade sempre a do SKU do anúncio (kit em kits).
  - "Espiral" (ads cortado + venda caindo) só acende se o anúncio TEVE estoque
    em todos os dias da semana (Full se o anúncio é Full, senão galpão > 0
    como sim/não — o galpão de kit vem dividido, nunca é somado). Sem estoque,
    a causa é ruptura e o sinal sai no bloco FULL (caso Garrafa LTT-CT3821,
    ML-Nala, ago/26).
  - Experiência por FAMÍLIA = mesma peça base pela composição dos kits
    (dim_kit_composicao), na mesma loja. "Todos os alicates" exige campo novo
    no cadastro: v2.
  - Custo por venda de ads = CPC ÷ conversão = gasto ÷ unidades de ads
    (units_quantity do coletor), comparado com a margem em R$ por unidade
    (margem_total ÷ quantidade, margem ANTES de mídia e Full). Só anúncio com
    UM SKU vendido na janela: anúncio com variações de SKU diferentes não
    divide R$ por unidades de produtos diferentes.
  - ACOS ≠ TACOS: a espiral mostra TACOS (gasto ÷ venda TOTAL do anúncio).
  - Estoque: só quantidade, nunca availability_type ("in_stock" vem com zero).
  - "O coletor rodou?" sai das datas de captura no banco de vendas (sem
    segredo novo para o banco do coletor).
  - Visitas: lê a VIEW vw_visitas_dia (frente [VISITAS ML]); enquanto ela não
    existir, "coleta de visitas em construção".

Reaproveita de estoque_peca.py só a leitura auxiliar (composição, mapeamento de
SKU, nomes, contagem de fotos parciais, atraso da foto); NÃO usa montar() (é
cobertura em peça da empresa inteira, outra pergunta) nem cobertura_full.py /
v_cobertura_full (não auditados).

Os limiares abaixo são CONSTANTES para a gestora e o [MESTRE ANÁLISES]
calibrarem nas 2 primeiras semanas.

CIENTE (v1.1, plano aprovado pelo Thiago em 05/10/2026)
  - A gestora responde ao sinal: motivo (falta no fornecedor / decisão
    proposital / já em andamento / outro), nota livre e "silenciar até"
    (padrão 7 dias, máx. 60). Grava em sinal_ciente (quem, quando, por quê):
    é a "intenção" que o [MESTRE ANÁLISES] lê para não recomendar o que foi
    decidido de propósito.
  - Chave = regra + loja + objeto (não o tipo: silenciar "Full em 5 dias" não
    pode calar "Full zerado com ads"). Objeto = estoque_id nos sinais de Full
    de estoque, MLB nos de anúncio, "peça:X" na família; evento de um dia só
    (mudança de ROAS/orçamento, envio que entrou) leva a data no objeto, para
    o ciente não calar o próximo evento.
  - O silenciado sai ANTES do corte de 10 e vai para "Silenciados". VOLTA
    sozinho quando passa a data ou quando PIORA (piorou(), limiares PIORA_*).
    Vencer e piorar são calculados na leitura: abrir a tela não grava nada.
    Só gente encerra um ciente (reativar, ou um ciente novo que substitui).
  - Quem dá ciente: nível 'completo' ou 'parcial' no módulo (ADMIN,
    CONTROLADORIA, COMPRAS; GESTOR só nas lojas dele). DIRETOR só vê.
  - Antes do SQL aplicado (sql/sinais_ciente.sql), a tela funciona como a v1.

v1.2 (06/10/2026)
  - Velocidade: o banco fica em São Paulo e o app no Streamlit Cloud; cada
    ida e volta custa ~130 ms e a v1 fazia ~39 por carga, a cada clique. Agora:
    a leitura do dia (tudo menos os cientes) fica em memória por
    TTL_LEITURA_S, com botão "Atualizar dados"; leitura incompleta (alguma
    fonte com erro) NÃO vai para a memória; os cientes são relidos a cada
    carga; rollback só no erro; cada loja é um st.fragment (marcar uma caixa
    não recarrega as outras); quadro "⏱ tempos desta carga" só para ADMIN.
  - "Venda de ontem" e "Mês até ontem" sem selo "parcial" (dia fechado), com
    "?" explicando o que ainda muda. A exclusão de D-1/D-2 nas comparações de
    VENDAS fica até a v1.3 medir (pedido pago depois pode entrar com data de
    ontem).
  - Coluna "Período" em cada sinal (janela analisada ou foto usada).
  - Excel de ida e volta: baixar os sinais com a chave e Status/Motivo/Nota/
    Até vazios; subir preenchido com PRÉVIA; grava só as linhas ok, numa
    transação, pela MESMA gravar_cientes do formulário. Medida e texto do
    sinal vêm do cálculo de hoje, nunca do arquivo. DIRETOR baixa, não sobe.

RESSALVAS CONHECIDAS (auditor-tecnico, 05/10/2026, aprovado com ressalvas)
  - R1: o mesmo SKU em DOIS estoques do mesmo anúncio conta a venda nos dois
    (a venda não diz de qual variação saiu). O alarme fica pessimista (cada
    estoque vê a venda inteira). Em 05/10 era 1 caso, fora do Full.
  - R2: "Full zerado com ads" usa o gasto do ANÚNCIO inteiro (a API de ads
    não separa por variação); o texto do sinal diz isso.
  - R3: cada bloco (vendas, Full, ads, experiência, visitas) é montado e
    cada loja é desenhada isoladamente: um dado inesperado derruba só aquele
    bloco/loja, com st.error próprio.
  - R4: o aviso de "coletor não rodou" assume que data_processamento,
    data_captura e data_importacao estão em horário de Brasília sem fuso (é
    como os coletores gravam hoje). Se um coletor passar a gravar em UTC, o
    aviso erra em 3 horas.
"""

import hashlib
import io
import re
import time
from datetime import date, datetime, timedelta, timezone

import pandas as pd

import estoque_peca as ep

MARKETPLACE = 'MERCADO LIVRE'

# ============================================================
# LIMIARES (calibrar aqui)
# ============================================================

DIAS_ATRASO_API = 2              # D-1 e D-2 ainda recebem venda da API
SEMANAS_BASE = 4                 # base = média semanal das 4 semanas anteriores
QUEDA_VENDA = 0.50               # semana < 50% da média semanal → queda
ALTA_VENDA = 2.0                 # semana > 2× a média semanal → alta
RECEITA_MIN_RELEVANTE = 1000.0   # R$ nas 4 semanas de base: anúncio "relevante"

JANELA_RITMO_FULL = 7            # venda recente (dias, terminando em D-3)
JANELA_PISO_FULL = 30            # piso: média de 30 dias (terminando em D-3)
FULL_URGENTE_DIAS = 7
FULL_ATENCAO_DIAS = 14

DIAS_GASTO_SEM_VENDA = 3         # dias seguidos com gasto e zero venda de ads
ESPIRAL_QUEDA_GASTO = 0.50       # gasto da semana < 50% da média semanal
ESPIRAL_QUEDA_VENDA = 0.80       # venda da semana < 80% da média semanal
MIN_UNID_ADS_CUSTO = 3           # unidades de ads (30d) p/ julgar custo por venda

DIAS_REF_EXPERIENCIA = 7         # nota de hoje × nota de 7 dias atrás
MIN_ANUNCIOS_FAMILIA = 2         # anúncios da mesma peça que pioraram

# Ciente: quando o silenciado volta antes da data (calibrar aqui)
PIORA_FATOR_RS = 1.5             # R$ em jogo / prejuízo por unidade cresceu 50%
PIORA_FATOR_GASTO = 2.0          # gasto de ads dobrou
PIORA_QUEDA_PP = 0.20            # queda aprofundou 20 pontos percentuais
DIAS_SILENCIO_PADRAO = 7
DIAS_SILENCIO_MAX = 60           # igual à CHECK sinal_ciente_prazo
DIAS_SELO_VENCIDO = 7            # mostra "voltou: venceu" por 7 dias

MOTIVOS = {'falta_fornecedor': 'Falta no fornecedor',
           'proposital': 'Decisão proposital',
           'em_andamento': 'Já em andamento',
           'outro': 'Outro'}

REGRAS = {'vendas_queda': 'Venda caiu', 'vendas_alta': 'Venda subiu',
          'full_cobertura': 'Full acabando', 'full_zerado_ads': 'Full zerado com ads',
          'full_envio': 'Envio entrou no Full', 'full_ruptura': 'Ruptura',
          'ads_config': 'ROAS/orçamento mudou', 'ads_sem_venda': 'Ads sem venda',
          'ads_custo': 'Custo de ads > margem', 'ads_espiral': 'Espiral de ads',
          'exp_anuncio': 'Experiência piorou', 'exp_familia': 'Experiência da família',
          'visitas_queda': 'Visitas caíram', 'visitas_alta': 'Visitas subiram'}

MAX_SINAIS_POR_LOJA = 10
TTL_LEITURA_S = 600              # leitura do dia em memória por 10 min
HORA_COLETA_COMPLETA = 10        # depois das 10h (Brasília) o dado do dia já devia ter chegado
DIAS_LEITURA_FRESCOR = 10

BRT = timezone(timedelta(hours=-3))


# ============================================================
# JANELAS
# ============================================================

def janelas(hoje):
    """Todas as datas que a tela usa, a partir de HOJE (Brasília)."""
    ontem = hoje - timedelta(days=1)
    fim = hoje - timedelta(days=DIAS_ATRASO_API + 1)                 # D-3
    sem_ini = fim - timedelta(days=6)                                # D-9
    base_fim = sem_ini - timedelta(days=1)                           # D-10
    base_ini = base_fim - timedelta(days=7 * SEMANAS_BASE - 1)       # D-37
    ritmo_ini = fim - timedelta(days=JANELA_RITMO_FULL - 1)
    piso_ini = fim - timedelta(days=JANELA_PISO_FULL - 1)
    mes_ini = ontem.replace(day=1)
    mesmo_dia = [ontem - timedelta(days=7 * k) for k in range(1, SEMANAS_BASE + 1)]
    return {
        'hoje': hoje, 'ontem': ontem, 'fim': fim,
        'sem_ini': sem_ini, 'base_ini': base_ini, 'base_fim': base_fim,
        'ritmo_ini': ritmo_ini, 'piso_ini': piso_ini,
        'mes_ini': mes_ini, 'mesmo_dia': mesmo_dia,
        'ini_resumo': min(mes_ini, mesmo_dia[-1]),
        'ini_vendas': min(base_ini, piso_ini),
        'recente_ini': hoje - timedelta(days=DIAS_GASTO_SEM_VENDA),
        'ini_curto': hoje - timedelta(days=DIAS_LEITURA_FRESCOR),
    }


# ============================================================
# SQL — parâmetros ligados (psycopg2); sem schema de propósito: o teste usa
# tabelas TEMPORÁRIAS com os mesmos nomes. Toda leitura filtra por data em
# coluna com índice e pela lista de lojas permitidas.
# ============================================================

SQL_LOJAS_ML = """
    SELECT loja FROM dim_lojas WHERE marketplace = %(marketplace)s ORDER BY loja
"""

SQL_RESUMO = """
    SELECT loja_origem AS loja,
           SUM(valor_venda_efetivo) FILTER (WHERE data_venda = %(ontem)s)          AS rec_ontem,
           SUM(valor_venda_efetivo) FILTER (WHERE data_venda = ANY(%(mesmo_dia)s)) AS rec_mesmo_dia,
           SUM(valor_venda_efetivo) FILTER (WHERE data_venda >= %(mes_ini)s)       AS rec_mes
      FROM fact_vendas_snapshot
     WHERE marketplace_origem = %(marketplace)s
       AND loja_origem = ANY(%(lojas)s)
       AND data_venda >= %(ini_resumo)s AND data_venda <= %(ontem)s
     GROUP BY loja_origem
"""

SQL_METAS = """
    SELECT loja_origem, meta_receita
      FROM dim_metas_loja
     WHERE ano_mes = %(ano_mes)s AND loja_origem = ANY(%(lojas)s)
"""

# Por (loja, anúncio, SKU): a unidade nunca mistura SKUs.
SQL_VENDAS_ANUNCIO = """
    SELECT loja_origem AS loja, codigo_anuncio AS anuncio, sku,
           SUM(valor_venda_efetivo) FILTER (WHERE data_venda >= %(sem_ini)s)   AS rec_sem,
           SUM(valor_venda_efetivo) FILTER (WHERE data_venda >= %(base_ini)s
                                              AND data_venda <= %(base_fim)s) AS rec_base,
           SUM(quantidade) FILTER (WHERE data_venda >= %(ritmo_ini)s)          AS qtd_ritmo,
           SUM(quantidade) FILTER (WHERE data_venda >= %(piso_ini)s)           AS qtd_30,
           SUM(valor_venda_efetivo) FILTER (WHERE data_venda >= %(piso_ini)s)  AS rec_30,
           SUM(margem_total) FILTER (WHERE data_venda >= %(piso_ini)s)         AS margem_30
      FROM fact_vendas_snapshot
     WHERE marketplace_origem = %(marketplace)s
       AND loja_origem = ANY(%(lojas)s)
       AND codigo_anuncio IS NOT NULL
       AND data_venda >= %(ini_vendas)s AND data_venda <= %(fim)s
     GROUP BY loja_origem, codigo_anuncio, sku
"""

SQL_PONTE = """
    SELECT DISTINCT loja, anuncio_id, estoque_id, sku, COALESCE(em_full, false) AS em_full
      FROM dim_estoque_anuncio
     WHERE marketplace = %(marketplace)s AND loja = ANY(%(lojas)s)
       AND estoque_id IS NOT NULL AND anuncio_id IS NOT NULL
"""

# Última foto de cada loja, um estoque por linha.
SQL_ESTOQUE_FOTO = """
    WITH ult AS (
        SELECT loja, max(data) AS data
          FROM fact_estoque_diario
         WHERE marketplace = %(marketplace)s AND loja = ANY(%(lojas)s)
           AND data >= %(ini_curto)s
         GROUP BY loja)
    SELECT f.loja, f.estoque_id, f.data,
           COALESCE(f.full_disponivel, 0)       AS full_disponivel,
           COALESCE(f.full_em_transferencia, 0) AS full_em_transferencia,
           COALESCE(f.unidades_recebidas, 0)    AS unidades_recebidas
      FROM fact_estoque_diario f
      JOIN ult u ON u.loja = f.loja AND u.data = f.data
     WHERE f.marketplace = %(marketplace)s
"""

# Estoque dia a dia na semana avaliada (para a espiral).
SQL_ESTOQUE_SEMANA = """
    SELECT loja, estoque_id, data,
           COALESCE(full_disponivel, 0) AS full_disponivel, galpao_disponivel
      FROM fact_estoque_diario
     WHERE marketplace = %(marketplace)s AND loja = ANY(%(lojas)s)
       AND data >= %(sem_ini)s AND data <= %(fim)s
"""

SQL_ADS = """
    SELECT loja, codigo_anuncio AS anuncio,
           SUM(gasto_ads) FILTER (WHERE data >= %(sem_ini)s AND data <= %(fim)s)      AS gasto_sem,
           SUM(gasto_ads) FILTER (WHERE data >= %(base_ini)s AND data <= %(base_fim)s) AS gasto_base,
           SUM(gasto_ads) FILTER (WHERE data = %(ontem)s)                             AS gasto_ontem,
           SUM(gasto_ads) FILTER (WHERE data >= %(recente_ini)s)                      AS gasto_recente,
           COUNT(DISTINCT data) FILTER (WHERE data >= %(recente_ini)s
                                          AND gasto_ads > 0)                          AS dias_com_gasto,
           SUM(vendas) FILTER (WHERE data >= %(recente_ini)s)                         AS unid_recente,
           SUM(gasto_ads) FILTER (WHERE data >= %(piso_ini)s AND data <= %(fim)s)     AS gasto_30,
           SUM(vendas) FILTER (WHERE data >= %(piso_ini)s AND data <= %(fim)s)        AS unid_ads_30,
           SUM(cliques) FILTER (WHERE data >= %(piso_ini)s AND data <= %(fim)s)       AS cliques_30
      FROM fact_ads_performance
     WHERE marketplace = %(marketplace)s AND loja = ANY(%(lojas)s)
       AND codigo_anuncio IS NOT NULL
       AND data >= %(ini_vendas)s AND data <= %(ontem)s
     GROUP BY loja, codigo_anuncio
"""

# Foto de config de hoje × a foto anterior de cada campanha. A foto sai às
# ~05h30: o que a gestora mudou ONTEM aparece na foto de HOJE.
SQL_CONFIG_MUDOU = """
    WITH c AS (
        SELECT DISTINCT ON (loja, id_campanha, data_captura::date)
               loja, id_campanha, campanha, data_captura::date AS dia,
               roas_objetivo, orcamento_diario
          FROM fact_ads_campanha_config
         WHERE marketplace = %(marketplace)s AND loja = ANY(%(lojas)s)
           AND data_captura >= %(ini_curto)s
         ORDER BY loja, id_campanha, data_captura::date, data_captura DESC),
    r AS (
        SELECT c.*, row_number() OVER (PARTITION BY loja, id_campanha ORDER BY dia DESC) AS k
          FROM c)
    SELECT a.loja, a.id_campanha, a.campanha, b.dia AS dia_antes, a.dia AS dia_agora,
           b.roas_objetivo AS roas_antes, a.roas_objetivo AS roas_agora,
           b.orcamento_diario AS orc_antes, a.orcamento_diario AS orc_agora
      FROM r a
      JOIN r b ON b.loja = a.loja AND b.id_campanha = a.id_campanha AND b.k = 2
     WHERE a.k = 1 AND a.dia = %(hoje)s
       AND (a.roas_objetivo IS DISTINCT FROM b.roas_objetivo
            OR a.orcamento_diario IS DISTINCT FROM b.orcamento_diario)
"""

SQL_EXPERIENCIA = """
    WITH d AS (
        SELECT loja, max(data) AS data
          FROM fact_saude_anuncio
         WHERE marketplace = %(marketplace)s AND loja = ANY(%(lojas)s)
           AND tipo_sinal = 'experiencia_compra'
           AND data >= %(ini_curto)s AND data <= %(hoje)s
         GROUP BY loja)
    SELECT s.loja, s.codigo_anuncio, s.data, s.valor, s.texto,
           r.valor AS valor_ref, r.texto AS texto_ref
      FROM fact_saude_anuncio s
      JOIN d ON d.loja = s.loja AND d.data = s.data
      LEFT JOIN fact_saude_anuncio r
             ON r.marketplace = s.marketplace AND r.loja = s.loja
            AND r.codigo_anuncio = s.codigo_anuncio AND r.tipo_sinal = s.tipo_sinal
            AND r.data = s.data - %(dias_ref)s
     WHERE s.marketplace = %(marketplace)s AND s.tipo_sinal = 'experiencia_compra'
       AND s.data >= %(ini_curto)s AND s.codigo_anuncio IS NOT NULL
"""

SQL_EXISTE_VISITAS = "SELECT to_regclass('vw_visitas_dia') IS NOT NULL"

SQL_VISITAS = """
    SELECT loja, codigo_anuncio AS anuncio,
           SUM(visitas) FILTER (WHERE data >= %(sem_ini)s)                       AS vis_sem,
           SUM(visitas) FILTER (WHERE data >= %(base_ini)s AND data <= %(base_fim)s) AS vis_base
      FROM vw_visitas_dia
     WHERE marketplace = %(marketplace)s AND loja = ANY(%(lojas)s)
       AND data >= %(base_ini)s AND data <= %(fim)s
     GROUP BY loja, codigo_anuncio
"""

# Última chegada de cada fonte, por loja (horário de Brasília, sem fuso).
SQL_FRESCOR = """
    SELECT 'Vendas' AS fonte, loja_origem AS loja, max(data_processamento) AS ultima
      FROM fact_vendas_snapshot
     WHERE marketplace_origem = %(marketplace)s AND loja_origem = ANY(%(lojas)s)
       AND data_venda >= %(ini_curto)s
     GROUP BY loja_origem
    UNION ALL
    SELECT 'Estoque', loja, max(data_captura)
      FROM fact_estoque_diario
     WHERE marketplace = %(marketplace)s AND loja = ANY(%(lojas)s) AND data >= %(ini_curto)s
     GROUP BY loja
    UNION ALL
    SELECT 'Ads', loja, max(data_importacao)
      FROM fact_ads_performance
     WHERE marketplace = %(marketplace)s AND loja = ANY(%(lojas)s) AND data >= %(ini_curto)s
     GROUP BY loja
    UNION ALL
    SELECT 'Config. de ads', loja, max(data_captura)
      FROM fact_ads_campanha_config
     WHERE marketplace = %(marketplace)s AND loja = ANY(%(lojas)s)
       AND data_captura >= %(ini_curto)s
     GROUP BY loja
    UNION ALL
    SELECT 'Experiência', loja, max(data_importacao)
      FROM fact_saude_anuncio
     WHERE marketplace = %(marketplace)s AND loja = ANY(%(lojas)s) AND data >= %(ini_curto)s
     GROUP BY loja
"""

# ---- Ciente (v1.1): a única escrita do módulo ----------------------------
SQL_EXISTE_CIENTE = "SELECT to_regclass('sinal_ciente') IS NOT NULL"

# Abertos (não encerrados por gente). Vencer e piorar são julgados no Python.
SQL_CIENTES_ABERTOS = """
    SELECT id, loja, regra, objeto, motivo, nota, silenciar_ate, medida_no_ciente,
           texto_no_ciente, criado_por, criado_em
      FROM sinal_ciente
     WHERE marketplace = %(marketplace)s AND loja = ANY(%(lojas)s)
       AND encerrado_em IS NULL
"""

SQL_FECHAR_ABERTO = """
    UPDATE sinal_ciente
       SET encerrado_em = (now() AT TIME ZONE 'America/Sao_Paulo'),
           encerrado_por = %(usuario)s, como_encerrou = 'substituido'
     WHERE marketplace = %(marketplace)s AND loja = %(loja)s AND regra = %(regra)s
       AND objeto = %(objeto)s AND encerrado_em IS NULL
"""

SQL_INSERIR_CIENTE = """
    INSERT INTO sinal_ciente
           (marketplace, loja, regra, objeto, motivo, nota, silenciar_ate,
            medida_no_ciente, texto_no_ciente, criado_por)
    VALUES (%(marketplace)s, %(loja)s, %(regra)s, %(objeto)s, %(motivo)s, %(nota)s,
            %(silenciar_ate)s, %(medida)s, %(texto)s, %(usuario)s)
"""

SQL_REATIVAR = """
    UPDATE sinal_ciente
       SET encerrado_em = (now() AT TIME ZONE 'America/Sao_Paulo'),
           encerrado_por = %(usuario)s, como_encerrou = 'reativado'
     WHERE id = %(id)s AND marketplace = %(marketplace)s AND loja = ANY(%(lojas)s)
       AND encerrado_em IS NULL
"""

FONTES_FRESCOR = ('Vendas', 'Estoque', 'Ads', 'Config. de ads', 'Experiência')

TODAS_AS_SQL = (SQL_LOJAS_ML, SQL_RESUMO, SQL_METAS, SQL_VENDAS_ANUNCIO, SQL_PONTE,
                SQL_ESTOQUE_FOTO, SQL_ESTOQUE_SEMANA, SQL_ADS, SQL_CONFIG_MUDOU,
                SQL_EXPERIENCIA, SQL_EXISTE_VISITAS, SQL_VISITAS, SQL_FRESCOR,
                SQL_EXISTE_CIENTE, SQL_CIENTES_ABERTOS, SQL_FECHAR_ABERTO,
                SQL_INSERIR_CIENTE, SQL_REATIVAR)


def params(hoje, lojas):
    p = janelas(hoje)
    p.update({'marketplace': MARKETPLACE, 'lojas': list(lojas),
              'dias_ref': DIAS_REF_EXPERIENCIA,
              'ano_mes': p['ontem'].strftime('%Y-%m')})
    return p


# ============================================================
# LEITURA (só SELECT; cada fonte isolada: uma que falhe não derruba as outras)
# ============================================================

def _ler(conn, sql, p):
    """Uma consulta. Rollback SÓ no erro (a transação abortada não pode
    derrubar as leituras seguintes); no sucesso não gasta a ida e volta: a
    conexão volta ao pool com rollback no close."""
    cur = conn.cursor()
    try:
        cur.execute(sql, p)
        return cur.fetchall()
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()


FONTES = {
    'resumo': SQL_RESUMO,
    'metas': SQL_METAS,
    'vendas': SQL_VENDAS_ANUNCIO,
    'ponte': SQL_PONTE,
    'foto': SQL_ESTOQUE_FOTO,
    'estoque_semana': SQL_ESTOQUE_SEMANA,
    'ads': SQL_ADS,
    'config': SQL_CONFIG_MUDOU,
    'experiencia': SQL_EXPERIENCIA,
    'frescor': SQL_FRESCOR,
}


def ler_cientes(conn, lojas):
    """Cientes abertos das lojas; None se a tabela não existe. Erro sobe."""
    if not _ler(conn, SQL_EXISTE_CIENTE, {})[0][0]:
        return None
    return _ler(conn, SQL_CIENTES_ABERTOS, {'marketplace': MARKETPLACE, 'lojas': list(lojas)})


def ler_tudo(conn, hoje, lojas):
    """{fonte: linhas} e {fonte: exceção}: a leitura do DIA (vai para a
    memória). Visitas só se a view existir (None = ainda não existe). Os
    cientes NÃO entram aqui: mudam a cada clique (ler_cientes)."""
    p = params(hoje, lojas)
    dados, erros = {}, {}
    for nome, sql in FONTES.items():
        try:
            dados[nome] = _ler(conn, sql, p)
        except Exception as e:  # noqa: BLE001 — vira aviso por bloco na tela
            erros[nome] = e
    try:
        dados['composicao'] = _ler(conn, ep.SQL_COMPOSICAO, {})
        dados['mapa'] = dict(_ler(conn, ep.SQL_MAPEAMENTO, {}))
    except Exception as e:  # noqa: BLE001
        erros['composicao'] = e
    try:
        if _ler(conn, SQL_EXISTE_VISITAS, {})[0][0]:
            dados['visitas'] = _ler(conn, SQL_VISITAS, p)
        else:
            dados['visitas'] = None
    except Exception as e:  # noqa: BLE001
        erros['visitas'] = e
    return dados, erros


# ============================================================
# CONTAS (sem banco)
# ============================================================

def _f(v):
    return float(v) if v is not None else 0.0


def _brl(v):
    return ep._brl(v)


def _dec(v, casas=1):
    return f'{v:.{casas}f}'.replace('.', ',')


def _pct(v):
    return f'{v * 100:+.0f}%'.replace('.', ',')


def anuncio_da_campanha(id_campanha, campanha=None):
    """No ML cada campanha é um anúncio só: 'MLB4639959197 Torx + Allen'."""
    for t in (id_campanha, campanha):
        m = re.match(r'\s*(MLB\d+)', t or '')
        if m:
            return m.group(1)
    return None


def objeto_campanha(anuncio, campanha, dia):
    """Objeto do ciente de mudança de ROAS/orçamento: MLB@data; campanha sem
    MLB vira camp:<hash curto do nome>@data (nome comprido não pode estourar o
    varchar(120) e derrubar o lote)."""
    data = f"{pd.Timestamp(dia):%Y-%m-%d}"
    if anuncio:
        return f"{anuncio}@{data}"
    h = hashlib.sha1((campanha or '').encode('utf-8')).hexdigest()[:12]
    return f"camp:{h}@{data}"


def agregar_vendas(vendas):
    """[(loja, anuncio, sku, rec_sem, rec_base, qtd_ritmo, qtd_30, rec_30, margem_30)]
    -> (por_anuncio, por_anuncio_sku).
    por_anuncio[(loja, anuncio)]: R$ somados (R$ pode somar), SKUs, e
    unidades/margem SÓ quando o anúncio vendeu um SKU único."""
    por_sku = {}
    an = {}
    for loja, anuncio, sku, rs, rb, qr, q30, r30, m30 in vendas:
        por_sku[(loja, anuncio, sku)] = {'qtd_ritmo': _f(qr), 'qtd_30': _f(q30),
                                         'rec_30': _f(r30)}
        a = an.setdefault((loja, anuncio), {'rec_sem': 0.0, 'rec_base': 0.0, 'rec_30': 0.0,
                                            'skus': set(), 'qtd_30': 0.0, 'margem_30': None})
        a['rec_sem'] += _f(rs)
        a['rec_base'] += _f(rb)
        a['rec_30'] += _f(r30)
        if _f(q30) > 0 or _f(rs) > 0 or _f(rb) > 0:
            a['skus'].add(sku)
        if _f(q30) > 0:
            a['qtd_30'] += _f(q30)
            if m30 is not None:
                a['margem_30'] = (a['margem_30'] or 0.0) + float(m30)
    for a in an.values():
        a['sku_unico'] = len(a['skus']) == 1
        if not a['sku_unico']:
            a['qtd_30'], a['margem_30'] = None, None
    return an, por_sku


def _d(x):
    return f"{pd.Timestamp(x):%d/%m}"


def _periodo_semanas(j):
    return f"{_d(j['sem_ini'])}–{_d(j['fim'])} × {_d(j['base_ini'])}–{_d(j['base_fim'])}"


def _sinal(loja, tipo, anuncio, skus, numero, sugestao, em_jogo, urgente=False,
           familia='', regra='', objeto=None, medida=None, periodo=''):
    """regra + loja + objeto = a chave do Ciente; medida = o número que diz se
    o sinal piorou depois do ciente (ver piorou())."""
    return {'loja': loja, 'tipo': tipo, 'anuncio': anuncio or '',
            'skus': tuple(sorted(skus or ())), 'numero': numero,
            'sugestao': sugestao, 'em_jogo': float(em_jogo or 0),
            'urgente': urgente, 'familia': familia, 'regra': regra,
            'objeto': objeto if objeto is not None else (anuncio or ''),
            'medida': None if medida is None else float(medida), 'voltou': '',
            'periodo': periodo}


def _rel(a):
    return a['rec_base'] >= RECEITA_MIN_RELEVANTE


def sinais_vendas(por_anuncio, j):
    out = []
    for (loja, anuncio), a in por_anuncio.items():
        if not _rel(a):
            continue
        base_sem = a['rec_base'] / SEMANAS_BASE
        sem = a['rec_sem']
        periodo = f"{j['sem_ini']:%d/%m}–{j['fim']:%d/%m}"
        if sem < QUEDA_VENDA * base_sem:
            out.append(_sinal(
                loja, 'VENDAS', anuncio, a['skus'],
                f"Queda: semana {periodo} {_brl(sem)} × média {_brl(base_sem)}/sem "
                f"({_pct(sem / base_sem - 1)})",
                "Ver Full/estoque, preço, posição e se o ads foi mexido",
                base_sem - sem, regra='vendas_queda', medida=sem / base_sem - 1,
                periodo=_periodo_semanas(j)))
        elif sem > ALTA_VENDA * base_sem:
            out.append(_sinal(
                loja, 'VENDAS', anuncio, a['skus'],
                f"Alta: semana {periodo} {_brl(sem)} × média {_brl(base_sem)}/sem "
                f"({_pct(sem / base_sem - 1)})",
                "Garantir Full para o novo ritmo; campanha não travar no orçamento",
                sem - base_sem, regra='vendas_alta', medida=sem / base_sem - 1,
                periodo=_periodo_semanas(j)))
    return out


def _ponte_mapeada(ponte, mapa):
    """[(loja, anuncio, estoque_id, sku, em_full)] com o SKU corrigido
    (dim_sku_mapeamento, um salto, igual às vendas)."""
    return [(l, a, e, mapa.get(s, s) if s else s, bool(f)) for l, a, e, s, f in ponte]


def estoques(ponte):
    """{(loja, estoque_id): {'anuncios': set, 'skus': set, 'em_full': bool}}"""
    out = {}
    for loja, anuncio, eid, sku, em_full in ponte:
        e = out.setdefault((loja, eid), {'anuncios': set(), 'skus': set(), 'em_full': False})
        e['anuncios'].add(anuncio)
        if sku:
            e['skus'].add(sku)
        e['em_full'] |= em_full
    return out


def sinais_full(foto, ponte, por_sku, ads_por_anuncio, j):
    """Cobertura do Full POR estoque_id. foto: [(loja, eid, data, full, transf,
    recebidas)]; ponte já mapeada."""
    est = estoques(ponte)
    out = []
    for loja, eid, data, full, transf, recebidas in foto:
        e = est.get((loja, eid))
        if not e or not e['em_full']:
            continue
        # venda do estoque: cada anúncio dele × cada SKU dele (mesmo produto)
        vend = {}
        for a in e['anuncios']:
            for s in e['skus']:
                v = por_sku.get((loja, a, s))
                if v:
                    vend[a] = vend.get(a, {'qtd_ritmo': 0.0, 'qtd_30': 0.0, 'rec_30': 0.0})
                    for k in vend[a]:
                        vend[a][k] += v[k]
        qtd_ritmo = sum(v['qtd_ritmo'] for v in vend.values())
        qtd_30 = sum(v['qtd_30'] for v in vend.values())
        rec_30 = sum(v['rec_30'] for v in vend.values())
        dia = max(qtd_ritmo / JANELA_RITMO_FULL, qtd_30 / JANELA_PISO_FULL)
        # o anúncio que vende (o de maior venda em 30d), senão qualquer um
        anuncio = (max(vend, key=lambda a: vend[a]['qtd_30']) if vend
                   else sorted(e['anuncios'])[0])
        ads_ontem = sum(_f(ads_por_anuncio.get((loja, a), {}).get('gasto_ontem'))
                        for a in e['anuncios'])
        full, transf, recebidas = int(full), int(transf), int(recebidas)
        if full == 0 and ads_ontem > 0:
            out.append(_sinal(
                loja, 'FULL', anuncio, e['skus'],
                f"Full zerado com ads ligado (gasto ontem {_brl(ads_ontem)} = ads do "
                f"anúncio, todas as variações"
                + (f"; {transf} em transferência" if transf else '') + ")",
                "Baixar/pausar o ads até o Full voltar", max(rec_30, ads_ontem),
                urgente=True, regra='full_zerado_ads', objeto=eid, medida=ads_ontem,
                periodo=f"foto {_d(data)} · ads {_d(j['ontem'])}"))
        elif dia > 0:
            dias = (full + transf) / dia
            if dias < FULL_ATENCAO_DIAS:
                urg = dias < FULL_URGENTE_DIAS
                out.append(_sinal(
                    loja, 'FULL', anuncio, e['skus'],
                    f"{dias:.0f} dias de Full ({full} disp. + {transf} em transf. ÷ "
                    f"{_dec(dia)}/dia; 7d {_dec(qtd_ritmo / JANELA_RITMO_FULL)}/dia, "
                    f"30d {_dec(qtd_30 / JANELA_PISO_FULL)}/dia)",
                    "Enviar ao Full agora" if urg else "Programar envio ao Full",
                    rec_30, urgente=urg, regra='full_cobertura', objeto=eid, medida=full,
                    periodo=f"foto {_d(data)} · venda {_d(j['ritmo_ini'])}–{_d(j['fim'])}"))
        if recebidas > 0:
            out.append(_sinal(
                loja, 'FULL', anuncio, e['skus'],
                f"Envio entrou no Full em {pd.Timestamp(data):%d/%m}: +{recebidas} un. "
                f"(Full agora {full})",
                "Conferir se o ads/preço voltou ao normal", rec_30, regra='full_envio',
                objeto=f"{eid}@{pd.Timestamp(data):%Y-%m-%d}", medida=recebidas,
                periodo=f"foto {_d(data)}"))
    return out


def dias_com_estoque(estoque_semana, ponte):
    """{(loja, anuncio): nº de dias da semana com estoque}. Anúncio Full:
    full_disponivel > 0; fora do Full: galpão > 0 (só sim/não, nunca somado)."""
    por_est = {}
    for loja, anuncio, eid, _sku, em_full in ponte:
        por_est.setdefault((loja, eid), []).append((anuncio, em_full))
    dias = {}
    for loja, eid, data, full, galpao in estoque_semana:
        for anuncio, em_full in por_est.get((loja, eid), ()):
            tem = (full or 0) > 0 if em_full else (galpao or 0) > 0
            if tem:
                dias.setdefault((loja, anuncio), set()).add(data)
    return {k: len(v) for k, v in dias.items()}


def sinais_ads(ads, config, por_anuncio, dias_estoque, j, com_estoque=True):
    """ads: {(loja, anuncio): dict}; config: linhas de SQL_CONFIG_MUDOU.
    com_estoque=False (a leitura do estoque da semana falhou): a espiral não é
    julgada, porque sem estoque não dá para separar espiral de ruptura."""
    out = []
    for loja, idc, camp, d_antes, d_agora, r0, r1, o0, o1 in config:
        anuncio = anuncio_da_campanha(idc, camp)
        a = por_anuncio.get((loja, anuncio), {})
        partes = []
        if r0 != r1:
            partes.append(f"ROAS objetivo {_f(r0):.1f} → {_f(r1):.1f}".replace('.', ','))
        if o0 != o1:
            partes.append(f"orçamento {_brl(_f(o0))} → {_brl(_f(o1))}/dia")
        out.append(_sinal(
            loja, 'ADS', anuncio or camp, a.get('skus'),
            "Mudou ontem: " + '; '.join(partes) + f" (foto de {pd.Timestamp(d_antes):%d/%m} → hoje)",
            "Acompanhar venda e TACOS nos próximos 3 dias", a.get('rec_30', 0.0),
            regra='ads_config', objeto=objeto_campanha(anuncio, camp, d_agora),
            periodo=f"foto {_d(d_antes)} → {_d(d_agora)}"))

    semana = 7
    for (loja, anuncio), d in ads.items():
        a = por_anuncio.get((loja, anuncio), {'rec_sem': 0.0, 'rec_base': 0.0,
                                               'rec_30': 0.0, 'skus': set(),
                                               'qtd_30': None, 'margem_30': None,
                                               'sku_unico': False})
        # gasto com zero venda de ads há N dias
        if int(d.get('dias_com_gasto') or 0) >= DIAS_GASTO_SEM_VENDA and _f(d.get('unid_recente')) == 0:
            out.append(_sinal(
                loja, 'ADS', anuncio, a['skus'],
                f"{DIAS_GASTO_SEM_VENDA} dias com gasto e zero venda de ads "
                f"({_brl(_f(d.get('gasto_recente')))})",
                "Rever ROAS objetivo ou pausar a campanha", _f(d.get('gasto_recente')),
                regra='ads_sem_venda', medida=_f(d.get('gasto_recente')),
                periodo=f"{_d(j['recente_ini'])}–{_d(j['ontem'])}"))

        # custo por venda acima da margem por unidade (só SKU único)
        unid, gasto, cliques = _f(d.get('unid_ads_30')), _f(d.get('gasto_30')), _f(d.get('cliques_30'))
        if (a.get('sku_unico') and a.get('qtd_30') and a.get('margem_30') is not None
                and unid >= MIN_UNID_ADS_CUSTO):
            custo_un = gasto / unid
            margem_un = a['margem_30'] / a['qtd_30']
            if custo_un > margem_un:
                cpc = gasto / cliques if cliques else None
                conv = unid / cliques if cliques else None
                txt = (f"Custo por venda de ads {_brl(custo_un)}/un"
                       + (f" (CPC {_brl(cpc)} ÷ conversão {conv * 100:.1f}%)".replace('.', ',')
                          if cpc is not None else '')
                       + f" × margem {_brl(margem_un)}/un (30d)")
                out.append(_sinal(loja, 'ADS', anuncio, a['skus'], txt,
                                  "Subir ROAS objetivo (cada venda de ads dá prejuízo)",
                                  (custo_un - margem_un) * unid, regra='ads_custo',
                                  medida=custo_un - margem_un,
                                  periodo=f"{_d(j['piso_ini'])}–{_d(j['fim'])}"))

        # espiral: gasto cortado E venda caindo — só com estoque a semana toda
        gasto_base_sem = _f(d.get('gasto_base')) / SEMANAS_BASE
        rec_base_sem = a['rec_base'] / SEMANAS_BASE
        if (com_estoque and gasto_base_sem > 0 and _rel(a)
                and _f(d.get('gasto_sem')) < ESPIRAL_QUEDA_GASTO * gasto_base_sem
                and a['rec_sem'] < ESPIRAL_QUEDA_VENDA * rec_base_sem):
            tacos_sem = _f(d.get('gasto_sem')) / a['rec_sem'] if a['rec_sem'] else None
            tacos_base = gasto_base_sem / rec_base_sem if rec_base_sem else None
            dias = dias_estoque.get((loja, anuncio), 0)
            queda = rec_base_sem - a['rec_sem']
            if dias >= semana:
                out.append(_sinal(
                    loja, 'ADS', anuncio, a['skus'],
                    f"Espiral: gasto de ads {_brl(_f(d.get('gasto_sem')))} × média "
                    f"{_brl(gasto_base_sem)}/sem e venda {_pct(a['rec_sem'] / rec_base_sem - 1)}"
                    + (f"; TACOS {tacos_sem * 100:.1f}% (era {tacos_base * 100:.1f}%)".replace('.', ',')
                       if tacos_sem is not None and tacos_base is not None else ''),
                    "Com estoque a semana toda: devolver o investimento aos poucos (+20–30%/ciclo)",
                    queda, regra='ads_espiral', medida=queda, periodo=_periodo_semanas(j)))
            else:
                out.append(_sinal(
                    loja, 'FULL', anuncio, a['skus'],
                    f"Ruptura: estoque em só {dias} de 7 dias da semana; venda "
                    f"{_pct(a['rec_sem'] / rec_base_sem - 1)} e ads cortado — a causa é estoque",
                    "Repor estoque antes de mexer no ads", queda, urgente=True,
                    regra='full_ruptura', medida=queda, periodo=_periodo_semanas(j)))
    return out


def familias(ponte, comp):
    """{(loja, anuncio): set(peças)} pela composição (kit → peças; o resto é
    ele mesmo)."""
    out = {}
    for loja, anuncio, _eid, sku, _f_ in ponte:
        if sku:
            out.setdefault((loja, anuncio), set()).update(p for p, _q in ep.pecas_do_sku(sku, comp))
    return out


def sinais_experiencia(saude, ponte, comp, por_anuncio):
    """saude: [(loja, anuncio, data, valor, texto, valor_ref, texto_ref)]."""
    fam = familias(ponte, comp)
    nota = {(l, a): (v, t, vr, tr) for l, a, _dt, v, t, vr, tr in saude}
    data_nota = {(l, a): dt for l, a, dt, _v, _t, _vr, _tr in saude}

    def periodo(k):
        dt = data_nota.get(k)
        return (f"nota {_d(pd.Timestamp(dt) - timedelta(days=DIAS_REF_EXPERIENCIA))} → {_d(dt)}"
                if dt is not None else '')

    def piorou(k):
        v, _t, vr, _tr = nota[k]
        return v is not None and vr is not None and float(v) < float(vr)

    # membros de cada família (mesma loja, mesma peça) que têm nota
    membros = {}
    for k in nota:
        for peca in fam.get(k, ()):
            membros.setdefault((k[0], peca), set()).add(k)

    def principal(loja, peca, ks):
        # o anúncio da própria peça; senão o de maior receita
        proprios = [k for k in ks if por_anuncio.get(k, {}).get('skus') == {peca}]
        cand = proprios or list(ks)
        return max(cand, key=lambda k: (por_anuncio.get(k, {}).get('rec_30', 0.0), k[1]))

    def texto_familia(loja, peca):
        ks = membros[(loja, peca)]
        if len(ks) < 2:
            return ''
        distintas = sorted({float(nota[k][0]) for k in ks if nota[k][0] is not None})
        p = principal(loja, peca, ks)
        v_p = nota[p][0]
        return (f"{peca}: {len(ks)} anúncios, notas {', '.join(f'{n:.0f}' for n in distintas)}"
                f" · principal {p[1]} em {'—' if v_p is None else f'{float(v_p):.0f}'}")

    out, cobertos = [], set()
    for (loja, peca), ks in sorted(membros.items(), key=lambda x: -len(x[1])):
        pioraram = sorted(k for k in ks if piorou(k) and k not in cobertos)
        if len(pioraram) >= MIN_ANUNCIOS_FAMILIA:
            cobertos |= set(pioraram)
            em_jogo = sum(por_anuncio.get(k, {}).get('rec_30', 0.0) for k in pioraram)
            de_para = '; '.join(f"{k[1]} {float(nota[k][2]):.0f}→{float(nota[k][0]):.0f}"
                                for k in pioraram)
            out.append(_sinal(
                loja, 'EXPERIÊNCIA', ', '.join(k[1] for k in pioraram), {peca},
                f"{len(pioraram)} anúncios da mesma peça pioraram em "
                f"{DIAS_REF_EXPERIENCIA} dias: {de_para}",
                "Ver o que muda entre os anúncios (kit, embalagem, reclamação, prazo)",
                em_jogo, familia=texto_familia(loja, peca), regra='exp_familia',
                objeto=f"peça:{peca}", medida=len(pioraram), periodo=periodo(pioraram[0])))
    for k in sorted(nota):
        if k in cobertos or not piorou(k):
            continue
        v, t, vr, tr = nota[k]
        a = por_anuncio.get(k, {})
        textos = [texto_familia(k[0], p) for p in sorted(fam.get(k, ()))
                  if (k[0], p) in membros]
        out.append(_sinal(
            k[0], 'EXPERIÊNCIA', k[1], a.get('skus') or set(),
            f"Experiência de compra {float(vr):.0f} ({tr or '—'}) → {float(v):.0f} ({t or '—'}) "
            f"em {DIAS_REF_EXPERIENCIA} dias",
            "Abrir o anúncio no ML e ver a ação principal sugerida",
            a.get('rec_30', 0.0), familia=' | '.join(x for x in textos if x),
            regra='exp_anuncio', medida=float(v), periodo=periodo(k)))
    return out


def sinais_visitas(visitas, por_anuncio, j):
    out = []
    for loja, anuncio, vs, vb in visitas:
        a = por_anuncio.get((loja, anuncio))
        if not a or not _rel(a):
            continue
        base_sem = _f(vb) / SEMANAS_BASE
        sem = _f(vs)
        if base_sem <= 0:
            continue
        if sem < QUEDA_VENDA * base_sem:
            txt, sug, regra = ("Queda de visitas",
                               "Ver posição, preço, ads e se o anúncio foi pausado", 'visitas_queda')
        elif sem > ALTA_VENDA * base_sem:
            txt, sug, regra = ("Alta de visitas",
                               "Ver se a conversão acompanha (preço, estoque)", 'visitas_alta')
        else:
            continue
        out.append(_sinal(
            loja, 'VISITAS', anuncio, a['skus'],
            f"{txt}: semana {sem:.0f} × média {base_sem:.0f}/sem ({_pct(sem / base_sem - 1)})",
            sug, a['rec_base'] / SEMANAS_BASE, regra=regra, medida=sem / base_sem - 1,
            periodo=_periodo_semanas(j)))
    return out


def ads_por_anuncio(linhas):
    cols = ('gasto_sem', 'gasto_base', 'gasto_ontem', 'gasto_recente', 'dias_com_gasto',
            'unid_recente', 'gasto_30', 'unid_ads_30', 'cliques_30')
    return {(r[0], r[1]): dict(zip(cols, r[2:])) for r in linhas}


def _bloco(nome, erros, padrao, f, *args, **kw):
    """Roda uma parte da montagem; se quebrar, guarda o erro em `erros[nome]`
    e devolve `padrao`. Um bloco quebrado não derruba os outros (R3)."""
    try:
        return f(*args, **kw)
    except Exception as e:  # noqa: BLE001 — vira st.error do bloco na tela
        erros[nome] = e
        return padrao


def montar_sinais(dados, hoje):
    """(sinais, erros_por_bloco), a partir do que ler_tudo leu. Fonte que não
    veio (erro de leitura) só deixa de gerar os seus sinais; bloco que quebra
    na conta aparece em erros_por_bloco e os outros seguem."""
    j = janelas(hoje)
    erros = {}
    mapa = dados.get('mapa') or {}
    comp = _bloco('experiencia', erros, {}, ep.agrupar_composicao,
                  dados.get('composicao') or [])
    por_anuncio, por_sku = _bloco('vendas', erros, ({}, {}), agregar_vendas,
                                  dados.get('vendas') or [])
    ponte = _bloco('full', erros, [], _ponte_mapeada, dados.get('ponte') or [], mapa)
    ads = _bloco('ads', erros, {}, ads_por_anuncio, dados.get('ads') or [])

    sinais = []
    if 'vendas' in dados and 'vendas' not in erros:
        sinais += _bloco('vendas', erros, [], sinais_vendas, por_anuncio, j)
    if 'foto' in dados and 'ponte' in dados and 'vendas' in dados and 'full' not in erros:
        sinais += _bloco('full', erros, [], sinais_full, dados['foto'], ponte, por_sku, ads, j)
    if ('ads' in dados or 'config' in dados) and 'ads' not in erros:
        dias_est = _bloco('ads', erros, {}, dias_com_estoque,
                          dados.get('estoque_semana') or [], ponte)
        sinais += _bloco('ads', erros, [], sinais_ads, ads, dados.get('config') or [],
                         por_anuncio, dias_est, j,
                         com_estoque='estoque_semana' in dados and 'ponte' in dados)
    if 'experiencia' in dados and 'experiencia' not in erros:
        sinais += _bloco('experiencia', erros, [], sinais_experiencia,
                         dados['experiencia'], ponte, comp, por_anuncio)
    if dados.get('visitas'):
        sinais += _bloco('visitas', erros, [], sinais_visitas, dados['visitas'], por_anuncio, j)

    # Espiral/ruptura explicam a queda: o sinal genérico de VENDAS do mesmo
    # anúncio sai, para não ocupar duas linhas das 10.
    explicados = {(s['loja'], s['anuncio']) for s in sinais
                  if s['regra'] in ('ads_espiral', 'full_ruptura')}
    sinais = [s for s in sinais
              if not (s['regra'] == 'vendas_queda'
                      and (s['loja'], s['anuncio']) in explicados)]
    return sinais, erros


def piorou(regra, agora, antes):
    """O sinal silenciado piorou desde o ciente? (medida de hoje × medida no
    ciente). Alta, envio que entrou e mudança de ROAS/orçamento nunca
    "pioram": só voltam pela data."""
    if agora is None or antes is None:
        return False
    agora, antes = float(agora), float(antes)
    if regra == 'full_cobertura':                       # un. no Full: zerou
        return antes > 0 and agora <= 0
    if regra in ('full_zerado_ads', 'ads_sem_venda'):   # gasto de ads
        return agora > antes and agora >= PIORA_FATOR_GASTO * antes
    if regra in ('full_ruptura', 'ads_espiral', 'ads_custo'):  # R$
        return agora > antes and agora >= PIORA_FATOR_RS * antes
    if regra in ('vendas_queda', 'visitas_queda'):      # variação (−0,55 = −55%)
        return agora <= antes - PIORA_QUEDA_PP or (agora <= -1.0 < antes)
    if regra == 'exp_anuncio':                          # nota
        # CALIBRAR ([MESTRE ANÁLISES], auditor 05/10): hoje QUALQUER queda de
        # nota reativa (75 -> 65 já volta). Pode virar "caiu de faixa".
        return agora < antes
    if regra == 'exp_familia':                          # nº de anúncios piorados
        return agora > antes
    return False


def texto_medida(regra, v):
    if v is None:
        return '—'
    v = float(v)
    if regra == 'full_cobertura':
        return f"{int(v)} un. no Full"
    if regra in ('vendas_queda', 'visitas_queda', 'vendas_alta', 'visitas_alta'):
        return _pct(v)
    if regra == 'exp_anuncio':
        return f"nota {v:.0f}"
    if regra == 'exp_familia':
        return f"{int(v)} anúncios"
    if regra == 'full_envio':
        return f"{int(v)} un."
    return _brl(v)


CIENTE_COLS = ('id', 'loja', 'regra', 'objeto', 'motivo', 'nota', 'silenciar_ate',
               'medida_no_ciente', 'texto_no_ciente', 'criado_por', 'criado_em')


def aplicar_cientes(sinais, cientes, hoje):
    """(ativos, silenciados). Sinal com ciente aberto e dentro da data sai dos
    ativos; volta com selo quando a data passou ou quando piorou."""
    por_chave = {}
    for linha in cientes or []:
        c = dict(zip(CIENTE_COLS, linha))
        por_chave[(c['loja'], c['regra'], c['objeto'])] = c
    ativos, silenciados = [], []
    for s in sinais:
        c = por_chave.get((s['loja'], s['regra'], s['objeto']))
        if c is None:
            ativos.append(s)
            continue
        ate = pd.Timestamp(c['silenciar_ate']).date()
        quando = f"{pd.Timestamp(c['criado_em']):%d/%m}"
        if piorou(s['regra'], s['medida'], c['medida_no_ciente']):
            ativos.append({**s, 'ciente_voltou': c['id'], 'voltou': (
                f"🔁 piorou: {texto_medida(s['regra'], c['medida_no_ciente'])} → "
                f"{texto_medida(s['regra'], s['medida'])} (ciente de {quando}, "
                f"{MOTIVOS.get(c['motivo'], c['motivo'])})")})
        elif ate < hoje:
            selo = ''
            if (hoje - ate).days <= DIAS_SELO_VENCIDO:
                selo = (f"🔁 silêncio venceu em {ate:%d/%m} "
                        f"({MOTIVOS.get(c['motivo'], c['motivo'])})")
            ativos.append({**s, 'voltou': selo})
        else:
            silenciados.append({**s, 'ciente': c})
    return ativos, silenciados


def validar_ciente(motivo, nota, ate, hoje):
    erros = []
    if motivo not in MOTIVOS:
        erros.append("Escolha o motivo.")
    if motivo == 'outro' and not (nota or '').strip():
        erros.append('Motivo "Outro" pede uma nota.')
    if ate is None or ate < hoje:
        erros.append("A data de silêncio não pode ser no passado.")
    elif ate > hoje + timedelta(days=DIAS_SILENCIO_MAX):
        erros.append(f"Silêncio de no máximo {DIAS_SILENCIO_MAX} dias.")
    return erros


def gravar_cientes(conn, itens, usuario, hoje, lojas_permitidas):
    """A ÚNICA gravação de ciente (formulário da tela e Excel).
    itens = [(sinal, motivo, nota, ate)]: cada um com o seu motivo/nota/data.
    Valida tudo antes de tocar no banco; grava numa transação só (ciente
    aberto na mesma chave é encerrado como 'substituido'). Medida e texto vêm
    do SINAL (cálculo de hoje), nunca de quem chama. Devolve quantos gravou."""
    erros = []
    if not (usuario or '').strip():
        erros.append("Usuário não identificado.")
    for s, motivo, nota, ate in itens:
        for e in validar_ciente(motivo, nota, ate, hoje):
            erros.append(f"{s['loja']} · {s['objeto']}: {e}")
        if len(s['objeto']) > 120:
            erros.append(f"Chave do sinal longa demais: {s['objeto'][:40]}…")
    if erros:
        raise ValueError(' '.join(erros))
    fora = sorted({s['loja'] for s, *_r in itens if s['loja'] not in lojas_permitidas})
    if fora:
        raise PermissionError(f"Sem permissão nas lojas: {', '.join(fora)}")
    cur = conn.cursor()
    try:
        for s, motivo, nota, ate in itens:
            p = {'marketplace': MARKETPLACE, 'loja': s['loja'], 'regra': s['regra'],
                 'objeto': s['objeto'], 'motivo': motivo,
                 'nota': (nota or '').strip() or None, 'silenciar_ate': ate,
                 'medida': s['medida'], 'texto': s['numero'], 'usuario': usuario}
            cur.execute(SQL_FECHAR_ABERTO, p)
            cur.execute(SQL_INSERIR_CIENTE, p)
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
    return len(itens)


# ============================================================
# EXCEL DE IDA E VOLTA (v1.2)
# ============================================================

COLUNAS_EXCEL = ('Loja', 'Regra', 'Objeto', 'Tipo', 'O que é', 'SKU', 'Produto', 'Anúncio',
                 'O que disparou', 'Período', 'R$ em jogo',
                 'Status', 'Motivo', 'Nota', 'Silenciar até')
EDITAVEIS_EXCEL = ('Status', 'Motivo', 'Nota', 'Silenciar até')
OBRIGATORIAS_EXCEL = ('Loja', 'Regra', 'Objeto') + EDITAVEIS_EXCEL


def excel_sinais(sinais, nomes, hoje):
    """xlsx com os sinais: a chave (Loja, Regra, Objeto) e os dados do sinal
    travados; Status/Motivo/Nota/Silenciar até livres, com lista suspensa."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = 'Sinais'
    ws.append(list(COLUNAS_EXCEL))
    for sn in sinais:
        produto = ' / '.join(nomes.get(x) or '(sem cadastro)' for x in sn['skus'])
        ws.append([sn['loja'], sn['regra'], sn['objeto'], sn['tipo'],
                   REGRAS.get(sn['regra'], sn['regra']), ', '.join(sn['skus']), produto,
                   sn['anuncio'], sn['numero'], sn.get('periodo', ''),
                   round(sn['em_jogo'], 2), None, None, None, None])
    ultima = max(len(sinais) + 1, 2)
    col = {c: get_column_letter(i + 1) for i, c in enumerate(COLUNAS_EXCEL)}
    dv_status = DataValidation(type='list', formula1='"Ciente"', allow_blank=True)
    dv_motivo = DataValidation(type='list', formula1='"' + ','.join(MOTIVOS.values()) + '"',
                               allow_blank=True)
    ws.add_data_validation(dv_status)
    ws.add_data_validation(dv_motivo)
    dv_status.add(f"{col['Status']}2:{col['Status']}{ultima}")
    dv_motivo.add(f"{col['Motivo']}2:{col['Motivo']}{ultima}")
    livre = PatternFill('solid', fgColor='FFF2CC')
    for nome in EDITAVEIS_EXCEL:
        for r in range(2, ultima + 1):
            c = ws[f"{col[nome]}{r}"]
            c.protection = Protection(locked=False)
            c.fill = livre
            if nome == 'Silenciar até':
                c.number_format = 'DD/MM/YYYY'
    for r in range(2, ultima + 1):
        ws[f"{col['R$ em jogo']}{r}"].number_format = '#,##0.00'
    for c in ws[1]:
        c.font = Font(bold=True)
    larguras = {'Loja': 12, 'Regra': 16, 'Objeto': 22, 'Tipo': 12, 'O que é': 22, 'SKU': 16,
                'Produto': 30, 'Anúncio': 16, 'O que disparou': 60, 'Período': 24,
                'R$ em jogo': 12, 'Status': 10, 'Motivo': 22, 'Nota': 30, 'Silenciar até': 14}
    for nome, w in larguras.items():
        ws.column_dimensions[col[nome]].width = w
    ws.freeze_panes = 'A2'
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.autoFilter = False
    ws.auto_filter.ref = f"A1:{col['Silenciar até']}{ultima}"

    ins = wb.create_sheet('Instruções')
    for linha in (
            [f"Sinais do Dia — Mercado Livre — gerado em {hoje:%d/%m/%Y}"], [],
            ["Preencha só as colunas amarelas (Status, Motivo, Nota, Silenciar até)."],
            ['Status: "Ciente" para responder ao sinal; vazio = linha ignorada.'],
            ["Motivo: " + ' / '.join(MOTIVOS.values()) + '. "Outro" pede Nota.'],
            [f"Silenciar até: dd/mm/aaaa, de hoje até {DIAS_SILENCIO_MAX} dias; vazio = "
             f"{DIAS_SILENCIO_PADRAO} dias."],
            ["Não mude Loja, Regra e Objeto: são a chave do sinal."],
            ["Ao subir, a tela mostra uma PRÉVIA e grava só as linhas ok. Sinal que não "
             "existe mais hoje é recusado."]):
        ins.append(linha)
    ins.column_dimensions['A'].width = 100
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def ler_excel(conteudo):
    """[{coluna: valor, '_linha': n}] das linhas não vazias da aba Sinais."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    ws = wb['Sinais'] if 'Sinais' in wb.sheetnames else wb.active
    linhas = list(ws.iter_rows(values_only=True))
    if not linhas:
        raise ValueError("Planilha vazia.")
    cab = [str(c).strip() if c is not None else '' for c in linhas[0]]
    faltam = [c for c in OBRIGATORIAS_EXCEL if c not in cab]
    if faltam:
        raise ValueError("Faltam as colunas: " + ', '.join(faltam))
    out = []
    for n, r in enumerate(linhas[1:], start=2):
        if all(v is None or str(v).strip() == '' for v in r):
            continue
        d = dict(zip(cab, r))
        d['_linha'] = n
        out.append(d)
    return out


def _motivo_excel(v):
    t = str(v or '').strip().lower()
    for cod, rot in MOTIVOS.items():
        if t in (cod, rot.lower()):
            return cod
    return None


def _data_excel(v, hoje):
    if v is None or str(v).strip() == '':
        return hoje + timedelta(days=DIAS_SILENCIO_PADRAO)
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = str(v).strip()
    for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d/%m/%y'):
        try:
            return datetime.strptime(t, fmt).date()
        except ValueError:
            pass
    raise ValueError(t)


def previa_excel(linhas, sinais_hoje, lojas_permitidas, hoje):
    """(itens para gravar_cientes, DataFrame da prévia). Só linhas com Status.
    O sinal (medida, texto) vem de `sinais_hoje`, nunca do arquivo."""
    por_chave = {(x['loja'], x['regra'], x['objeto']): x for x in sinais_hoje}
    itens, previa, vistas = [], [], {}
    for d in linhas:
        loja, regra, objeto = (str(d.get(k) or '').strip() for k in ('Loja', 'Regra', 'Objeto'))
        status = str(d.get('Status') or '').strip()
        if not status:
            continue
        base = {'Linha': d['_linha'], 'Loja': loja, 'Regra': REGRAS.get(regra, regra),
                'Objeto': objeto, 'Motivo': '', 'Até': ''}

        def recusa(motivo_recusa):
            previa.append({**base, 'Resultado': '❌ ' + motivo_recusa})

        if status.lower() != 'ciente':
            recusa('Status deve ser "Ciente" (ou vazio para ignorar a linha)')
            continue
        if loja not in lojas_permitidas:
            recusa('loja fora do seu perfil')
            continue
        chave = (loja, regra, objeto)
        if chave in vistas:
            recusa(f'sinal repetido no arquivo (já na linha {vistas[chave]})')
            continue
        vistas[chave] = d['_linha']
        sinal = por_chave.get(chave)
        if sinal is None:
            recusa('o sinal não existe mais hoje')
            continue
        motivo = _motivo_excel(d.get('Motivo'))
        try:
            ate = _data_excel(d.get('Silenciar até'), hoje)
        except ValueError:
            recusa('"Silenciar até" ilegível (use dd/mm/aaaa)')
            continue
        nota = str(d.get('Nota') or '').strip()
        erros = validar_ciente(motivo, nota, ate, hoje)
        if erros:
            recusa(' '.join(erros))
            continue
        itens.append((sinal, motivo, nota, ate))
        previa.append({**base, 'Motivo': MOTIVOS[motivo], 'Até': f"{ate:%d/%m/%Y}",
                       'Resultado': '✅ ok'})
    return itens, pd.DataFrame(previa, columns=['Linha', 'Loja', 'Regra', 'Objeto', 'Motivo',
                                                'Até', 'Resultado'])


def reativar_ciente(conn, id_ciente, usuario, lojas_permitidas):
    """Encerra um ciente ('reativado'): o sinal volta a aparecer. Só nas lojas
    do usuário. Devolve quantas linhas mudou (0 ou 1)."""
    if not (usuario or '').strip():
        raise ValueError("Usuário não identificado.")
    cur = conn.cursor()
    try:
        cur.execute(SQL_REATIVAR, {'id': id_ciente, 'usuario': usuario,
                                   'marketplace': MARKETPLACE,
                                   'lojas': list(lojas_permitidas)})
        n = cur.rowcount
        conn.commit()
        return n
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()


def pode_dar_ciente():
    from permissoes import get_nivel_acesso
    return get_nivel_acesso('sinais') in ('completo', 'parcial')


def sinais_da_loja(sinais, loja, limite=MAX_SINAIS_POR_LOJA):
    """(os `limite` primeiros por R$ em jogo, o resto). Urgente desempata."""
    da_loja = sorted((s for s in sinais if s['loja'] == loja),
                     key=lambda s: (-s['em_jogo'], not s['urgente'], s['tipo'], s['anuncio']))
    return da_loja[:limite], da_loja[limite:]


def resumo_lojas(linhas_resumo, metas, lojas, hoje):
    j = janelas(hoje)
    por = {r[0]: r for r in linhas_resumo}
    meta = {l: _f(m) for l, m in metas}
    import calendar
    dias_mes = calendar.monthrange(j['ontem'].year, j['ontem'].month)[1]
    out = {}
    for loja in lojas:
        _l, ontem, mesmo_dia, mes = por.get(loja, (loja, None, None, None))
        media = _f(mesmo_dia) / SEMANAS_BASE
        m = meta.get(loja)
        out[loja] = {
            'ontem': _f(ontem),
            'media_mesmo_dia': media,
            'var_ontem': (_f(ontem) / media - 1) if media > 0 else None,
            'mes': _f(mes),
            'meta': m if m else None,
            'pct_meta': (_f(mes) / m) if m else None,
            'esperado_linear': (m * j['ontem'].day / dias_mes) if m else None,
        }
    return out


def avaliar_frescor(linhas, lojas, agora):
    """[(nível, texto)] por loja: 'atrasado' quando a fonte não chegou hoje e já
    passou da HORA_COLETA_COMPLETA; 'cedo' antes disso. Fonte sem nenhum dado
    em 10 dias (ex.: loja sem ads) não acusa atraso."""
    ult = {(l, f): u for f, l, u in linhas}
    out = {}
    for loja in lojas:
        avisos = []
        for fonte in FONTES_FRESCOR:
            u = ult.get((loja, fonte))
            if u is None:
                continue
            if pd.Timestamp(u).date() < agora.date():
                nivel = 'atrasado' if agora.hour >= HORA_COLETA_COMPLETA else 'cedo'
                avisos.append((nivel, f"{fonte}: último dado {pd.Timestamp(u):%d/%m %H:%M}"))
        out[loja] = avisos
    return out


# ============================================================
# PERMISSÃO
# ============================================================

def restricao_de_lojas(engine):
    """None = vê todas as lojas; lista = só estas. Lista VAZIA = usuário
    restrito sem loja cadastrada: NÃO vê nada (o bug conhecido do card aberto,
    em que lista vazia vira "sem restrição", não se repete aqui)."""
    from permissoes import ve_todas_lojas, get_lojas_usuario
    if ve_todas_lojas():
        return None
    return list(get_lojas_usuario(engine) or [])


def lojas_visiveis(lojas_ml, restricao):
    return list(lojas_ml) if restricao is None else [l for l in lojas_ml if l in restricao]


# ============================================================
# TELA
# ============================================================

ROTULO_FONTE = {
    'resumo': 'resumo de vendas', 'metas': 'metas', 'vendas': 'vendas por anúncio',
    'ponte': 'anúncios × estoque', 'foto': 'foto do Full', 'estoque_semana': 'estoque da semana',
    'ads': 'ads', 'config': 'mudanças de ROAS/orçamento', 'experiencia': 'experiência de compra',
    'frescor': 'datas de chegada', 'composicao': 'composição de kits', 'visitas': 'visitas',
    'ciente': 'cientes registrados',
}

ROTULO_BLOCO = {'vendas': 'VENDAS', 'full': 'FULL', 'ads': 'ADS',
                'experiencia': 'EXPERIÊNCIA', 'visitas': 'VISITAS', 'ciente': 'CIENTE'}

ICONE = {'VENDAS': '📉', 'FULL': '📦', 'ADS': '📣', 'EXPERIÊNCIA': '⭐', 'VISITAS': '👀'}


def tabela_sinais(sinais, nomes):
    def produto(skus):
        return ' / '.join(nomes.get(s) or '(sem cadastro)' for s in skus) if skus else ''
    return pd.DataFrame({
        'Loja': [s['loja'] for s in sinais],
        'Tipo': [('🔴 ' if s['urgente'] else '') + ICONE.get(s['tipo'], '') + ' ' + s['tipo']
                 for s in sinais],
        'SKU': [', '.join(s['skus']) for s in sinais],
        'Produto': [produto(s['skus']) for s in sinais],
        'Anúncio': [s['anuncio'] for s in sinais],
        'O que disparou': [s['numero'] for s in sinais],
        'Período': [s.get('periodo', '') for s in sinais],
        'Sugestão': [s['sugestao'] for s in sinais],
        'R$ em jogo': [_brl(s['em_jogo']) for s in sinais],
        'Família (mesma peça, mesma loja)': [s['familia'] for s in sinais],
        'Situação': [s.get('voltou', '') for s in sinais],
    })


def cientes_silenciando(cientes_loja, ativos, hoje):
    """Cientes que estão calando algo hoje: abertos, dentro da data e cujo
    sinal NÃO voltou por piora (esse já está nos ativos com o selo)."""
    voltaram = {x.get('ciente_voltou') for x in ativos}
    return [c for c in cientes_loja
            if pd.Timestamp(c['silenciar_ate']).date() >= hoje and c['id'] not in voltaram]


def tabela_silenciados(silenciados, cientes_loja, hoje, nomes, ativos=()):
    """Cientes que estão calando algo, com o motivo à vista; diz se o sinal
    ainda dispara hoje."""
    disparando = {(x['regra'], x['objeto']) for x in silenciados}
    linhas = []
    for c in cientes_silenciando(cientes_loja, ativos, hoje):
        linhas.append({
            'Regra': REGRAS.get(c['regra'], c['regra']),
            'Objeto': c['objeto'],
            'Motivo': MOTIVOS.get(c['motivo'], c['motivo']),
            'Nota': c['nota'] or '',
            'Silenciado até': f"{pd.Timestamp(c['silenciar_ate']):%d/%m/%Y}",
            'Quem': c['criado_por'],
            'Quando': f"{pd.Timestamp(c['criado_em']):%d/%m %H:%M}",
            'Hoje': ('ainda dispara' if (c['regra'], c['objeto']) in disparando
                     else 'não disparou hoje'),
            'No ciente': c['texto_no_ciente'] or '',
        })
    return pd.DataFrame(linhas)


AJUDA_DIA_FECHADO = ("Pode subir com pagamentos aprovados depois (boleto/Pix) e cair "
                     "com cancelamentos e devoluções.")


def _render_resumo(st, loja, r, ontem):
    c1, c2, c3 = st.columns(3)
    c1.metric(f"Venda de ontem ({ontem:%d/%m})", _brl(r['ontem']),
              None if r['var_ontem'] is None else _pct(r['var_ontem']) + " vs mesmo dia da semana",
              help=f"Dia fechado. {AJUDA_DIA_FECHADO} Comparada com a média dos "
                   f"{SEMANAS_BASE} mesmos dias da semana anteriores "
                   f"({_brl(r['media_mesmo_dia'])}).")
    c2.metric("Mês até ontem", _brl(r['mes']), help=f"Dias fechados. {AJUDA_DIA_FECHADO}")
    if r['meta']:
        c3.metric("Meta do mês", _brl(r['meta']),
                  f"{r['pct_meta'] * 100:.0f}% atingido · esperado linear "
                  f"{_brl(r['esperado_linear'])}", delta_color='off')
    else:
        c3.metric("Meta do mês", "não cadastrada", help="Cadastrar na aba Performance.")


def _no_streamlit():
    """True só rodando dentro do Streamlit (nos testes: sem memória, sem fragmento)."""
    try:
        from streamlit.runtime import exists
        return exists()
    except Exception:  # noqa: BLE001
        return False


class _LeituraIncompleta(Exception):
    """Leitura com alguma fonte em erro: devolvida, mas NÃO guardada em memória."""

    def __init__(self, pacote):
        super().__init__('leitura incompleta')
        self.pacote = pacote


def _ler_lojas_ml(_engine):
    conn = _engine.raw_connection()
    try:
        return [r[0] for r in _ler(conn, SQL_LOJAS_ML, {'marketplace': MARKETPLACE})]
    finally:
        conn.close()


def _ler_pacote(_engine, hoje, lojas):
    """A leitura do DIA (tudo menos os cientes) + os sinais montados."""
    t0 = time.perf_counter()
    conn = _engine.raw_connection()
    try:
        dados, erros = ler_tudo(conn, hoje, list(lojas))
        parciais = []
        try:
            parciais = [x for x in ep.fotos_parciais(ep.ler_contagens(conn)) if x[0] in lojas]
        except Exception:  # noqa: BLE001
            pass
        sinais, erros_bloco = montar_sinais(dados, hoje)
        try:
            nomes = ep.ler_nomes(conn, {x for sn in sinais for x in sn['skus']})
        except Exception:  # noqa: BLE001
            nomes = {}
    finally:
        conn.close()
    return {'dados': dados, 'erros': {k: str(v) for k, v in erros.items()},
            'erros_bloco': {k: str(v) for k, v in erros_bloco.items()},
            'parciais': parciais, 'sinais': sinais, 'nomes': nomes,
            'lido_em': datetime.now(BRT), 'segundos': time.perf_counter() - t0}


def _ler_pacote_completo(_engine, hoje, lojas):
    pacote = _ler_pacote(_engine, hoje, lojas)
    if pacote['erros'] or pacote['erros_bloco']:
        raise _LeituraIncompleta(pacote)
    return pacote


def _lojas_para_memoria(_engine):
    return _ler_lojas_ml(_engine)


def _pacote_para_memoria(_engine, hoje, lojas):
    return _ler_pacote_completo(_engine, hoje, lojas)


_MEMORIA = {}


def _memoria(nome):
    """st.cache_data criado uma vez por função (chave: hoje + lojas; o engine
    não entra na chave)."""
    f = _MEMORIA.get(nome)
    if f is None:
        import streamlit
        alvo = {'lojas': _lojas_para_memoria, 'pacote': _pacote_para_memoria}[nome]
        f = streamlit.cache_data(ttl=TTL_LEITURA_S, show_spinner=False, max_entries=64)(alvo)
        _MEMORIA[nome] = f
    return f


def limpar_memoria():
    for f in _MEMORIA.values():
        f.clear()


def obter_lojas_ml(engine):
    return _memoria('lojas')(engine) if _no_streamlit() else _ler_lojas_ml(engine)


def obter_pacote(engine, hoje, lojas):
    """(pacote, veio_da_memoria)."""
    if not _no_streamlit():
        return _ler_pacote(engine, hoje, tuple(lojas)), False
    antes = datetime.now(BRT)
    try:
        pacote = _memoria('pacote')(engine, hoje, tuple(lojas))
    except _LeituraIncompleta as e:
        return e.pacote, False
    return pacote, pacote['lido_em'] < antes


def _como_fragmento(f):
    """Cada loja num st.fragment: marcar uma caixa recarrega só aquela loja."""
    if _no_streamlit():
        import streamlit
        return streamlit.fragment(f)
    return f


def _render_mercado_livre(st, engine):
    t_ini = time.perf_counter()
    tempos = []
    agora = datetime.now(BRT)
    hoje = agora.date()

    restricao = restricao_de_lojas(engine)
    if restricao is not None and not restricao:
        st.caption("Nenhuma loja atribuída ao seu perfil.")
        return

    t = time.perf_counter()
    try:
        lojas_ml = obter_lojas_ml(engine)
    except Exception:  # noqa: BLE001
        st.error("Não consegui ler a lista de lojas agora. Tente de novo em instantes.")
        return
    tempos.append(('Lista de lojas', time.perf_counter() - t))
    lojas = lojas_visiveis(lojas_ml, restricao)
    if not lojas:
        st.caption("Nenhuma loja do Mercado Livre atribuída ao seu perfil.")
        return

    t = time.perf_counter()
    pacote, da_memoria = obter_pacote(engine, hoje, lojas)
    tempos.append(('Leitura do dia ' + (f"(da memória, lida às {pacote['lido_em']:%H:%M})"
                                        if da_memoria else "(do banco)"),
                   time.perf_counter() - t))
    dados, erros, erros_bloco = pacote['dados'], pacote['erros'], dict(pacote['erros_bloco'])
    sinais, nomes, parciais = pacote['sinais'], pacote['nomes'], pacote['parciais']

    t = time.perf_counter()
    cientes_linhas, erro_ciente = None, False
    try:
        conn = engine.raw_connection()
        try:
            cientes_linhas = ler_cientes(conn, lojas)
        finally:
            conn.close()
    except Exception:  # noqa: BLE001
        erro_ciente = True
        erros = {**erros, 'ciente': 'erro'}
    tempos.append(('Cientes (relidos a cada carga)', time.perf_counter() - t))
    ativos, silenciados = _bloco('ciente', erros_bloco, (sinais, []), aplicar_cientes,
                                 sinais, cientes_linhas or [], hoje)

    j = janelas(hoje)
    c_txt, c_bt = st.columns([5, 1])
    c_txt.caption(
        f"Dado de {j['ontem']:%d/%m} (ontem, dia fechado). Os sinais de venda comparam a "
        f"semana {j['sem_ini']:%d/%m}–{j['fim']:%d/%m} com a média semanal de "
        f"{j['base_ini']:%d/%m}–{j['base_fim']:%d/%m}: D-1 e D-2 ficam fora porque pedido "
        "pago depois (boleto/Pix) ainda pode entrar com a data deles. "
        f"Leitura de {pacote['lido_em']:%H:%M}.")
    if c_bt.button("🔄 Atualizar dados", key="sin_atualizar"):
        limpar_memoria()
        st.rerun()
    if erros:
        st.error("Não consegui ler agora: " + ', '.join(ROTULO_FONTE.get(k, k) for k in erros)
                 + ". Os sinais dessas fontes ficaram de fora; o resto está abaixo.")
    for bloco in erros_bloco:
        st.error(f"Os sinais de {ROTULO_BLOCO.get(bloco, bloco)} não puderam ser montados "
                 "agora; os outros blocos estão abaixo.")
    if dados.get('visitas') is None and 'visitas' not in erros:
        st.info("👀 Visitas: coleta de visitas em construção.")
    com_ciente = cientes_linhas is not None and not erro_ciente and 'ciente' not in erros_bloco
    pode = com_ciente and pode_dar_ciente()
    if cientes_linhas is None and not erro_ciente and pode_dar_ciente():
        st.caption("Ciente: aguardando a tabela sinal_ciente (SQL da v1.1 ainda não aplicado).")
    cientes = [dict(zip(CIENTE_COLS, r)) for r in cientes_linhas or []]

    frescor = avaliar_frescor(dados.get('frescor') or [], lojas, agora)
    resumo = resumo_lojas(dados.get('resumo') or [], dados.get('metas') or [], lojas, hoje)
    datas_foto = {r[0]: r[2] for r in dados.get('foto') or []}
    atrasadas = set(ep.lojas_atrasadas(datas_foto, hoje, agora))
    base_ctx = {'frescor': frescor, 'atrasadas': atrasadas, 'datas_foto': datas_foto,
                'parciais': parciais, 'erros': erros, 'resumo': resumo, 'nomes': nomes,
                'j': j, 'hoje': hoje, 'engine': engine, 'lojas': lojas, 'pode': pode,
                'com_ciente': com_ciente, 'silenciados': silenciados, 'ativos': ativos,
                'lido_em': pacote['lido_em']}

    try:
        _render_excel(st, ativos, silenciados, nomes, base_ctx)
    except Exception:  # noqa: BLE001
        st.error("Excel indisponível agora; os sinais abaixo seguem valendo.")

    t = time.perf_counter()
    fragmento = _como_fragmento(_render_loja_seguro)
    for loja in lojas:
        st.markdown("---")
        st.subheader(loja)
        fragmento(st, loja, ativos, {**base_ctx,
                                     'cientes': [c for c in cientes if c['loja'] == loja]})
    tempos.append(('Tela (4 lojas)' if len(lojas) == 4 else f'Tela ({len(lojas)} lojas)',
                   time.perf_counter() - t))
    tempos.append(('Total desta carga', time.perf_counter() - t_ini))
    _render_tempos(st, tempos, pacote, da_memoria)


def _render_tempos(st, tempos, pacote, da_memoria):
    from permissoes import _get_role
    if _get_role() != 'ADMIN':
        return
    with st.expander("⏱ tempos desta carga (só ADMIN)"):
        linhas = [{'Etapa': e, 'Segundos': f"{x:.2f}"} for e, x in tempos]
        if not da_memoria:
            linhas.append({'Etapa': '  (a leitura do banco em si)',
                           'Segundos': f"{pacote['segundos']:.2f}"})
        st.dataframe(pd.DataFrame(linhas), hide_index=True, use_container_width=True)
        st.caption("Marcar uma caixa recarrega só a loja (fragmento): não passa por aqui.")


def _excel_em_cache(ativos, nomes, ctx):
    """Bytes do Excel guardados na sessão até a leitura ou os cientes mudarem
    (gerar a cada clique deixaria a tela lenta de novo)."""
    import streamlit
    versao = streamlit.session_state.get('sin_versao', 0)
    chave = (ctx['lido_em'].isoformat(), versao, tuple(ctx['lojas']), len(ativos))
    guardado = streamlit.session_state.get('sin_xlsx')
    if not guardado or guardado[0] != chave:
        ordem = sorted(ativos, key=lambda x: (x['loja'], -x['em_jogo']))
        guardado = (chave, excel_sinais(ordem, nomes, ctx['hoje']))
        streamlit.session_state['sin_xlsx'] = guardado
    return guardado[1]


def _render_excel(st, ativos, silenciados, nomes, ctx):
    import streamlit
    hoje = ctx['hoje']
    versao = streamlit.session_state.get('sin_versao', 0)
    with st.expander("📥 Excel: baixar os sinais e subir cientes"):
        st.download_button(
            "Baixar Excel dos sinais", data=_excel_em_cache(ativos, nomes, ctx),
            file_name=f"sinais_ML_{hoje:%Y-%m-%d}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="sin_xlsx_baixar")
        if not ctx['pode']:
            st.caption("Seu perfil pode baixar, mas não subir cientes.")
            return
        arq = st.file_uploader("Subir o Excel preenchido (Status, Motivo, Nota, Silenciar até)",
                               type=['xlsx'], key=f"sin_xlsx_subir_{versao}")
        if arq is None:
            return
        try:
            linhas = ler_excel(arq.getvalue())
        except Exception as e:  # noqa: BLE001
            st.error(f"Não consegui ler o arquivo: {e}")
            return
        itens, previa = previa_excel(linhas, ativos + silenciados, ctx['lojas'], hoje)
        if previa.empty:
            st.info('Nenhuma linha com Status "Ciente" no arquivo.')
            return
        st.markdown("**Prévia** (nada foi gravado ainda)")
        st.dataframe(previa, hide_index=True, use_container_width=True)
        recusadas = len(previa) - len(itens)
        if not itens:
            st.warning("Nenhuma linha ok para gravar.")
            return
        rotulo = f"Gravar {len(itens)} ciente(s)"
        if recusadas:
            rotulo += f" — as {recusadas} recusada(s) ficam de fora"
        if st.button(rotulo, key=f"sin_xlsx_gravar_{versao}"):
            usuario = (streamlit.session_state.get('usuario') or {}).get('username', '')
            conn = ctx['engine'].raw_connection()
            try:
                n = gravar_cientes(conn, itens, usuario, hoje, ctx['lojas'])
            except Exception:  # noqa: BLE001
                st.error("Não consegui gravar agora; nada foi gravado. Tente de novo.")
                return
            finally:
                conn.close()
            streamlit.session_state['sin_versao'] = versao + 1
            st.success(f"Ciente registrado em {n} sinal(is).")
            st.rerun()


def _render_loja_seguro(st, loja, ativos, ctx):
    try:
        _render_loja(st, loja, ativos, ctx)
    except Exception:  # noqa: BLE001 — uma loja quebrada não esconde as outras
        st.error(f"{loja}: não consegui mostrar esta loja agora; as outras seguem.")


def _editor(st, sinais, nomes, chave, pode):
    """Tabela de sinais; com permissão, uma coluna "Ciente" para marcar.
    Devolve os sinais marcados."""
    df = tabela_sinais(sinais, nomes)
    if not pode:
        st.dataframe(df, use_container_width=True, hide_index=True)
        return []
    df.insert(0, 'Ciente', False)
    ed = st.data_editor(df, key=chave, hide_index=True, use_container_width=True,
                        disabled=[c for c in df.columns if c != 'Ciente'])
    return [sinais[i] for i in range(len(sinais)) if bool(ed['Ciente'].iloc[i])]


def _render_form_ciente(st, loja, marcados, ctx):
    import streamlit
    hoje = ctx['hoje']
    with st.form(key=f"sin_form_{loja}"):
        st.markdown(f"**Ciente em {len(marcados)} sinal(is) da {loja}**")
        motivo = st.selectbox("Motivo", list(MOTIVOS), format_func=MOTIVOS.get,
                              key=f"sin_mot_{loja}")
        nota = st.text_input("Nota (obrigatória em \"Outro\")", key=f"sin_nota_{loja}")
        ate = st.date_input("Silenciar até", value=hoje + timedelta(days=DIAS_SILENCIO_PADRAO),
                            min_value=hoje, max_value=hoje + timedelta(days=DIAS_SILENCIO_MAX),
                            format="DD/MM/YYYY", key=f"sin_ate_{loja}")
        if not st.form_submit_button("Registrar ciente"):
            return
    erros = validar_ciente(motivo, nota, ate, hoje)
    if erros:
        st.error(' '.join(erros))
        return
    usuario = (streamlit.session_state.get('usuario') or {}).get('username', '')
    conn = ctx['engine'].raw_connection()
    try:
        n = gravar_cientes(conn, [(x, motivo, nota, ate) for x in marcados], usuario, hoje,
                           ctx['lojas'])
    except Exception:  # noqa: BLE001
        st.error("Não consegui registrar o ciente agora; nada foi gravado. Tente de novo.")
        return
    finally:
        conn.close()
    streamlit.session_state['sin_versao'] = streamlit.session_state.get('sin_versao', 0) + 1
    st.success(f"Ciente registrado em {n} sinal(is).")
    st.rerun()


def _render_silenciados(st, loja, ctx):
    import streamlit
    sil = [x for x in ctx['silenciados'] if x['loja'] == loja]
    ativos = [x for x in ctx['ativos'] if x['loja'] == loja]
    df = tabela_silenciados(sil, ctx['cientes'], ctx['hoje'], ctx['nomes'], ativos)
    if df.empty:
        return
    with st.expander(f"🔕 Silenciados ({len(df)})"):
        st.dataframe(df, use_container_width=True, hide_index=True)
        if not ctx['pode']:
            return
        abertos = cientes_silenciando(ctx['cientes'], ativos, ctx['hoje'])
        escolha = st.selectbox(
            "Reativar (o sinal volta a aparecer)", [c['id'] for c in abertos],
            format_func=lambda i: next(f"{REGRAS.get(c['regra'], c['regra'])} · {c['objeto']}"
                                       for c in abertos if c['id'] == i),
            key=f"sin_reat_{loja}")
        if st.button("Reativar", key=f"sin_reat_btn_{loja}"):
            usuario = (streamlit.session_state.get('usuario') or {}).get('username', '')
            conn = ctx['engine'].raw_connection()
            try:
                reativar_ciente(conn, escolha, usuario, ctx['lojas'])
            except Exception:  # noqa: BLE001
                st.error("Não consegui reativar agora; nada mudou.")
                return
            finally:
                conn.close()
            st.rerun()


def _render_loja(st, loja, ativos, ctx):
    frescor, atrasadas, datas_foto = ctx['frescor'], ctx['atrasadas'], ctx['datas_foto']
    parciais, erros, resumo, nomes, j = (ctx['parciais'], ctx['erros'], ctx['resumo'],
                                         ctx['nomes'], ctx['j'])
    avisos = frescor.get(loja, [])
    atr = [t for n, t in avisos if n == 'atrasado']
    cedo = [t for n, t in avisos if n == 'cedo']
    if atr:
        st.warning("⚠️ O coletor de hoje ainda não trouxe: " + '; '.join(atr))
    elif cedo:
        st.caption("⏳ O coletor do dia pode ainda não ter rodado: " + '; '.join(cedo))
    if loja in atrasadas:
        st.warning(f"⚠️ Foto do estoque atrasada: {pd.Timestamp(datas_foto[loja]):%d/%m}.")
    for l, n_ult, ult, n_ant, ant in parciais:
        if l == loja:
            st.warning(f"⚠️ Foto do estoque possivelmente PARCIAL: {n_ult} estoques em "
                       f"{pd.Timestamp(ult):%d/%m} contra {n_ant} em {pd.Timestamp(ant):%d/%m}.")

    if 'resumo' in erros:
        st.error("Resumo indisponível agora.")
    else:
        try:
            _render_resumo(st, loja, resumo[loja], j['ontem'])
        except Exception:  # noqa: BLE001
            st.error("Resumo indisponível agora.")

    import streamlit
    versao = streamlit.session_state.get('sin_versao', 0)
    try:
        top, resto = sinais_da_loja(ativos, loja)
        if not top:
            st.success("Nenhum sinal hoje.")
            marcados = []
        else:
            marcados = _editor(st, top, nomes, f"sin_ed_{loja}_{versao}", ctx['pode'])
            if resto:
                with st.expander(f"Mais {len(resto)} sinais (menor R$ em jogo)"):
                    marcados += _editor(st, resto, nomes, f"sin_ed_resto_{loja}_{versao}",
                                        ctx['pode'])
    except Exception:  # noqa: BLE001
        st.error("Tabela de sinais indisponível agora para esta loja.")
        return
    if ctx['com_ciente']:
        try:
            if marcados:
                _render_form_ciente(st, loja, marcados, ctx)
            _render_silenciados(st, loja, ctx)
        except Exception:  # noqa: BLE001
            st.error("Ciente indisponível agora para esta loja; os sinais acima seguem valendo.")


def render(engine):
    import streamlit as st

    st.title("📡 Sinais do Dia")
    st.caption("Painel diário: o que mudou ontem em cada loja, ordenado pelo R$ em jogo. "
               "Marque \"Ciente\" para responder a um sinal e silenciá-lo até uma data; "
               "os limiares estão em calibragem.")
    aba_ml, aba_shopee = st.tabs(["Mercado Livre", "Shopee"])
    with aba_ml:
        try:
            _render_mercado_livre(st, engine)
        except Exception:  # noqa: BLE001
            st.error("Os sinais do Mercado Livre não carregaram agora. Tente de novo em "
                     "instantes; se persistir, avise o time do sistema.")
    with aba_shopee:
        st.info("Shopee: em seguida (v2).")
