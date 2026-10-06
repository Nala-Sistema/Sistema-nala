"""
Testa o painel Sinais do Dia (sinais_dia.py, frente [SINAIS DO DIA], v1).

Duas partes:
  - Sem banco: as regras de cada sinal com os casos reais da calibragem do
    [MESTRE ANÁLISES] (05/10/2026) — K-L-0421-A (Full por estoque_id, nunca
    somado por SKU), Garrafa LTT-CT3821 (sem estoque não é espiral, é
    ruptura), família de alicates (experiência por peça base) —, janelas com
    D-1/D-2 fora, corte de 10 por loja, permissão que NÃO abre tudo para
    usuário sem loja, e % solto nas SQL.
  - Ciente (v1.1): chave regra + loja + objeto, quando o silenciado volta
    (data ou piora, por regra), validação, permissão de quem grava; com banco,
    o ciclo gravar -> substituir -> reativar numa TEMP criada com o DDL LIDO de
    sql/sinais_ciente.sql (a tabela ainda não existe em produção; depois de
    aplicada, um teste confere que as CHECK reais são as do arquivo).
  - v1.2: Período em cada sinal, "?" no lugar de "parcial", leitura com
    rollback só no erro e em memória só quando completa, quadro de tempos só
    para ADMIN, Excel de ida e volta (gera, lê, prévia com cada recusa,
    grava pela mesma gravar_cientes; DIRETOR baixa e não sobe).
  - Aba Shopee: anúncio pela ponte do espelho (números do snapshot), Litstore
    por SKU, ROAS 0 = automático, custo por PEDIDO × margem por PEDIDO,
    espiral só com estoque no banco, views de página D × D-30 em data exata,
    nota só das avaliações novas, carga pendente só para loja de API; com
    banco, o cenário Shopee inteiro pelas SQL.
  - Full × galpão (regra do Thiago, 06/10, igual no ML e na Shopee) com os
    casos reais de 05/10: K10-L-7248-A (Shopee Nala, Full 0, galpão 18, ads
    ligado), L-0321 (Shopee LPT, Full 29, galpão 2.328), K-L-0421-A (Shopee
    LPT, Full 0, galpão 364) e um caso ML com Full 0 e galpão > 0; upload
    atrasado da Litstore (R2); nota inválida com rating 0 (R3).
  - v1.3 "Aparece desde": a conta da sequência (dia sem execução do job não
    quebra; "🆕 hoje"), a montagem única da tela e do job, o job do repo
    público (só contagens na saída, erro só com o tipo, nada antes das 10h,
    backfill), a gravação com upsert numa TEMP criada com o DDL LIDO de
    sql/sinais_historico.sql, e o teto de data das leituras de "última foto"
    (o backfill de um dia passado lê a foto e a config daquele dia).
  - Com banco (NALA_TEST_DB_URL, usuário de permissão mínima): EXECUTA as SQL
    de verdade em tabelas TEMPORÁRIAS com os nomes das reais, com as CHECK
    COPIADAS da tabela real (lição de 05/10: a TEMP sem CHECK aceitava o que
    produção recusa), conferidas em pg_temp antes de qualquer INSERT, e SEMPRE
    ROLLBACK.

Rodar:  python -m pytest tests/test_sinais_dia.py
"""

import io
import os
import re
import sys
import unittest
from datetime import date, datetime, timedelta
from unittest import mock

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

import sinais_dia as sd  # noqa: E402

DB_URL = os.environ.get('NALA_TEST_DB_URL')
HOJE = date(2026, 10, 5)
LPT, NALA, RJ, SP = 'ML-LPT', 'ML-Nala', 'ML-YanniRJ', 'ML-YanniSP'
S_LPT, S_NALA, S_YANNI = 'Shopee-LPT', 'Shopee Lithouse(Nala)', 'Shopee Litstore(Yanni)'


def _colunas(n):
    return [mock.MagicMock() for _ in range(n if isinstance(n, int) else len(n))]


def _gravar(conn, sinais, motivo, nota, ate, usuario, hoje, lojas):
    """Atalho dos testes antigos: o mesmo motivo/nota/data para todos."""
    return sd.gravar_cientes(conn, [(x, motivo, nota, ate) for x in sinais], usuario, hoje,
                             lojas)


def _vend(loja, anuncio, sku, rec_sem=0, rec_base=0, qtd_ritmo=0, qtd_30=0,
          rec_30=0, margem_30=None):
    return (loja, anuncio, sku, rec_sem, rec_base, qtd_ritmo, qtd_30, rec_30, margem_30)


def _ads(gasto_sem=0, gasto_base=0, gasto_ontem=0, gasto_recente=0, dias_com_gasto=0,
         unid_recente=0, gasto_30=0, unid_ads_30=0, cliques_30=0):
    return dict(gasto_sem=gasto_sem, gasto_base=gasto_base, gasto_ontem=gasto_ontem,
                gasto_recente=gasto_recente, dias_com_gasto=dias_com_gasto,
                unid_recente=unid_recente, gasto_30=gasto_30, unid_ads_30=unid_ads_30,
                cliques_30=cliques_30)


class Janelas(unittest.TestCase):
    def test_janelas_terminam_em_d1_semanas_inteiras_sem_sobrepor(self):
        j = sd.janelas(HOJE)
        self.assertEqual(j['fim'], date(2026, 10, 4))          # D-1 (medido: D-1 não muda)
        self.assertEqual(j['sem_ini'], date(2026, 9, 28))      # D-7
        self.assertEqual(j['base_fim'], date(2026, 9, 27))     # D-8
        self.assertEqual(j['base_ini'], date(2026, 8, 31))     # D-35
        self.assertEqual((j['base_fim'] - j['base_ini']).days + 1, 28)
        self.assertEqual(j['ritmo_ini'], date(2026, 9, 28))
        self.assertEqual(j['piso_ini'], date(2026, 9, 5))
        self.assertEqual(j['mesmo_dia'], [date(2026, 9, 27), date(2026, 9, 20),
                                          date(2026, 9, 13), date(2026, 9, 6)])
        self.assertTrue(all(d.weekday() == j['ontem'].weekday() for d in j['mesmo_dia']))

    def test_dia_1_usa_o_mes_de_ontem(self):
        j = sd.janelas(date(2026, 11, 1))
        self.assertEqual(j['mes_ini'], date(2026, 10, 1))
        self.assertEqual(sd.params(date(2026, 11, 1), [LPT])['ano_mes'], '2026-10')


class Full(unittest.TestCase):
    # K-L-0421-A na ML-LPT, números reais de 04/10/2026
    PONTE = [(LPT, 'MLB5183384677', 'MLBU5071722523', 'K-L-0421-A', True),
             (LPT, 'MLB4793561323', 'MLBU4131286394', 'K-L-0421-A', True)]
    FOTO = [(LPT, 'MLBU5071722523', date(2026, 10, 4), 6, 2, 0),
            (LPT, 'MLBU4131286394', date(2026, 10, 4), 89, 0, 0)]
    VENDAS = [_vend(LPT, 'MLB5183384677', 'K-L-0421-A', qtd_ritmo=21, qtd_30=65, rec_30=2000),
              _vend(LPT, 'MLB4793561323', 'K-L-0421-A', qtd_ritmo=2, qtd_30=12, rec_30=380)]

    def _full(self, foto=None, ponte=None, vendas=None, ads=None):
        _a, por_sku = sd.agregar_vendas(vendas if vendas is not None else self.VENDAS)
        return sd.sinais_full(foto or self.FOTO, ponte or self.PONTE, por_sku, ads or {},
                              sd.janelas(HOJE))

    def test_k_l_0421_a_por_estoque_id_nunca_somado(self):
        s = self._full()
        self.assertEqual(len(s), 1)
        self.assertEqual(s[0]['anuncio'], 'MLB5183384677')
        self.assertTrue(s[0]['urgente'])
        # (6 + 2) ÷ max(21/7, 65/30) = 2,7 dias; a soma do SKU (95) nunca aparece
        self.assertIn('3 dias de Full (6 disp. + 2 em transf.', s[0]['numero'])
        self.assertNotIn('95', s[0]['numero'])
        self.assertNotIn('97', s[0]['numero'])

    def test_piso_da_media_de_30_quando_a_semana_e_fraca(self):
        vendas = [_vend(LPT, 'MLB5183384677', 'K-L-0421-A', qtd_ritmo=7, qtd_30=90)]
        s = self._full(vendas=vendas)
        self.assertIn('÷ 3,0/dia', s[0]['numero'])          # 90/30 = 3 > 7/7 = 1

    def test_dois_anuncios_no_mesmo_estoque_somam_a_venda(self):
        ponte = [(LPT, 'MLB1', 'E1', 'L-0320', True), (LPT, 'MLB2', 'E1', 'L-0320', True)]
        foto = [(LPT, 'E1', date(2026, 10, 4), 20, 0, 0)]
        vendas = [_vend(LPT, 'MLB1', 'L-0320', qtd_ritmo=14, qtd_30=30),
                  _vend(LPT, 'MLB2', 'L-0320', qtd_ritmo=14, qtd_30=30)]
        s = self._full(foto, ponte, vendas)
        self.assertEqual(len(s), 1)
        self.assertIn('5 dias de Full', s[0]['numero'])     # 20 ÷ (28/7)

    def test_venda_de_outro_sku_no_mesmo_anuncio_nao_entra(self):
        ponte = [(LPT, 'MLB1', 'E1', 'L-0320', True)]
        foto = [(LPT, 'E1', date(2026, 10, 4), 20, 0, 0)]
        vendas = [_vend(LPT, 'MLB1', 'L-0320', qtd_ritmo=7, qtd_30=30),
                  _vend(LPT, 'MLB1', 'K2-L-0320', qtd_ritmo=700, qtd_30=3000)]
        s = self._full(foto, ponte, vendas)
        self.assertEqual(s, [])                               # 20 ÷ 1 = 20 dias

    def test_full_zerado_com_ads_ligado_e_envio_que_entrou(self):
        foto = [(LPT, 'MLBU5071722523', date(2026, 10, 4), 0, 30, 0),
                (LPT, 'MLBU4131286394', date(2026, 10, 4), 89, 0, 40)]
        ads = {(LPT, 'MLB5183384677'): _ads(gasto_ontem=12.5)}
        s = self._full(foto=foto, ads=ads)
        textos = [x['numero'] for x in s]
        self.assertTrue(any(t.startswith('Full zerado com ads ligado') and '30 em transferência' in t
                            and 'ads do anúncio, todas as variações' in t
                            for t in textos))
        self.assertTrue(any(t.startswith('Envio entrou no Full em 04/10: +40') for t in textos))

    def test_anuncio_fora_do_full_nao_tem_cobertura_de_full(self):
        ponte = [(LPT, 'MLB1', 'E1', 'L-0320', False)]
        foto = [(LPT, 'E1', date(2026, 10, 4), 0, 0, 0)]
        self.assertEqual(self._full(foto, ponte, [_vend(LPT, 'MLB1', 'L-0320', 7, 0, 7, 30)]), [])


class Ads(unittest.TestCase):
    def _ads(self, ads, vendas, dias_est, config=()):
        por_an, _ = sd.agregar_vendas(vendas)
        return sd.sinais_ads(ads, list(config), por_an, dias_est, sd.janelas(HOJE))

    def test_espiral_so_com_estoque_a_semana_toda(self):
        vendas = [_vend(LPT, 'MLB1', 'L-0320', rec_sem=300, rec_base=4000)]
        ads = {(LPT, 'MLB1'): _ads(gasto_sem=20, gasto_base=400)}
        s = self._ads(ads, vendas, {(LPT, 'MLB1'): 7})
        self.assertEqual([x['tipo'] for x in s], ['ADS'])
        self.assertTrue(s[0]['numero'].startswith('Espiral'))
        self.assertIn('TACOS', s[0]['numero'])

    def test_garrafa_ltt_ct3821_sem_full_e_ruptura_nao_espiral(self):
        vendas = [_vend(NALA, 'MLB7175749876', 'LTT-CT3821', rec_sem=150, rec_base=3200)]
        ads = {(NALA, 'MLB7175749876'): _ads(gasto_sem=10, gasto_base=300)}
        s = self._ads(ads, vendas, {(NALA, 'MLB7175749876'): 3})
        self.assertEqual(len(s), 1)
        self.assertEqual(s[0]['tipo'], 'FULL')
        self.assertTrue(s[0]['numero'].startswith('Ruptura: estoque (Full ou galpão) em só 3 de 7'))
        self.assertFalse(any(x['numero'].startswith('Espiral') for x in s))

    def test_sem_leitura_de_estoque_a_espiral_nao_e_julgada(self):
        vendas = [_vend(LPT, 'MLB1', 'L-0320', rec_sem=300, rec_base=4000)]
        ads = {(LPT, 'MLB1'): _ads(gasto_sem=20, gasto_base=400)}
        por_an, _ = sd.agregar_vendas(vendas)
        s = sd.sinais_ads(ads, [], por_an, {}, sd.janelas(HOJE), com_estoque=False)
        self.assertEqual(s, [])

    def test_gasto_com_zero_venda_por_3_dias(self):
        ads = {(LPT, 'MLB1'): _ads(gasto_recente=31.2, dias_com_gasto=3, unid_recente=0),
               (LPT, 'MLB2'): _ads(gasto_recente=31.2, dias_com_gasto=2, unid_recente=0),
               (LPT, 'MLB3'): _ads(gasto_recente=31.2, dias_com_gasto=3, unid_recente=1)}
        s = self._ads(ads, [], {})
        self.assertEqual([x['anuncio'] for x in s], ['MLB1'])
        self.assertAlmostEqual(s[0]['em_jogo'], 31.2)

    def test_custo_por_venda_acima_da_margem_por_unidade(self):
        vendas = [_vend(LPT, 'MLB1', 'L-0320', qtd_30=10, margem_30=50)]   # R$ 5/un
        ads = {(LPT, 'MLB1'): _ads(gasto_30=80, unid_ads_30=10, cliques_30=200)}  # R$ 8/un
        s = self._ads(ads, vendas, {})
        self.assertEqual(len(s), 1)
        self.assertIn('R$ 8,00/un', s[0]['numero'])
        self.assertIn('margem R$ 5,00/un', s[0]['numero'])
        self.assertIn('CPC R$ 0,40', s[0]['numero'])
        self.assertAlmostEqual(s[0]['em_jogo'], 30.0)

    def test_custo_por_venda_nao_divide_unidades_de_skus_diferentes(self):
        vendas = [_vend(LPT, 'MLB1', 'L-0250-A', qtd_30=5, margem_30=10),
                  _vend(LPT, 'MLB1', 'L-0250-B', qtd_30=5, margem_30=10)]
        ads = {(LPT, 'MLB1'): _ads(gasto_30=80, unid_ads_30=10, cliques_30=200)}
        self.assertEqual(self._ads(ads, vendas, {}), [])

    def test_mudanca_de_roas_e_orcamento(self):
        config = [(LPT, 'MLB5379140666 Pote Divisoria', 'MLB5379140666 Pote Divisoria',
                   date(2026, 10, 4), HOJE, 14, 10, 12, 20)]
        vendas = [_vend(LPT, 'MLB5379140666', 'L-0400', rec_30=1500)]
        s = self._ads({}, vendas, {}, config)
        self.assertEqual(s[0]['anuncio'], 'MLB5379140666')
        self.assertIn('ROAS objetivo 14,0 → 10,0', s[0]['numero'])
        self.assertIn('orçamento R$ 12,00 → R$ 20,00/dia', s[0]['numero'])
        self.assertAlmostEqual(s[0]['em_jogo'], 1500)

    def test_anuncio_da_campanha(self):
        self.assertEqual(sd.anuncio_da_campanha('MLB4639959197 Torx + Allen'), 'MLB4639959197')
        self.assertIsNone(sd.anuncio_da_campanha('Campanha geral', None))


class Experiencia(unittest.TestCase):
    # Alicates da ML-LPT: 4 kits caíram de 75–100 para 65; o principal em 100.
    COMP = [('K2-L-0380', 'L-0380', 2), ('K3-L-0380', 'L-0380', 3),
            ('K4-L-0380', 'L-0380', 4), ('K5-L-0380', 'L-0380', 5)]
    PONTE = [(LPT, 'MLB10', 'E10', 'L-0380', True), (LPT, 'MLB2', 'E2', 'K2-L-0380', True),
             (LPT, 'MLB3', 'E3', 'K3-L-0380', True), (LPT, 'MLB4', 'E4', 'K4-L-0380', True),
             (LPT, 'MLB5', 'E5', 'K5-L-0380', True), (LPT, 'MLB9', 'E9', 'L-0999', True)]
    D = date(2026, 10, 5)

    def _exp(self, saude):
        import estoque_peca as ep
        por_an, _ = sd.agregar_vendas([_vend(LPT, 'MLB10', 'L-0380', rec_30=900),
                                       _vend(LPT, 'MLB2', 'K2-L-0380', rec_30=400)])
        return sd.sinais_experiencia(saude, self.PONTE, ep.agrupar_composicao(self.COMP), por_an)

    def test_familia_de_alicates_vira_um_sinal_com_o_principal(self):
        saude = [(LPT, 'MLB10', self.D, 100, 'Boa', 100, 'Boa'),
                 (LPT, 'MLB2', self.D, 65, 'Média', 100, 'Boa'),
                 (LPT, 'MLB3', self.D, 65, 'Média', 75, 'Boa'),
                 (LPT, 'MLB4', self.D, 65, 'Média', 100, 'Boa'),
                 (LPT, 'MLB5', self.D, 65, 'Média', 75, 'Boa')]
        s = self._exp(saude)
        self.assertEqual(len(s), 1)
        self.assertTrue(s[0]['numero'].startswith('4 anúncios da mesma peça pioraram'))
        self.assertEqual(s[0]['skus'], ('L-0380',))
        self.assertIn('principal MLB10 em 100', s[0]['familia'])
        self.assertIn('5 anúncios', s[0]['familia'])
        self.assertAlmostEqual(s[0]['em_jogo'], 400)

    def test_um_so_piorou_sai_sozinho_com_a_familia_ao_lado(self):
        saude = [(LPT, 'MLB10', self.D, 100, 'Boa', 100, 'Boa'),
                 (LPT, 'MLB2', self.D, 65, 'Média', 100, 'Boa'),
                 (LPT, 'MLB3', self.D, 75, 'Boa', 75, 'Boa')]
        s = self._exp(saude)
        self.assertEqual(len(s), 1)
        self.assertEqual(s[0]['anuncio'], 'MLB2')
        self.assertIn('100 (Boa) → 65 (Média)', s[0]['numero'])
        self.assertIn('principal MLB10 em 100', s[0]['familia'])

    def test_sem_nota_nao_e_piora(self):
        saude = [(LPT, 'MLB2', self.D, None, None, 100, 'Boa'),
                 (LPT, 'MLB3', self.D, 65, 'Média', None, None),
                 (LPT, 'MLB9', self.D, 100, 'Boa', 30, 'Ruim')]
        self.assertEqual(self._exp(saude), [])


class VendasEOrdem(unittest.TestCase):
    def test_queda_alta_e_relevancia(self):
        vendas = [_vend(LPT, 'MLB1', 'A', rec_sem=400, rec_base=4000),    # 1000/sem → -60%
                  _vend(LPT, 'MLB2', 'B', rec_sem=2500, rec_base=4000),   # +150%
                  _vend(LPT, 'MLB3', 'C', rec_sem=0, rec_base=900),       # irrelevante
                  _vend(LPT, 'MLB4', 'D', rec_sem=900, rec_base=4000)]    # -10%: nada
        por_an, _ = sd.agregar_vendas(vendas)
        s = sd.sinais_vendas(por_an, sd.janelas(HOJE))
        self.assertEqual(sorted(x['anuncio'] for x in s), ['MLB1', 'MLB2'])
        q = next(x for x in s if x['anuncio'] == 'MLB1')
        self.assertIn('semana 28/09–04/10', q['numero'])
        self.assertAlmostEqual(q['em_jogo'], 600)

    def test_ruptura_tira_a_queda_generica_do_mesmo_anuncio(self):
        dados = {'vendas': [_vend(NALA, 'MLB7', 'LTT-CT3821', rec_sem=150, rec_base=3200)],
                 'ads': [(NALA, 'MLB7', 10, 300, 0, 0, 0, 0, 0, 0, 0)],
                 'estoque_semana': [], 'ponte': [], 'config': [],
                 # a loja tem estoque no banco (outro anúncio); o MLB7 não teve na semana
                 'foto': [(NALA, 'OUTRO', date(2026, 10, 4), 5, 0, 0)]}
        s, erros = sd.montar_sinais(dados, HOJE)
        self.assertEqual(erros, {})
        self.assertEqual([(x['tipo'], x['numero'].split(':')[0]) for x in s], [('FULL', 'Ruptura')])

    def test_corte_de_10_por_loja_ordenado_por_r(self):
        sinais = [sd._sinal(LPT, 'VENDAS', f'MLB{i}', ['A'], 'x', 's', i) for i in range(15)]
        sinais.append(sd._sinal(NALA, 'VENDAS', 'MLB99', ['A'], 'x', 's', 999))
        top, resto = sd.sinais_da_loja(sinais, LPT)
        self.assertEqual(len(top), 10)
        self.assertEqual(len(resto), 5)
        self.assertEqual([x['em_jogo'] for x in top], list(range(14, 4, -1)))
        self.assertNotIn('MLB99', [x['anuncio'] for x in top + resto])

    def test_fonte_com_erro_nao_derruba_as_outras(self):
        dados = {'vendas': [_vend(LPT, 'MLB1', 'A', rec_sem=0, rec_base=4000)]}
        s, erros = sd.montar_sinais(dados, HOJE)   # sem ponte/foto/ads/experiência
        self.assertEqual([x['tipo'] for x in s], ['VENDAS'])
        self.assertEqual(erros, {})

    def test_bloco_que_quebra_nao_derruba_os_outros(self):
        dados = {'vendas': [_vend(LPT, 'MLB1', 'A', rec_sem=0, rec_base=4000)],
                 'foto': [(LPT, 'E1', date(2026, 10, 4), 0, 0, 0)],
                 'ponte': [(LPT, 'MLB1', 'E1', 'A', True)],
                 'experiencia': [(LPT, 'MLB1', HOJE, 30, 'Ruim', 100, 'Boa')],
                 'ads': [], 'config': [], 'estoque_semana': []}
        with mock.patch.object(sd, 'sinais_full', side_effect=KeyError('dado inesperado')):
            s, erros = sd.montar_sinais(dados, HOJE)
        self.assertEqual(set(erros), {'full'})
        self.assertEqual(sorted({x['tipo'] for x in s}), ['EXPERIÊNCIA', 'VENDAS'])


class ResumoEFrescor(unittest.TestCase):
    def test_resumo_mesmo_dia_da_semana_e_meta(self):
        r = sd.resumo_lojas([(LPT, 3000, 16000, 12000)], [(LPT, 200000)], [LPT, NALA], HOJE)
        self.assertAlmostEqual(r[LPT]['media_mesmo_dia'], 4000)
        self.assertAlmostEqual(r[LPT]['var_ontem'], -0.25)
        self.assertAlmostEqual(r[LPT]['pct_meta'], 0.06)
        self.assertAlmostEqual(r[LPT]['esperado_linear'], 200000 * 4 / 31)
        self.assertIsNone(r[NALA]['meta'])                    # meta não cadastrada
        self.assertIsNone(r[NALA]['var_ontem'])

    def test_frescor_cedo_atrasado_e_loja_sem_ads(self):
        linhas = [('Vendas', LPT, datetime(2026, 10, 4, 6, 4)),
                  ('Estoque', LPT, datetime(2026, 10, 5, 4, 8))]
        cedo = sd.avaliar_frescor(linhas, [LPT, SP], datetime(2026, 10, 5, 7, 0))
        tarde = sd.avaliar_frescor(linhas, [LPT, SP], datetime(2026, 10, 5, 11, 0))
        self.assertEqual(cedo[LPT], [('cedo', 'Vendas: último dado 04/10 06:04')])
        self.assertEqual(tarde[LPT], [('atrasado', 'Vendas: último dado 04/10 06:04')])
        self.assertEqual(tarde[SP], [])                       # sem dado = sem alarme


class Permissao(unittest.TestCase):
    def test_restrito_sem_loja_nao_ve_nada_e_nao_consulta(self):
        import permissoes
        engine = mock.Mock()
        with mock.patch.object(permissoes, 've_todas_lojas', return_value=False), \
             mock.patch.object(permissoes, 'get_lojas_usuario', return_value=[]):
            st = mock.Mock()
            sd._render_mercado_livre(st, engine)
        engine.raw_connection.assert_not_called()
        st.caption.assert_called_with("Nenhuma loja atribuída ao seu perfil.")

    def test_tela_mostra_erro_do_bloco_e_as_outras_lojas_seguem(self):
        dados = {'vendas': [_vend(LPT, 'MLB1', 'A', rec_sem=0, rec_base=4000),
                            _vend(NALA, 'MLB2', 'B', rec_sem=0, rec_base=4000)],
                 'foto': [], 'ponte': [], 'resumo': [], 'metas': [], 'frescor': [],
                 'visitas': None}
        st = mock.MagicMock()
        st.columns.side_effect = _colunas
        engine = mock.MagicMock()
        tabela_real = sd.tabela_sinais

        def tabela_que_quebra_na_lpt(sinais, nomes):
            if sinais and sinais[0]['loja'] == LPT:
                raise ValueError('dado inesperado')
            return tabela_real(sinais, nomes)
        with mock.patch.object(sd, 'restricao_de_lojas', return_value=None), \
             mock.patch.object(sd, '_ler', return_value=[(LPT,), (NALA,)]), \
             mock.patch.object(sd, 'ler_tudo', return_value=(dados, {})), \
             mock.patch.object(sd, 'ler_cientes', return_value=None), \
             mock.patch.object(sd.ep, 'ler_contagens', return_value=[]), \
             mock.patch.object(sd.ep, 'ler_nomes', return_value={}), \
             mock.patch.object(sd, 'sinais_vendas', side_effect=[KeyError('x')]):
            sd._render_mercado_livre(st, engine)
        erros = [c.args[0] for c in st.error.call_args_list]
        self.assertTrue(any('VENDAS' in e for e in erros), erros)
        self.assertEqual(st.success.call_count, 2)      # as duas lojas seguem: "Nenhum sinal"

        st = mock.MagicMock()
        st.columns.side_effect = _colunas
        with mock.patch.object(sd, 'restricao_de_lojas', return_value=None), \
             mock.patch.object(sd, '_ler', return_value=[(LPT,), (NALA,)]), \
             mock.patch.object(sd, 'ler_tudo', return_value=(dados, {})), \
             mock.patch.object(sd, 'ler_cientes', return_value=None), \
             mock.patch.object(sd.ep, 'ler_contagens', return_value=[]), \
             mock.patch.object(sd.ep, 'ler_nomes', return_value={}), \
             mock.patch.object(sd, 'tabela_sinais', side_effect=tabela_que_quebra_na_lpt):
            sd._render_mercado_livre(st, engine)
        erros = [c.args[0] for c in st.error.call_args_list]
        self.assertEqual(erros, ["Tabela de sinais indisponível agora para esta loja."])
        self.assertEqual(st.dataframe.call_count, 1)    # a ML-Nala apareceu

    def test_restricao_e_lojas_visiveis(self):
        import permissoes
        with mock.patch.object(permissoes, 've_todas_lojas', return_value=True):
            self.assertIsNone(sd.restricao_de_lojas(None))
        with mock.patch.object(permissoes, 've_todas_lojas', return_value=False), \
             mock.patch.object(permissoes, 'get_lojas_usuario', return_value=[NALA, 'Shopee-LPT']):
            r = sd.restricao_de_lojas(None)
        self.assertEqual(sd.lojas_visiveis([LPT, NALA, RJ], r), [NALA])
        self.assertEqual(sd.lojas_visiveis([LPT, NALA], None), [LPT, NALA])

    def test_menu_e_perfis(self):
        import permissoes
        self.assertEqual(permissoes.MENU_MODULOS['📡 Sinais do Dia'], 'sinais')
        for perfil, perms in permissoes.PERMISSOES.items():
            self.assertIn('sinais', perms, perfil)
        self.assertEqual(permissoes.PERMISSOES['GESTOR']['sinais'], 'parcial')


class Sql(unittest.TestCase):
    def test_sql_nao_fixa_schema_nem_tem_porcento_solto(self):
        for sql in sd.TODAS_AS_SQL:
            self.assertNotIn('public.', sql)
            self.assertEqual(re.findall(r'%(?!\(\w+\)s)', sql), [], sql[:80])

    def test_todo_parametro_da_sql_existe(self):
        p = sd.params(HOJE, [LPT])
        escrita = (sd.SQL_FECHAR_ABERTO, sd.SQL_INSERIR_CIENTE, sd.SQL_REATIVAR,
                   sd.SQL_GRAVAR_HISTORICO)
        p_escrita = {'marketplace', 'loja', 'regra', 'objeto', 'motivo', 'nota',
                     'silenciar_ate', 'medida', 'texto', 'usuario', 'id', 'lojas',
                     'data', 'em_jogo'}
        for sql in sd.TODAS_AS_SQL:
            for nome in re.findall(r'%\((\w+)\)s', sql):
                self.assertIn(nome, p_escrita if sql in escrita else p, nome)


# ============================================================
# CIENTE (v1.1) — sem banco
# ============================================================

SQL_CIENTE = os.path.join(RAIZ, 'sql', 'sinais_ciente.sql')


def _ddl_ciente_em_temp():
    with open(SQL_CIENTE, encoding='utf-8') as f:
        texto = f.read()
    tabela = re.findall(r'^CREATE TABLE public\.sinal_ciente \(.*?^\);', texto, re.S | re.M)
    indice = re.findall(r'^CREATE UNIQUE INDEX .*?;$', texto, re.M)
    assert len(tabela) == 1 and len(indice) == 1, 'DDL de sql/sinais_ciente.sql mudou'
    return [x.replace('public.', 'pg_temp.') for x in tabela + indice]


SQL_HISTORICO_ARQ = os.path.join(RAIZ, 'sql', 'sinais_historico.sql')


def _ddl_historico_em_temp():
    with open(SQL_HISTORICO_ARQ, encoding='utf-8') as f:
        texto = f.read()
    tabela = re.findall(r'^CREATE TABLE public\.sinal_historico \(.*?^\);', texto, re.S | re.M)
    indice = re.findall(r'^CREATE INDEX ix_sinal_historico.*?;$', texto, re.M)
    assert len(tabela) == 1 and len(indice) == 1, 'DDL de sql/sinais_historico.sql mudou'
    return [x.replace('public.', 'pg_temp.') for x in tabela + indice]


def _ciente(loja, regra, objeto, motivo='falta_fornecedor', ate=date(2026, 10, 12),
            medida=None, id_=1, texto='', criado=datetime(2026, 10, 5, 9, 0)):
    return (id_, loja, regra, objeto, motivo, None, ate, medida, texto, 'larissa', criado)


class Ciente(unittest.TestCase):
    def _full(self, full, eid='E1', regra='full_cobertura'):
        return sd._sinal(LPT, 'FULL', 'MLB1', ['K-L-0421-A'], 'x', 's', 100,
                         regra=regra, objeto=eid, medida=full)

    def test_silenciado_sai_dos_ativos_e_vai_para_silenciados(self):
        ativos, sil = sd.aplicar_cientes([self._full(6)],
                                         [_ciente(LPT, 'full_cobertura', 'E1', medida=6)], HOJE)
        self.assertEqual(ativos, [])
        self.assertEqual(len(sil), 1)
        self.assertEqual(sil[0]['ciente']['motivo'], 'falta_fornecedor')

    def test_volta_quando_piora_full_zerou(self):
        ativos, sil = sd.aplicar_cientes([self._full(0)],
                                         [_ciente(LPT, 'full_cobertura', 'E1', medida=6)], HOJE)
        self.assertEqual(sil, [])
        self.assertIn('piorou: 6 un. no Full → 0 un. no Full', ativos[0]['voltou'])
        self.assertIn('Falta no fornecedor', ativos[0]['voltou'])

    def test_volta_quando_a_data_passa_com_selo_por_7_dias(self):
        c = [_ciente(LPT, 'full_cobertura', 'E1', medida=6, ate=date(2026, 10, 1))]
        ativos, _ = sd.aplicar_cientes([self._full(6)], c, HOJE)
        self.assertIn('silêncio venceu em 01/10', ativos[0]['voltou'])
        ativos, _ = sd.aplicar_cientes([self._full(6)], c, date(2026, 10, 20))
        self.assertEqual(ativos[0]['voltou'], '')
        ativos, sil = sd.aplicar_cientes([self._full(6)], c, date(2026, 10, 1))
        self.assertEqual(len(sil), 1)                      # no último dia ainda calado

    def test_chave_e_regra_nao_tipo(self):
        zerado = self._full(0, regra='full_zerado_ads')
        ativos, sil = sd.aplicar_cientes([zerado],
                                         [_ciente(LPT, 'full_cobertura', 'E1', medida=6)], HOJE)
        self.assertEqual(len(ativos), 1)                   # outra regra: não é calada
        outra_loja = sd._sinal(NALA, 'FULL', 'MLB1', [], 'x', 's', 1, regra='full_cobertura',
                               objeto='E1', medida=6)
        ativos, _ = sd.aplicar_cientes([outra_loja],
                                       [_ciente(LPT, 'full_cobertura', 'E1', medida=6)], HOJE)
        self.assertEqual(len(ativos), 1)

    def test_piora_por_regra(self):
        casos = [('full_cobertura', 0, 6, True), ('full_cobertura', 2, 6, False),
                 ('full_cobertura', 0, 0, False),
                 ('full_zerado_ads', 20, 10, True), ('full_zerado_ads', 19, 10, False),
                 ('ads_sem_venda', 60, 30, True), ('ads_sem_venda', 50, 30, False),
                 ('ads_espiral', 150, 100, True), ('ads_espiral', 140, 100, False),
                 ('full_ruptura', 150, 100, True), ('ads_custo', 3, 2, True),
                 ('vendas_queda', -0.75, -0.55, True), ('vendas_queda', -0.70, -0.55, False),
                 ('vendas_queda', -1.0, -0.90, True), ('visitas_queda', -0.80, -0.55, True),
                 ('exp_anuncio', 30, 65, True), ('exp_anuncio', 65, 65, False),
                 ('exp_familia', 5, 4, True), ('exp_familia', 4, 4, False),
                 ('vendas_alta', 9, 1, False), ('visitas_alta', 9, 1, False),
                 ('full_envio', 9, 1, False), ('ads_config', None, None, False)]
        for regra, agora, antes, esperado in casos:
            self.assertEqual(sd.piorou(regra, agora, antes), esperado, (regra, agora, antes))

    def test_objeto_de_evento_leva_a_data(self):
        config = [(LPT, 'MLB5 Pote', 'MLB5 Pote', date(2026, 10, 4), HOJE, 14, 10, 12, 12)]
        s = sd.sinais_ads({}, config, {}, {}, sd.janelas(HOJE))
        self.assertEqual((s[0]['regra'], s[0]['objeto']), ('ads_config', 'MLB5@2026-10-05'))
        _a, por_sku = sd.agregar_vendas([])
        f = sd.sinais_full([(LPT, 'E1', date(2026, 10, 4), 50, 0, 40)],
                           [(LPT, 'MLB1', 'E1', 'A', True)], por_sku, {}, sd.janelas(HOJE))
        self.assertEqual((f[0]['regra'], f[0]['objeto']), ('full_envio', 'E1@2026-10-04'))

    def test_todo_sinal_tem_regra_conhecida(self):
        dados = {'vendas': [_vend(NALA, 'MLB7', 'X', rec_sem=150, rec_base=3200),
                            _vend(LPT, 'MLB1', 'A', rec_sem=0, rec_base=4000)],
                 'ads': [(NALA, 'MLB7', 10, 300, 0, 0, 0, 0, 0, 0, 0)],
                 'estoque_semana': [], 'ponte': [], 'foto': [], 'config': [],
                 'experiencia': [(LPT, 'MLB1', HOJE, 30, 'Ruim', 100, 'Boa')]}
        s, _ = sd.montar_sinais(dados, HOJE)
        self.assertTrue(s)
        for x in s:
            self.assertIn(x['regra'], sd.REGRAS)
            self.assertTrue(x['objeto'])

    def test_validacao(self):
        self.assertEqual(sd.validar_ciente('proposital', '', HOJE, HOJE), [])
        self.assertTrue(sd.validar_ciente('outro', '  ', HOJE, HOJE))
        self.assertTrue(sd.validar_ciente('xpto', 'n', HOJE, HOJE))
        self.assertTrue(sd.validar_ciente('proposital', '', date(2026, 10, 4), HOJE))
        self.assertEqual(sd.validar_ciente('proposital', '', date(2026, 12, 4), HOJE), [])
        self.assertTrue(sd.validar_ciente('proposital', '', date(2026, 12, 5), HOJE))

    def test_gravar_recusa_sem_tocar_no_banco(self):
        conn = mock.Mock()
        s = [self._full(6)]
        with self.assertRaises(ValueError):
            _gravar(conn, s, 'outro', '', HOJE, 'larissa', HOJE, [LPT])
        with self.assertRaises(PermissionError):
            _gravar(conn, s, 'proposital', '', HOJE, 'larissa', HOJE, [NALA])
        with self.assertRaises(ValueError):
            _gravar(conn, s, 'proposital', '', HOJE, '', HOJE, [LPT])
        conn.cursor.assert_not_called()

    def test_listas_do_codigo_iguais_as_check_do_sql(self):
        with open(SQL_CIENTE, encoding='utf-8') as f:
            texto = f.read()
        regra = re.search(r'sinal_ciente_regra_valida CHECK \(regra IN \((.*?)\)\)', texto, re.S)
        motivo = re.search(r'sinal_ciente_motivo_valido CHECK \(motivo IN \((.*?)\)\)', texto, re.S)
        self.assertEqual(set(re.findall(r"'(\w+)'", regra.group(1))), set(sd.REGRAS))
        self.assertEqual(set(re.findall(r"'(\w+)'", motivo.group(1))), set(sd.MOTIVOS))
        self.assertIn(f"criado_em::date + {sd.DIAS_SILENCIO_MAX})", texto)

    def test_quem_pode_dar_ciente(self):
        import permissoes
        for perfil, pode in (('ADMIN', True), ('CONTROLADORIA', True), ('COMPRAS', True),
                             ('GESTOR', True), ('DIRETOR', False), ('', False)):
            with mock.patch.object(permissoes, '_get_role', return_value=perfil):
                self.assertEqual(sd.pode_dar_ciente(), pode, perfil)

    def _render(self, perfil, cientes):
        import permissoes
        dados = {'vendas': [_vend(LPT, 'MLB1', 'A', rec_sem=0, rec_base=4000),
                            _vend(LPT, 'MLB2', 'B', rec_sem=0, rec_base=8000)],
                 'resumo': [], 'metas': [], 'frescor': [], 'visitas': None,
                 }
        st = mock.MagicMock()
        st.columns.side_effect = _colunas
        st.data_editor.side_effect = lambda df, **kw: df
        st.file_uploader.return_value = None
        with mock.patch.object(sd, 'restricao_de_lojas', return_value=None), \
             mock.patch.object(sd, '_ler', return_value=[(LPT,)]), \
             mock.patch.object(sd, 'ler_tudo', return_value=(dados, {})), \
             mock.patch.object(sd, 'ler_cientes', return_value=cientes), \
             mock.patch.object(sd.ep, 'ler_contagens', return_value=[]), \
             mock.patch.object(sd.ep, 'ler_nomes', return_value={}), \
             mock.patch.object(permissoes, '_get_role', return_value=perfil):
            sd._render_mercado_livre(st, mock.MagicMock())
        return st

    def test_sinal_que_piorou_nao_fica_em_silenciados(self):
        ativos, sil = sd.aplicar_cientes(
            [self._full(0)], [_ciente(LPT, 'full_cobertura', 'E1', medida=6, id_=7)], HOJE)
        cientes = [dict(zip(sd.CIENTE_COLS,
                            _ciente(LPT, 'full_cobertura', 'E1', medida=6, id_=7)))]
        self.assertEqual(ativos[0]['ciente_voltou'], 7)
        self.assertTrue(sd.tabela_silenciados(sil, cientes, HOJE, {}, ativos).empty)
        self.assertEqual(sd.cientes_silenciando(cientes, ativos, HOJE), [])
        # sem a piora, continua listado
        ativos, sil = sd.aplicar_cientes(
            [self._full(6)], [_ciente(LPT, 'full_cobertura', 'E1', medida=6, id_=7)], HOJE)
        self.assertEqual(len(sd.tabela_silenciados(sil, cientes, HOJE, {}, ativos)), 1)

    def test_campanha_sem_mlb_com_nome_comprido_cabe_na_chave(self):
        nome = 'Campanha de natal ' * 20
        config = [(LPT, nome, nome, date(2026, 10, 4), HOJE, 14, 10, 12, 12)]
        s = sd.sinais_ads({}, config, {}, {}, sd.janelas(HOJE))
        self.assertRegex(s[0]['objeto'], r'^camp:[0-9a-f]{12}@2026-10-05$')
        self.assertEqual(s[0]['objeto'], sd.objeto_campanha(None, nome, HOJE))
        self.assertNotEqual(s[0]['objeto'], sd.objeto_campanha(None, nome + 'x', HOJE))

    def test_reativar_exige_usuario(self):
        conn = mock.Mock()
        with self.assertRaises(ValueError):
            sd.reativar_ciente(conn, 1, '  ', [LPT])
        conn.cursor.assert_not_called()

    def test_tela_gestora_marca_ciente_diretor_so_ve(self):
        silenciado = [_ciente(LPT, 'vendas_queda', 'MLB2', medida=-1.0)]
        st = self._render('GESTOR', silenciado)
        df = st.data_editor.call_args.args[0]
        self.assertEqual(list(df['Anúncio']), ['MLB1'])     # o silenciado saiu
        self.assertEqual(list(df.columns[:1]), ['Ciente'])
        self.assertTrue(any('Silenciados (1)' in str(c.args[0])
                            for c in st.expander.call_args_list))
        st.file_uploader.assert_called()                    # gestora sobe Excel
        st = self._render('DIRETOR', silenciado)
        st.data_editor.assert_not_called()
        st.download_button.assert_called()                  # diretor baixa...
        st.file_uploader.assert_not_called()                # ...e não sobe
        self.assertEqual(st.dataframe.call_count, 2)        # sinais + silenciados

    def test_sem_tabela_a_tela_e_a_v1(self):
        st = self._render('GESTOR', None)
        st.data_editor.assert_not_called()
        self.assertEqual(list(st.dataframe.call_args_list[0].args[0]['Anúncio']), ['MLB2', 'MLB1'])


# ============================================================
# v1.2 — sem banco
# ============================================================

class _ConnConta:
    def __init__(self, falha=False):
        self.rollbacks, self.falha = 0, falha

    def cursor(self):
        c = mock.Mock()
        if self.falha:
            c.execute.side_effect = RuntimeError('x')
        c.fetchall.return_value = [(1,)]
        return c

    def rollback(self):
        self.rollbacks += 1


class V12(unittest.TestCase):
    def test_ler_tudo_fecha_a_transacao_no_fim(self):
        c = _ConnConta()
        with mock.patch.object(sd, '_ler', return_value=[(False,)]):
            sd.ler_tudo(c, HOJE, [LPT])
        self.assertEqual(c.rollbacks, 1)                    # solta os locks de leitura

    def test_memoria_por_lojas_do_usuario(self):
        """Teste do auditor (06/10): a chave da memória são as lojas visíveis;
        nenhum gestor recebe a leitura de outro perfil; incompleta não fica."""
        chamadas = []

        def falso(_engine, hoje, lojas, mkt=sd.MARKETPLACE):
            chamadas.append(tuple(lojas))
            return {'dados': {}, 'erros': {}, 'erros_bloco': {}, 'parciais': [],
                    'sinais': [{'loja': l} for l in lojas], 'nomes': {},
                    'lido_em': datetime.now(sd.BRT), 'segundos': 0}

        sd._MEMORIA.clear()
        try:
            with mock.patch.object(sd, '_ler_pacote', side_effect=falso):
                f = sd._memoria('pacote')
                f.clear()
                hoje = date(2026, 10, 6)
                f(object(), hoje, (LPT, NALA, RJ, SP))
                gestor = f(object(), hoje, (NALA,))
                f(object(), hoje, (NALA,))                  # mesma loja: memória
                outro = f(object(), hoje, (LPT,))
                f(object(), hoje, (LPT,), sd.SHOPEE)        # outra aba: outra chave
            self.assertEqual(len(chamadas), 4)
            self.assertEqual({x['loja'] for x in gestor['sinais']}, {NALA})
            self.assertEqual({x['loja'] for x in outro['sinais']}, {LPT})

            def incompleto(_engine, hoje, lojas, mkt=sd.MARKETPLACE):
                r = falso(_engine, hoje, lojas)
                r['erros'] = {'ads': 'x'}
                return r
            with mock.patch.object(sd, '_ler_pacote', side_effect=incompleto):
                for _ in range(2):
                    with self.assertRaises(sd._LeituraIncompleta):
                        f(object(), hoje, (RJ,))
            self.assertEqual(chamadas[-2:], [(RJ,), (RJ,)])  # releu: não ficou na memória
        finally:
            sd.limpar_memoria()
            sd._MEMORIA.clear()

    def test_ler_so_desfaz_a_transacao_no_erro(self):
        c = _ConnConta()
        self.assertEqual(sd._ler(c, 'SELECT 1', {}), [(1,)])
        self.assertEqual(c.rollbacks, 0)                    # sem ida e volta à toa
        c = _ConnConta(falha=True)
        with self.assertRaises(RuntimeError):
            sd._ler(c, 'SELECT 1', {})
        self.assertEqual(c.rollbacks, 1)                    # erro não derruba as próximas

    def test_leitura_incompleta_nao_vai_para_a_memoria(self):
        pacote = {'erros': {'ads': 'x'}, 'erros_bloco': {}}
        with mock.patch.object(sd, '_ler_pacote', return_value=pacote):
            with self.assertRaises(sd._LeituraIncompleta) as e:
                sd._ler_pacote_completo(None, HOJE, (LPT,))
        self.assertIs(e.exception.pacote, pacote)
        ok = {'erros': {}, 'erros_bloco': {}}
        with mock.patch.object(sd, '_ler_pacote', return_value=ok):
            self.assertIs(sd._ler_pacote_completo(None, HOJE, (LPT,)), ok)

    def test_fora_do_streamlit_le_direto_sem_memoria(self):
        with mock.patch.object(sd, '_ler_pacote', return_value={'x': 1}) as ler:
            self.assertEqual(sd.obter_pacote(None, HOJE, [LPT]), ({'x': 1}, False))
            sd.obter_pacote(None, HOJE, [LPT])
        self.assertEqual(ler.call_count, 2)

    def test_todo_sinal_tem_periodo(self):
        dados = {'vendas': [_vend(NALA, 'MLB7', 'X', rec_sem=150, rec_base=3200),
                            _vend(LPT, 'MLB1', 'A', rec_sem=0, rec_base=4000, qtd_ritmo=21,
                                  qtd_30=65)],
                 'ads': [(NALA, 'MLB7', 10, 300, 5, 31, 3, 0, 80, 10, 200)],
                 'estoque_semana': [], 'config': [
                     (LPT, 'MLB5 Pote', 'MLB5 Pote', date(2026, 10, 4), HOJE, 14, 10, 12, 12)],
                 'ponte': [(LPT, 'MLB1', 'E1', 'A', True)],
                 'foto': [(LPT, 'E1', date(2026, 10, 4), 6, 2, 3)],
                 'experiencia': [(LPT, 'MLB1', HOJE, 30, 'Ruim', 100, 'Boa')]}
        s, _ = sd.montar_sinais(dados, HOJE)
        self.assertGreaterEqual(len({x['regra'] for x in s}), 6)
        for x in s:
            self.assertTrue(x['periodo'], x['regra'])
        por = {x['regra']: x['periodo'] for x in s}
        self.assertEqual(por['vendas_queda'], '28/09–04/10 × 31/08–27/09')
        self.assertEqual(por['full_cobertura'], 'foto 04/10 · venda 28/09–04/10')
        self.assertEqual(por['full_envio'], 'foto 04/10')
        self.assertEqual(por['ads_config'], 'foto 04/10 → 05/10')
        self.assertEqual(por['ads_sem_venda'], '02/10–04/10')
        self.assertEqual(por['exp_anuncio'], 'nota 28/09 → 05/10')
        self.assertIn('Período', sd.tabela_sinais(s, {}).columns)

    def test_tempos_so_para_admin(self):
        import permissoes
        for perfil, ve in (('ADMIN', True), ('GESTOR', False), ('DIRETOR', False)):
            st = mock.MagicMock()
            with mock.patch.object(permissoes, '_get_role', return_value=perfil):
                sd._render_tempos(st, [('Total', 1.0)], {'segundos': 0.5}, False)
            self.assertEqual(st.expander.called, ve, perfil)


class Shopee(unittest.TestCase):
    def test_fontes_da_shopee(self):
        f = sd.fontes(sd.SHOPEE)
        self.assertIs(f['vendas'], sd.SQL_VENDAS_ANUNCIO_SHOPEE)
        self.assertNotIn('experiencia', f)
        self.assertIs(sd.fontes(sd.MARKETPLACE)['vendas'], sd.SQL_VENDAS_ANUNCIO)
        self.assertEqual(sd.params(HOJE, [S_LPT], sd.SHOPEE)['marketplace'], 'SHOPEE')

    def test_roas_zero_e_automatico_so_na_shopee(self):
        self.assertEqual(sd._roas(0, sd.SHOPEE), 'automático')
        self.assertEqual(sd._roas(0, sd.MARKETPLACE), '0,0')
        self.assertEqual(sd._roas(9.8, sd.SHOPEE), '9,8')

    def test_rating_zero_com_avaliacoes_e_dado_invalido(self):
        self.assertEqual(sd.nota_das_novas(0, 15, 4.8, 10), (None, 0))
        self.assertEqual(sd.nota_das_novas(4.5, 20, 0, 10), (None, 0))
        self.assertEqual(sd.nota_das_novas(4.5, 20, 0, 0), (4.5, 20))   # sem avaliação antes
        caiu = self._views((S_LPT, '2249', date(2026, 10, 4), 400, 4.5, 20, 400, 4.8, 10))
        self.assertIn('com comentário: aproximação', caiu[0]['numero'])

    def test_nota_so_das_avaliacoes_novas(self):
        nota, novas = sd.nota_das_novas(4.5, 20, 4.8, 10)
        self.assertAlmostEqual(nota, 4.2)
        self.assertEqual(novas, 10)
        self.assertEqual(sd.nota_das_novas(4.8, 10, 4.8, 10), (None, 0))
        self.assertEqual(sd.nota_das_novas(None, 10, 4.8, 10), (None, 0))

    def _views(self, linha, rec_base=4000):
        por_an, _ = sd.agregar_vendas([_vend(S_LPT, '2249', 'A', rec_base=rec_base, rec_30=900,
                                             qtd_30=10) + (9,)])
        return sd.sinais_views_shopee([linha], por_an, sd.janelas(HOJE))

    def test_views_d_contra_d30_exato(self):
        s = self._views((S_LPT, '2249', date(2026, 10, 4), 100, 4.8, 10, 400, 4.8, 10))
        self.assertEqual([x['regra'] for x in s], ['visitas_queda'])
        self.assertIn('Visualizações de página (30d)', s[0]['numero'])
        self.assertIn('pedidos por visualização', s[0]['numero'])
        self.assertNotIn('visita', s[0]['numero'].lower().replace('visualiza', ''))
        # sem a foto de D-30 (ref None): não acende
        self.assertEqual(self._views((S_LPT, '2249', date(2026, 10, 4), 100, 4.8, 10,
                                      None, None, None)), [])
        self.assertEqual(self._views((S_LPT, '2249', date(2026, 10, 4), 900, 4.8, 10,
                                      400, 4.8, 10))[0]['regra'], 'visitas_alta')

    def test_nota_precisa_de_3_novas_e_queda_de_meio_ponto(self):
        poucas = self._views((S_LPT, '2249', date(2026, 10, 4), 400, 4.0, 12, 400, 4.8, 10))
        self.assertEqual(poucas, [])                         # 2 avaliações novas
        pouca_queda = self._views((S_LPT, '2249', date(2026, 10, 4), 400, 4.75, 20, 400, 4.8, 10))
        self.assertEqual(pouca_queda, [])                    # novas 4,7: só 0,1 abaixo
        caiu = self._views((S_LPT, '2249', date(2026, 10, 4), 400, 4.5, 20, 400, 4.8, 10))
        self.assertEqual(caiu[0]['regra'], 'exp_anuncio')
        self.assertIn('10 avaliações novas', caiu[0]['numero'])

    def test_custo_por_pedido_na_shopee_aceita_varios_skus(self):
        vendas = [_vend(S_LPT, '2249', 'A', qtd_30=5, margem_30=20) + (5,),
                  _vend(S_LPT, '2249', 'B', qtd_30=5, margem_30=20) + (5,)]   # R$ 4/pedido
        ads = {(S_LPT, '2249'): _ads(gasto_30=80, unid_ads_30=8, cliques_30=100)}  # R$ 10/pedido
        por_an, _ = sd.agregar_vendas(vendas)
        s = sd.sinais_ads(ads, [], por_an, {}, sd.janelas(HOJE), mkt=sd.SHOPEE)
        self.assertEqual(len(s), 1)
        self.assertIn('Custo por PEDIDO de ads R$ 10,00', s[0]['numero'])
        self.assertIn('margem por PEDIDO R$ 4,00', s[0]['numero'])
        self.assertIn('pedidos de ads, não unidades', s[0]['numero'])
        # o mesmo cenário no ML (por unidade, SKU único): não julga
        self.assertEqual(sd.sinais_ads(ads, [], por_an, {}, sd.janelas(HOJE)), [])

    def test_sem_estoque_no_banco_espiral_nao_e_julgada(self):
        vendas = [_vend(S_LPT, '2249', 'A', rec_sem=100, rec_base=4000)]
        ads = {(S_LPT, '2249'): _ads(gasto_sem=5, gasto_base=400)}
        por_an, _ = sd.agregar_vendas(vendas)
        s = sd.sinais_ads(ads, [], por_an, {}, sd.janelas(HOJE), mkt=sd.SHOPEE,
                          lojas_com_estoque=set())
        self.assertEqual(s, [])                              # nem espiral, nem "ruptura"

    def test_config_da_shopee_usa_o_item_id_do_detalhe(self):
        config = [(S_LPT, '106156986', 'Kit Escova', date(2026, 10, 4), HOJE, 5, 0, 10, 10,
                   '2249')]
        s = sd.sinais_ads({}, config, {}, {}, sd.janelas(HOJE), mkt=sd.SHOPEE)
        self.assertEqual((s[0]['anuncio'], s[0]['objeto']), ('2249', '2249@2026-10-05'))
        self.assertIn('5,0 → automático', s[0]['numero'])

    def test_litstore_por_sku(self):
        x = sd._sinal(S_YANNI, 'VENDAS', 'SKU:L-0500', ['L-0500'], 'Queda', 's', 1,
                      regra='vendas_queda')
        self.assertEqual(x['objeto'], 'SKU:L-0500')
        self.assertEqual(list(sd.tabela_sinais([x], {})['Anúncio']), ['(por SKU)'])

    def test_carga_pendente_ignora_loja_de_upload(self):
        linhas = [('Vendas', S_LPT, datetime(2026, 10, 4, 7, 0)), ('Vendas', S_YANNI, None)]
        self.assertEqual(sd.lojas_carga_pendente(linhas, [S_LPT, S_YANNI], HOJE), {S_LPT})

    def test_aba_shopee_mostra_em_construcao_e_o_aviso_de_pedido(self):
        import permissoes
        dados = {'vendas': [_vend(S_LPT, '2249', 'A', rec_sem=0, rec_base=4000)],
                 'resumo': [], 'metas': [], 'frescor': [], 'visitas': None, 'foto': []}
        st = mock.MagicMock()
        st.columns.side_effect = _colunas
        st.file_uploader.return_value = None
        with mock.patch.object(sd, 'restricao_de_lojas', return_value=None), \
             mock.patch.object(sd, '_ler', return_value=[(S_LPT,)]), \
             mock.patch.object(sd, 'ler_tudo', return_value=(dados, {})) as lt, \
             mock.patch.object(sd, 'ler_cientes', return_value=[]) as lc, \
             mock.patch.object(sd.ep, 'ler_contagens') as contagens, \
             mock.patch.object(sd.ep, 'ler_nomes', return_value={}), \
             mock.patch.object(permissoes, '_get_role', return_value='GESTOR'):
            sd._render_marketplace(st, mock.MagicMock(), sd.SHOPEE)
        self.assertEqual(lt.call_args.args[3], sd.SHOPEE)
        self.assertEqual(lc.call_args.args[2], sd.SHOPEE)
        contagens.assert_not_called()                        # fotos parciais: só ML
        infos = ' '.join(str(c.args[0]) for c in st.info.call_args_list)
        self.assertIn('Estoque/Full da Shopee: em construção', infos)
        self.assertIn('Visualizações de página (30d) e nota da Shopee: em construção', infos)
        caps = ' '.join(str(c.args[0]) for c in st.caption.call_args_list)
        self.assertIn('ads informa PEDIDOS, não unidades', caps)
        self.assertEqual(st.subheader.call_args.args[0], S_LPT)


class FullGalpao(unittest.TestCase):
    """Regra do Thiago (06/10), igual no ML e na Shopee: sem Full o anúncio
    ainda vende pelo galpão, mas muito menos; Full E galpão zerados = ruptura."""

    def _full(self, foto, ponte, vendas, ads=None, mkt=sd.SHOPEE):
        _a, por_sku = sd.agregar_vendas(vendas)
        return sd.sinais_full(foto, ponte, por_sku, ads or {}, sd.janelas(HOJE), mkt)

    def test_k10_l_7248_a_full_zerado_com_ads_e_galpao_nao_manda_pausar(self):
        ponte = [(S_NALA, '23298770778', 'E7248', 'K10-L-7248-A', True)]
        foto = [(S_NALA, 'E7248', date(2026, 10, 4), 0, 0, 0, 18)]
        vendas = [_vend(S_NALA, '23298770778', 'K10-L-7248-A', qtd_ritmo=7, qtd_30=30,
                        rec_30=900)]
        ads = {(S_NALA, '23298770778'): _ads(gasto_ontem=16)}
        for mkt in (sd.SHOPEE, sd.MARKETPLACE):             # a MESMA regra nos dois
            s = self._full(foto, ponte, vendas, ads, mkt)
            self.assertEqual([x['regra'] for x in s], ['full_zerado_ads'], mkt)
            self.assertTrue(s[0]['urgente'])
            self.assertIn('galpão 18 un.', s[0]['numero'])
            self.assertIn('queda forte esperada', s[0]['numero'])
            self.assertIn('avalie reduzir', s[0]['sugestao'])
            self.assertNotIn('ausar o ads', s[0]['sugestao'].replace('não pausar', ''))
            self.assertEqual(s[0]['em_jogo'], 900)           # a venda de 30 dias

    def test_l_0321_full_curto_com_galpao_continua_urgente(self):
        ponte = [(S_LPT, '22494216736', 'E321', 'L-0321', True)]
        foto = [(S_LPT, 'E321', date(2026, 10, 4), 29, 0, 0, 2328)]
        vendas = [_vend(S_LPT, '22494216736', 'L-0321', qtd_ritmo=210, qtd_30=600,
                        rec_30=12000)]
        s = self._full(foto, ponte, vendas)
        self.assertEqual([x['regra'] for x in s], ['full_cobertura'])
        self.assertTrue(s[0]['urgente'])
        self.assertIn('1 dias de Full', s[0]['numero'])
        self.assertIn('galpão 2328 un. (sem Full vende muito menos', s[0]['numero'])
        self.assertIn('Repor o Full já', s[0]['sugestao'])
        self.assertEqual(s[0]['em_jogo'], 12000)

    def test_k_l_0421_a_full_zerado_com_galpao_e_ads(self):
        ponte = [(S_LPT, '46662845939', 'E421', 'K-L-0421-A', True)]
        foto = [(S_LPT, 'E421', date(2026, 10, 4), 0, 0, 0, 364)]
        ads = {(S_LPT, '46662845939'): _ads(gasto_ontem=5)}
        s = self._full(foto, ponte, [], ads)
        self.assertEqual([x['regra'] for x in s], ['full_zerado_ads'])
        self.assertIn('galpão 364 un.', s[0]['numero'])
        self.assertIn('não pausar', s[0]['sugestao'])

    def test_full_e_galpao_zerados_e_ruptura_e_ai_sim_pausar(self):
        ponte = [(S_NALA, 'I1', 'E1', 'A', True)]
        foto = [(S_NALA, 'E1', date(2026, 10, 4), 0, 0, 0, 0)]
        ads = {(S_NALA, 'I1'): _ads(gasto_ontem=10)}
        vendas = [_vend(S_NALA, 'I1', 'A', qtd_ritmo=7, qtd_30=30, rec_30=800)]
        for mkt in (sd.SHOPEE, sd.MARKETPLACE):
            s = self._full(foto, ponte, vendas, ads, mkt)
            self.assertEqual([x['regra'] for x in s], ['full_ruptura'], mkt)
            self.assertTrue(s[0]['numero'].startswith('RUPTURA — Full e galpão zerados'))
            self.assertIn('Pausar o ads até repor', s[0]['sugestao'])
            self.assertEqual(s[0]['em_jogo'], 800)
        sem_ads = self._full(foto, ponte, vendas)
        self.assertEqual(sem_ads[0]['sugestao'], 'Repor o estoque já')
        parado = self._full(foto, ponte, [])                 # sem venda e sem ads: nada
        self.assertEqual(parado, [])

    def test_caso_ml_full_zerado_com_galpao_nao_pausa(self):
        ponte = [(LPT, 'MLB5183384677', 'MLBU5071722523', 'K-L-0421-A', True)]
        foto = [(LPT, 'MLBU5071722523', date(2026, 10, 4), 0, 2, 0, 1315)]
        vendas = [_vend(LPT, 'MLB5183384677', 'K-L-0421-A', qtd_ritmo=21, qtd_30=65,
                        rec_30=2000)]
        ads = {(LPT, 'MLB5183384677'): _ads(gasto_ontem=19)}
        s = self._full(foto, ponte, vendas, ads, sd.MARKETPLACE)
        self.assertEqual(s[0]['regra'], 'full_zerado_ads')
        self.assertIn('2 em transferência', s[0]['numero'])
        self.assertIn('galpão 1315 un.', s[0]['numero'])
        self.assertIn('não pausar', s[0]['sugestao'])

    def test_galpao_desconhecido_nao_vira_ruptura(self):
        ponte = [(LPT, 'MLB1', 'E1', 'A', True)]
        foto = [(LPT, 'E1', date(2026, 10, 4), 0, 0, 0)]          # sem coluna de galpão
        ads = {(LPT, 'MLB1'): _ads(gasto_ontem=10)}
        s = self._full(foto, ponte, [], ads, sd.MARKETPLACE)
        self.assertEqual([x['regra'] for x in s], ['full_zerado_ads'])
        self.assertNotIn('galpão', s[0]['numero'])

    def test_espiral_e_ruptura_olham_o_estoque_total(self):
        ponte = [(S_LPT, 'I1', 'E1', 'A', True)]
        semana = [(S_LPT, 'E1', date(2026, 9, 28) + timedelta(days=i), 0, 50) for i in range(7)]
        self.assertEqual(sd.dias_com_estoque(semana, ponte), {(S_LPT, 'I1'): 7})
        self.assertEqual(sd.dias_com_estoque(semana, ponte, so_full=True), {})
        for mkt in (sd.SHOPEE, sd.MARKETPLACE):
            dados = {'vendas': [_vend(S_LPT, 'I1', 'A', rec_sem=100, rec_base=4000)],
                     'ads': [(S_LPT, 'I1', 5, 400, 0, 0, 0, 0, 0, 0, 0)], 'config': [],
                     'ponte': ponte, 'estoque_semana': semana,
                     'foto': [(S_LPT, 'E1', date(2026, 10, 4), 0, 0, 0, 50)]}
            s, _ = sd.montar_sinais(dados, HOJE, mkt)
            regras = {x['regra'] for x in s}
            self.assertIn('ads_espiral', regras, mkt)         # tinha estoque (galpão)
            self.assertNotIn('full_ruptura', regras, mkt)

    def test_texto_da_ruptura_mostra_os_dias_com_full(self):
        vendas = [_vend(NALA, 'MLB7', 'X', rec_sem=150, rec_base=3200)]
        ads = {(NALA, 'MLB7'): _ads(gasto_sem=10, gasto_base=300)}
        por_an, _ = sd.agregar_vendas(vendas)
        s = sd.sinais_ads(ads, [], por_an, {(NALA, 'MLB7'): 4}, sd.janelas(HOJE),
                          dias_full={(NALA, 'MLB7'): 2})
        self.assertIn('estoque (Full ou galpão) em só 4 de 7 dias da semana, Full em 2',
                      s[0]['numero'])


class RupturaEGalpaoZero(unittest.TestCase):
    def test_galpao_zero_com_full_diz_que_para_de_vender(self):
        # caso real: ML-LPT MLBU3753467095 (Full > 0, galpão 0)
        ponte = [(LPT, 'MLB1', 'MLBU3753467095', 'A', True)]
        foto = [(LPT, 'MLBU3753467095', date(2026, 10, 4), 6, 0, 0, 0)]
        _a, por_sku = sd.agregar_vendas([_vend(LPT, 'MLB1', 'A', qtd_ritmo=21, qtd_30=60)])
        s = sd.sinais_full(foto, ponte, por_sku, {}, sd.janelas(HOJE))
        self.assertIn('galpão 0: quando o Full acabar, para de vender', s[0]['numero'])
        self.assertNotIn('vende muito menos', s[0]['numero'])

    def test_ruptura_das_duas_origens_vira_uma_linha_por_anuncio(self):
        ponte = [(NALA, 'MLB7', 'E1', 'X-AZUL', True), (NALA, 'MLB7', 'E2', 'X-ROSA', True)]
        foto = [(NALA, 'E1', date(2026, 10, 4), 0, 0, 0, 0),
                (NALA, 'E2', date(2026, 10, 4), 0, 0, 0, 0)]
        dados = {'vendas': [_vend(NALA, 'MLB7', 'X-AZUL', rec_sem=100, rec_base=2400, qtd_ritmo=7,
                                  qtd_30=30, rec_30=700),
                            _vend(NALA, 'MLB7', 'X-ROSA', rec_sem=50, rec_base=1600, qtd_ritmo=7,
                                  qtd_30=30, rec_30=500)],
                 'ads': [(NALA, 'MLB7', 10, 300, 0, 0, 0, 0, 0, 0, 0)], 'config': [],
                 'ponte': ponte, 'foto': foto, 'estoque_semana': []}
        s, erros = sd.montar_sinais(dados, HOJE)
        self.assertEqual(erros, {})
        rup = [x for x in s if x['regra'] == 'full_ruptura']
        self.assertEqual(len(rup), 1)                       # foto (2 variações) + semana
        r = rup[0]
        self.assertEqual(r['objeto'], 'MLB7')                # o anúncio, nas duas origens
        self.assertTrue(r['numero'].startswith('RUPTURA — Full e galpão zerados'))
        self.assertIn('(+1 variação(ões) do anúncio)', r['numero'])
        self.assertIn('na semana: estoque (Full ou galpão) em só 0 de 7', r['numero'])
        self.assertEqual(r['em_jogo'], 1200)                 # 700 + 500 (> queda da semana)
        self.assertEqual(r['medida'], 1200)
        self.assertEqual(set(r['skus']), {'X-AZUL', 'X-ROSA'})
        self.assertFalse([x for x in s if x['regra'] == 'vendas_queda'])  # explicada

    def test_um_ciente_cala_a_ruptura_venha_de_onde_vier(self):
        ruptura_semana = sd._sinal(NALA, 'FULL', 'MLB7', [], 'Ruptura: x', 's', 100,
                                   regra='full_ruptura', objeto='MLB7', medida=100)
        ruptura_hoje = sd._sinal(NALA, 'FULL', 'MLB7', [], 'RUPTURA — y', 's', 100,
                                 regra='full_ruptura', objeto='MLB7', medida=100)
        ciente = [(1, NALA, 'full_ruptura', 'MLB7', 'falta_fornecedor', None,
                   date(2026, 10, 12), 100, '', 'larissa', datetime(2026, 10, 5, 9))]
        for x in (ruptura_semana, ruptura_hoje):
            ativos, sil = sd.aplicar_cientes([x], ciente, HOJE)
            self.assertEqual((ativos, len(sil)), ([], 1))


class UploadAtrasado(unittest.TestCase):
    def test_suspende_so_acima_de_7_dias(self):
        self.assertFalse(sd.upload_suspenso(date(2026, 9, 27), HOJE))   # 7 dias antes de ontem
        self.assertTrue(sd.upload_suspenso(date(2026, 9, 26), HOJE))

    def test_relê_a_litstore_com_janelas_ate_o_ultimo_upload(self):
        frescor = [('Vendas', S_YANNI, None),
                   ('Último dia de venda (upload)', S_YANNI, datetime(2026, 10, 1))]
        chamadas = []

        def ler(conn, sql, p):
            chamadas.append(p)
            return [_vend(S_YANNI, 'SKU:L-0500', 'L-0500', rec_sem=100, rec_base=4000)]
        antigo = sd._sinal(S_YANNI, 'VENDAS', 'SKU:L-0500', [], 'Queda falsa', 's', 1,
                           regra='vendas_queda', objeto='SKU:L-0500')
        outro = sd._sinal(S_LPT, 'VENDAS', '2249', [], 'x', 's', 1, regra='vendas_queda')
        with mock.patch.object(sd, '_ler', side_effect=ler):
            sinais, ate = sd.reler_vendas_upload(None, {'frescor': frescor}, [antigo, outro],
                                                 HOJE, [S_YANNI, S_LPT], sd.SHOPEE)
        self.assertEqual(ate, {S_YANNI: date(2026, 10, 1)})
        self.assertEqual(chamadas[0]['fim'], date(2026, 10, 1))   # janela até o último upload
        self.assertEqual(chamadas[0]['lojas'], [S_YANNI])
        self.assertNotIn(antigo, sinais)
        self.assertIn(outro, sinais)                          # outra loja: intacta
        nova = [x for x in sinais if x['loja'] == S_YANNI][0]
        self.assertEqual(nova['periodo'], '25/09–01/10 × 28/08–24/09')

    def test_upload_muito_atrasado_nao_e_relido(self):
        frescor = [('Último dia de venda (upload)', S_YANNI, datetime(2026, 9, 20))]
        with mock.patch.object(sd, '_ler') as ler:
            _s, ate = sd.reler_vendas_upload(None, {'frescor': frescor}, [], HOJE, [S_YANNI],
                                             sd.SHOPEE)
        ler.assert_not_called()
        self.assertEqual(ate, {})

    def test_litstore_com_ultima_venda_antes_de_ontem(self):
        linhas = [('Vendas', S_LPT, datetime(2026, 10, 5, 4, 20)),
                  ('Vendas', S_YANNI, None),
                  ('Último dia de venda (upload)', S_YANNI, datetime(2026, 10, 3))]
        self.assertEqual(sd.lojas_upload_atrasado(linhas, [S_LPT, S_YANNI], HOJE),
                         {S_YANNI: date(2026, 10, 3)})
        em_dia = linhas[:2] + [('Último dia de venda (upload)', S_YANNI, datetime(2026, 10, 4))]
        self.assertEqual(sd.lojas_upload_atrasado(em_dia, [S_LPT, S_YANNI], HOJE), {})

    def test_loja_de_api_nunca_e_upload_atrasado(self):
        linhas = [('Vendas', S_LPT, datetime(2026, 10, 5, 4, 20)),
                  ('Último dia de venda (upload)', S_LPT, datetime(2026, 9, 1))]
        self.assertEqual(sd.lojas_upload_atrasado(linhas, [S_LPT], HOJE), {})

    def test_suspende_as_regras_de_venda_da_litstore(self):
        x = sd._sinal(S_YANNI, 'VENDAS', 'SKU:L-0500', [], 'Queda', 's', 1, regra='vendas_queda')
        self.assertEqual(sd.suspender_venda_pendente([x], {S_YANNI: 'upload:03/10'}), [])


class ApareceDesde(unittest.TestCase):
    def _s(self, loja=LPT, regra='vendas_queda', objeto='MLB1'):
        return sd._sinal(loja, 'VENDAS', objeto, [], 'x', 's', 1, regra=regra, objeto=objeto)

    def _desde(self, historico, dias_job, hoje=date(2026, 10, 6), sinal=None):
        return sd.aparece_desde([sinal or self._s()], historico, dias_job, hoje)[0]

    def test_novo_hoje(self):
        x = self._desde([(date(2026, 10, 5), LPT, 'vendas_queda', 'MLB9')], [date(2026, 10, 5)])
        self.assertEqual((x['desde'], x['desde_txt']), (date(2026, 10, 6), '🆕 hoje'))

    def test_sequencia_ate_ontem(self):
        k = (LPT, 'vendas_queda', 'MLB1')
        hist = [(date(2026, 10, 4),) + k, (date(2026, 10, 5),) + k]
        x = self._desde(hist, [date(2026, 10, 3), date(2026, 10, 4), date(2026, 10, 5)])
        self.assertEqual(x['desde'], date(2026, 10, 4))
        self.assertEqual(x['desde_txt'], 'há 2 dias (desde 04/10)')

    def test_dia_sem_execucao_do_job_nao_quebra(self):
        k = (LPT, 'vendas_queda', 'MLB1')
        hist = [(date(2026, 10, 3),) + k, (date(2026, 10, 5),) + k]
        x = self._desde(hist, [date(2026, 10, 3), date(2026, 10, 5)])   # 04/10 sem job
        self.assertEqual(x['desde'], date(2026, 10, 3))

    def test_dia_com_job_e_sem_o_sinal_quebra(self):
        k = (LPT, 'vendas_queda', 'MLB1')
        hist = [(date(2026, 10, 3),) + k, (date(2026, 10, 5),) + k,
                (date(2026, 10, 4), LPT, 'vendas_queda', 'OUTRO')]
        x = self._desde(hist, [date(2026, 10, 3), date(2026, 10, 4), date(2026, 10, 5)])
        self.assertEqual(x['desde'], date(2026, 10, 5))

    def test_a_chave_e_regra_loja_e_objeto(self):
        hist = [(date(2026, 10, 5), NALA, 'vendas_queda', 'MLB1'),   # outra loja
                (date(2026, 10, 5), LPT, 'vendas_alta', 'MLB1')]     # outra regra
        x = self._desde(hist, [date(2026, 10, 5)])
        self.assertEqual(x['desde_txt'], '🆕 hoje')

    def test_coluna_na_tabela_e_no_excel(self):
        x = sd.aparece_desde([self._s()], [], [], date(2026, 10, 6))
        self.assertEqual(list(sd.tabela_sinais(x, {})['Aparece desde']), ['🆕 hoje'])
        self.assertEqual(list(sd.tabela_sinais([self._s()], {})['Aparece desde']), ['—'])
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(sd.excel_sinais(x, {}, HOJE)))['Sinais']
        cab = [c.value for c in ws[1]]
        self.assertEqual(ws.cell(row=2, column=cab.index('Aparece desde') + 1).value, '🆕 hoje')


class MontagemUnicaEJob(unittest.TestCase):
    def test_tela_e_job_usam_a_mesma_suspensao(self):
        pacote = {'dados': {'frescor': [('Vendas', NALA, datetime(2026, 10, 4, 6, 7))]},
                  'sinais': [sd._sinal(NALA, 'VENDAS', 'MLB1', [], 'x', 's', 1,
                                       regra='vendas_queda'),
                             sd._sinal(NALA, 'ADS', 'MLB1', [], 'x', 's', 1,
                                       regra='ads_sem_venda')],
                  'upload_ate': {}}
        sinais, pendentes, _u = sd.sinais_da_loja_hoje(pacote, [NALA], HOJE)
        self.assertEqual(pendentes, {NALA: 'carga'})
        self.assertEqual([x['regra'] for x in sinais], ['ads_sem_venda'])

    def test_job_nao_le_nomes_e_nao_grava_leitura_incompleta(self):
        ok = {'dados': {'frescor': []}, 'erros': {}, 'erros_bloco': {},
              'sinais': [sd._sinal(LPT, 'VENDAS', 'MLB1', [], 'x', 's', 1, regra='vendas_queda')],
              'upload_ate': {}}
        with mock.patch.object(sd, '_ler_lojas_ml', return_value=[LPT, NALA]), \
             mock.patch.object(sd, '_ler_pacote', return_value=ok) as lp:
            sinais, n = sd.sinais_hoje_para_historico(None, HOJE, sd.MARKETPLACE)
        self.assertEqual((len(sinais), n), (1, 2))
        self.assertIs(lp.call_args.kwargs['com_nomes'], False)
        ruim = dict(ok, erros={'ads': 'x'})
        with mock.patch.object(sd, '_ler_lojas_ml', return_value=[LPT]), \
             mock.patch.object(sd, '_ler_pacote', return_value=ruim):
            with self.assertRaises(sd._LeituraIncompleta):
                sd.sinais_hoje_para_historico(None, HOJE, sd.MARKETPLACE)

    def test_gravar_historico_upsert_numa_transacao(self):
        conn = mock.MagicMock()
        x = sd._sinal(LPT, 'VENDAS', 'MLB1', [], 'x', 's', 12.345, regra='vendas_queda',
                      medida=-0.5)
        self.assertEqual(sd.gravar_historico(conn, [x], HOJE, sd.MARKETPLACE), 1)
        sql, linhas = conn.cursor.return_value.executemany.call_args.args
        self.assertIn('ON CONFLICT (data, marketplace, loja, regra, objeto) DO UPDATE', sql)
        self.assertEqual(linhas[0], {'data': HOJE, 'marketplace': 'MERCADO LIVRE', 'loja': LPT,
                                     'regra': 'vendas_queda', 'objeto': 'MLB1', 'medida': -0.5,
                                     'em_jogo': 12.35})
        conn.commit.assert_called_once()


class JobHistorico(unittest.TestCase):
    def setUp(self):
        sys.path.insert(0, os.path.join(RAIZ, 'jobs'))
        import historico_sinais
        self.job = historico_sinais

    def _rodar(self, argv=(), agora=datetime(2026, 10, 6, 12, 5), efeito=None):
        saida = []
        with mock.patch.object(self.job, 'gravar_dia', side_effect=efeito or (lambda e, d, m: (7, 4))):
            cod = self.job.main(list(argv), agora=agora, engine=object(), saida=saida.append)
        return cod, '\n'.join(saida)

    def test_dias_para_gravar(self):
        f = self.job.dias_para_gravar
        self.assertEqual(f(datetime(2026, 10, 6, 9, 59)), [])          # antes das 10h
        self.assertEqual(f(datetime(2026, 10, 6, 12, 0)), [date(2026, 10, 6)])
        self.assertEqual(f(datetime(2026, 10, 6, 8, 0), date(2026, 10, 3)),
                         [date(2026, 10, 3), date(2026, 10, 4), date(2026, 10, 5)])
        self.assertEqual(f(datetime(2026, 10, 6, 8, 0), date(2026, 10, 4), date(2026, 10, 9)),
                         [date(2026, 10, 4), date(2026, 10, 5)])         # nunca hoje/futuro

    def test_saida_so_com_contagens(self):
        cod, txt = self._rodar()
        self.assertEqual(cod, 0)
        self.assertIn('06/10/2026 Mercado Livre: gravou 7 sinais em 4 lojas.', txt)
        self.assertIn('06/10/2026 Shopee: gravou 7 sinais em 4 lojas.', txt)
        for proibido in ('R$', 'MLB', 'ML-', 'Shopee-', 'L-0', 'postgres', '@'):
            self.assertNotIn(proibido, txt)

    def test_falha_mostra_so_o_tipo_do_erro(self):
        def quebra(e, d, m):
            raise RuntimeError('valor R$ 50 do SKU L-0320 na loja ML-LPT senha=xyz')
        cod, txt = self._rodar(efeito=quebra)
        self.assertEqual(cod, 1)
        self.assertIn('falhou (RuntimeError)', txt)
        for proibido in ('R$', 'L-0320', 'ML-LPT', 'senha', 'Traceback'):
            self.assertNotIn(proibido, txt)

    def test_antes_das_10h_nao_grava(self):
        cod, txt = self._rodar(agora=datetime(2026, 10, 6, 9, 0),
                               efeito=lambda e, d, m: self.fail('não devia gravar'))
        self.assertEqual(cod, 0)
        self.assertIn('Nada a gravar', txt)

    def test_backfill(self):
        dias = []
        cod, _t = self._rodar(['--desde', '2026-10-03', '--ate', '2026-10-04'],
                              efeito=lambda e, d, m: dias.append((d, m)) or (1, 1))
        self.assertEqual(cod, 0)
        self.assertEqual(dias, [(date(2026, 10, 3), sd.MARKETPLACE), (date(2026, 10, 3), sd.SHOPEE),
                                (date(2026, 10, 4), sd.MARKETPLACE), (date(2026, 10, 4), sd.SHOPEE)])

    def test_sem_a_variavel_de_conexao(self):
        saida = []
        with mock.patch.dict(os.environ, {}, clear=True):
            cod = self.job.main([], agora=datetime(2026, 10, 6, 12, 0), saida=saida.append)
        self.assertEqual(cod, 2)
        self.assertIn('SINAIS_HIST_DB_URL', saida[0])

    def test_workflow_do_repo_publico(self):
        with open(os.path.join(RAIZ, '.github', 'workflows', 'sinais_historico.yml'),
                  encoding='utf-8') as f:
            wf = f.read()
        self.assertIn("cron: '0 15 * * *'", wf)
        self.assertIn('workflow_dispatch', wf)
        self.assertIn('contents: read', wf)
        self.assertNotIn('pull_request', wf)
        self.assertIn('secrets.SINAIS_HIST_DB_URL', wf)

    def test_listas_de_regras_iguais_nas_duas_tabelas(self):
        with open(SQL_HISTORICO_ARQ, encoding='utf-8') as f:
            texto = f.read()
        regra = re.search(r'sinal_historico_regra_valida CHECK \(regra IN \((.*?)\)\)', texto, re.S)
        self.assertEqual(set(re.findall(r"'(\w+)'", regra.group(1))), set(sd.REGRAS))


class CargaPendente(unittest.TestCase):
    def test_loja_sem_a_carga_de_hoje_fica_pendente(self):
        linhas = [('Vendas', LPT, datetime(2026, 10, 5, 6, 4)),
                  ('Vendas', NALA, datetime(2026, 10, 4, 6, 7)),
                  ('Estoque', RJ, datetime(2026, 10, 4, 4, 0))]
        self.assertEqual(sd.lojas_carga_pendente(linhas, [LPT, NALA, RJ], HOJE), {NALA})

    def test_suspende_so_os_sinais_de_venda_da_loja_pendente(self):
        sinais = [sd._sinal(NALA, 'VENDAS', 'MLB1', [], 'x', 's', 1, regra='vendas_queda'),
                  sd._sinal(NALA, 'ADS', 'MLB1', [], 'x', 's', 1, regra='ads_sem_venda'),
                  sd._sinal(NALA, 'EXP', 'MLB1', [], 'x', 's', 1, regra='exp_anuncio'),
                  sd._sinal(LPT, 'VENDAS', 'MLB2', [], 'x', 's', 1, regra='vendas_queda')]
        r = sd.suspender_venda_pendente(sinais, {NALA})
        self.assertEqual([(x['loja'], x['regra']) for x in r],
                         [(NALA, 'ads_sem_venda'), (NALA, 'exp_anuncio'), (LPT, 'vendas_queda')])

    def test_resumo_mostra_carga_pendente_sem_comparar(self):
        st = mock.MagicMock()
        cols = [mock.MagicMock(), mock.MagicMock(), mock.MagicMock()]
        st.columns.return_value = cols
        r = sd.resumo_lojas([(NALA, 50, 16000, 12000)], [], [NALA], HOJE)[NALA]
        sd._render_resumo(st, NALA, r, date(2026, 10, 4), pendente=True)
        args, kw = cols[0].metric.call_args
        self.assertEqual(args[1], 'carga pendente')
        self.assertEqual(len(args), 2)                      # sem delta de comparação


class ResumoTexto(unittest.TestCase):
    def test_rotulos_sem_parcial_e_ajuda_do_dia_fechado(self):
        st = mock.MagicMock()
        cols = [mock.MagicMock(), mock.MagicMock(), mock.MagicMock()]
        st.columns.return_value = cols
        r = sd.resumo_lojas([(LPT, 3000, 16000, 12000)], [], [LPT], HOJE)[LPT]
        sd._render_resumo(st, LPT, r, date(2026, 10, 4))
        for col in cols[:2]:
            args, kw = col.metric.call_args
            self.assertNotIn('parcial', args[0].lower())
            self.assertIn('boleto/Pix', kw['help'])
            self.assertIn('cancelamentos', kw['help'])
        self.assertIn('Dia fechado', cols[0].metric.call_args.kwargs['help'])


class Excel(unittest.TestCase):
    def _sinais(self):
        return [sd._sinal(LPT, 'FULL', 'MLB1', ['K-L-0421-A'], '3 dias de Full', 's', 2000,
                          regra='full_cobertura', objeto='E1', medida=6, periodo='foto 04/10'),
                sd._sinal(LPT, 'VENDAS', 'MLB2', ['L-0320'], 'Queda', 's', 600,
                          regra='vendas_queda', medida=-0.6, periodo='x'),
                sd._sinal(NALA, 'ADS', 'MLB3', ['L-0303'], 'Sem venda', 's', 30,
                          regra='ads_sem_venda', medida=30, periodo='y')]

    def _preencher(self, conteudo, valores):
        """valores: {linha_excel: {coluna: valor}}"""
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(conteudo))
        ws = wb['Sinais']
        cab = [c.value for c in ws[1]]
        for linha, cols in valores.items():
            for nome, v in cols.items():
                ws.cell(row=linha, column=cab.index(nome) + 1, value=v)
        bio = io.BytesIO()
        wb.save(bio)
        return bio.getvalue()

    def test_arquivo_traz_a_chave_travada_e_colunas_livres(self):
        from openpyxl import load_workbook
        sinais = self._sinais()
        wb = load_workbook(io.BytesIO(sd.excel_sinais(sinais, {'L-0320': 'Lixeira'}, HOJE)))
        ws = wb['Sinais']
        cab = [c.value for c in ws[1]]
        self.assertEqual(tuple(cab), sd.COLUNAS_EXCEL)
        self.assertEqual([ws.cell(row=r, column=3).value for r in (2, 3, 4)],
                         ['E1', 'MLB2', 'MLB3'])
        self.assertTrue(ws.protection.sheet)
        self.assertTrue(ws.cell(row=2, column=cab.index('Objeto') + 1).protection.locked)
        for nome in sd.EDITAVEIS_EXCEL:
            self.assertFalse(ws.cell(row=2, column=cab.index(nome) + 1).protection.locked)
        listas = [dv.formula1 for dv in ws.data_validations.dataValidation]
        self.assertIn('"Ciente"', listas)
        self.assertTrue(any('Falta no fornecedor' in f for f in listas))
        self.assertIn('Instruções', wb.sheetnames)

    def test_ida_e_volta_com_previa_e_cada_recusa(self):
        sinais = self._sinais()
        hoje = HOJE
        arq = sd.excel_sinais(sinais, {}, hoje)
        arq = self._preencher(arq, {
            2: {'Status': 'Ciente', 'Motivo': 'Falta no fornecedor', 'Nota': 'abraçadeira',
                'Silenciar até': datetime(2026, 10, 20)},
            3: {'Status': 'ciente', 'Motivo': 'outro'},                     # sem nota
            4: {'Status': 'Ciente', 'Motivo': 'proposital'},                # loja fora
        })
        # linhas extras: sinal que não existe mais, repetida, status errado, datas ruins
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(arq))
        ws = wb['Sinais']
        ws.append([LPT, 'full_cobertura', 'E999'] + [None] * 9 + ['Ciente', 'Já em andamento', None, None])
        ws.append([LPT, 'full_cobertura', 'E1'] + [None] * 9 + ['Ciente', 'Já em andamento', None, None])
        ws.append([LPT, 'vendas_queda', 'MLB2'] + [None] * 9 + ['ok', None, None, None])
        ws.append([LPT, 'vendas_queda', 'MLB2'] + [None] * 9 + ['Ciente', 'Já em andamento', None, '31/02/2026'])
        ws.append([LPT, 'vendas_queda', 'MLB2'] + [None] * 9 + [None, 'Outro', 'ignorada', None])
        bio = io.BytesIO()
        wb.save(bio)
        linhas = sd.ler_excel(bio.getvalue())
        itens, previa = sd.previa_excel(linhas, sinais, [LPT], hoje)
        res = dict(zip(previa['Linha'], previa['Resultado']))
        self.assertEqual(res[2], '✅ ok')
        self.assertIn('pede uma nota', res[3])
        self.assertIn('loja fora do seu perfil', res[4])
        self.assertIn('não existe mais hoje', res[5])
        self.assertIn('repetido no arquivo (já na linha 2)', res[6])
        self.assertIn('Status deve ser "Ciente"', res[7])
        self.assertIn('repetido', res[8])                    # MLB2 já veio na linha 3
        self.assertNotIn(9, res)                             # sem Status: ignorada
        self.assertEqual(len(itens), 1)
        sinal, motivo, nota, ate = itens[0]
        self.assertIs(sinal, sinais[0])                      # medida/texto do cálculo de hoje
        self.assertEqual((motivo, nota, ate), ('falta_fornecedor', 'abraçadeira', date(2026, 10, 20)))

    def test_datas_e_motivo_por_codigo(self):
        sinais = self._sinais()[:2]
        arq = self._preencher(sd.excel_sinais(sinais, {}, HOJE), {
            2: {'Status': 'Ciente', 'Motivo': 'em_andamento'},              # data vazia
            3: {'Status': 'Ciente', 'Motivo': 'Decisão proposital', 'Silenciar até': '2026-12-25'}})
        itens, previa = sd.previa_excel(sd.ler_excel(arq), sinais, [LPT], HOJE)
        res = dict(zip(previa['Linha'], previa['Resultado']))
        self.assertEqual(itens[0][1:], ('em_andamento', '', HOJE + timedelta(days=7)))
        self.assertIn('no máximo 60 dias', res[3])
        arq = self._preencher(sd.excel_sinais(sinais, {}, HOJE), {
            2: {'Status': 'Ciente', 'Motivo': 'Decisão proposital', 'Silenciar até': '04/10/2026'},
            3: {'Status': 'Ciente', 'Motivo': 'Decisão proposital', 'Silenciar até': 'amanhã'}})
        itens, previa = sd.previa_excel(sd.ler_excel(arq), sinais, [LPT], HOJE)
        res = dict(zip(previa['Linha'], previa['Resultado']))
        self.assertIn('passado', res[2])
        self.assertIn('ilegível', res[3])
        self.assertEqual(itens, [])

    def test_limites_do_arquivo(self):
        with self.assertRaises(ValueError) as e:
            sd.ler_excel(b'x' * (sd.EXCEL_MAX_BYTES + 1))
        self.assertIn('limite', str(e.exception))
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = 'Sinais'
        ws.append(list(sd.COLUNAS_EXCEL))
        with mock.patch.object(sd, 'EXCEL_MAX_LINHAS', 3):
            for i in range(4):
                ws.append([LPT, 'vendas_queda', f'MLB{i}'])
            bio = io.BytesIO()
            wb.save(bio)
            with self.assertRaises(ValueError) as e:
                sd.ler_excel(bio.getvalue())
        self.assertIn('mais de 3 linhas', str(e.exception))

    def test_arquivo_sem_as_colunas(self):
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.append(['Loja', 'Regra'])
        bio = io.BytesIO()
        wb.save(bio)
        with self.assertRaises(ValueError):
            sd.ler_excel(bio.getvalue())

    def test_gravar_valida_cada_linha_antes_de_tocar_no_banco(self):
        conn = mock.Mock()
        s1, s2 = self._sinais()[:2]
        with self.assertRaises(ValueError):
            sd.gravar_cientes(conn, [(s1, 'proposital', '', HOJE), (s2, 'outro', '', HOJE)],
                              'larissa', HOJE, [LPT])
        conn.cursor.assert_not_called()


# ============================================================
# COM BANCO
# ============================================================

_DDL = [
    """CREATE TEMP TABLE fact_vendas_snapshot (
        marketplace_origem varchar, loja_origem varchar, numero_pedido varchar,
        data_venda date, sku varchar, codigo_anuncio varchar, quantidade integer,
        valor_venda_efetivo numeric, margem_total numeric,
        data_processamento timestamp, arquivo_origem varchar)""",
    """CREATE TEMP TABLE dim_metas_loja (
        loja_origem varchar, marketplace varchar, ano_mes varchar, meta_receita numeric)""",
    """CREATE TEMP TABLE dim_estoque_anuncio (
        marketplace varchar, loja varchar, anuncio_id varchar, variacao_id varchar,
        estoque_id varchar, sku varchar, ativo boolean, em_full boolean)""",
    """CREATE TEMP TABLE fact_estoque_diario (
        marketplace varchar, loja varchar, estoque_id varchar, data date,
        full_disponivel integer, full_em_transferencia integer, galpao_disponivel integer,
        unidades_recebidas integer, origem varchar, fonte_venda varchar,
        data_captura timestamp)""",
    """CREATE TEMP TABLE fact_ads_performance (
        data date, marketplace varchar, loja varchar, codigo_anuncio varchar,
        cliques integer, vendas integer, gasto_ads numeric, id_campanha varchar,
        data_importacao timestamp)""",
    """CREATE TEMP TABLE fact_ads_campanha_config (
        marketplace varchar, loja varchar, campanha varchar, data_captura timestamp,
        roas_objetivo numeric, orcamento_diario numeric, id_campanha varchar,
        detalhe jsonb)""",
    """CREATE TEMP TABLE fact_pedidos_itens_marketplace (
        marketplace varchar, loja varchar, numero_pedido varchar, sku varchar,
        id_anuncio_plataforma varchar, id_variacao varchar, quantidade integer,
        PRIMARY KEY (marketplace, loja, numero_pedido, sku))""",
    """CREATE TEMP TABLE fact_saude_anuncio (
        marketplace varchar, loja varchar, codigo_anuncio varchar, data date,
        tipo_sinal varchar, valor numeric, texto varchar, data_importacao timestamp)""",
    """CREATE TEMP TABLE dim_kit_composicao (
        kit_sku text, peca_sku text, quantidade integer)""",
    """CREATE TEMP TABLE dim_sku_mapeamento (
        sku_errado varchar PRIMARY KEY, sku_correto varchar NOT NULL)""",
    """CREATE TEMP TABLE dim_lojas (marketplace text, loja text)""",
]
TABELAS = ['fact_vendas_snapshot', 'dim_metas_loja', 'dim_estoque_anuncio',
           'fact_estoque_diario', 'fact_ads_performance', 'fact_ads_campanha_config',
           'fact_saude_anuncio', 'dim_kit_composicao', 'dim_sku_mapeamento', 'dim_lojas',
           'fact_pedidos_itens_marketplace']


class _Conexao:
    """Repassa o cursor da conexão do teste e IGNORA commit/rollback/close: o
    ROLLBACK de verdade é do teste (tearDown), senão as TEMP somem no meio.
    savepoint=True: o rollback do código volta ao SAVEPOINT (testa atomicidade)."""

    def __init__(self, conn, savepoint=False):
        self._conn = conn
        self._sp = savepoint
        if savepoint:
            conn.cursor().execute('SAVEPOINT antes_do_codigo')

    def cursor(self):
        return self._conn.cursor()

    def commit(self):
        pass

    def rollback(self):
        if self._sp:
            self._conn.cursor().execute('ROLLBACK TO SAVEPOINT antes_do_codigo')

    def close(self):
        pass


@unittest.skipUnless(DB_URL, 'defina NALA_TEST_DB_URL (usuário de permissão mínima)')
class ComBanco(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import psycopg2
        cls.conn = psycopg2.connect(DB_URL, connect_timeout=15)
        cls.conn.autocommit = False

    @classmethod
    def tearDownClass(cls):
        cls.conn.rollback()
        cls.conn.close()

    def _copiar_checks(self, tabela):
        """CHECK da tabela REAL (pg_constraint) na TEMP de mesmo nome."""
        self.cur.execute(
            "SELECT conname, pg_get_constraintdef(oid) FROM pg_constraint "
            "WHERE conrelid = %s::regclass AND contype = 'c'", (f'public.{tabela}',))
        checks = self.cur.fetchall()
        for nome, definicao in checks:
            self.cur.execute(f'ALTER TABLE pg_temp.{tabela} ADD CONSTRAINT "{nome}" {definicao}')
        return len(checks)

    def setUp(self):
        # o cenário (datas) foi montado com as janelas da v1, que terminavam em
        # D-3; a SQL não depende do deslocamento, então ele fica fixo aqui
        p = mock.patch.object(sd, 'DIAS_ATRASO_API', 2)
        p.start()
        self.addCleanup(p.stop)
        c = self.cur = self.conn.cursor()
        for ddl in _DDL:
            c.execute(ddl)
        # visão e tabela que JÁ podem existir em produção: TEMP vazias, para o
        # to_regclass do código achar a TEMP e nunca o objeto real
        c.execute("""CREATE TEMP VIEW vw_visitas_dia AS
            SELECT NULL::varchar AS marketplace, NULL::varchar AS loja,
                   NULL::varchar AS codigo_anuncio, NULL::date AS data,
                   NULL::integer AS visitas WHERE false""")
        c.execute("""CREATE TEMP TABLE vw_views_shopee_30d (
            marketplace varchar, loja varchar, codigo_anuncio varchar, data date,
            views_30d integer, rating_star numeric, comment_count integer,
            data_importacao timestamp)""")
        for ddl in _ddl_ciente_em_temp():
            c.execute(ddl)
        for ddl in _ddl_historico_em_temp():
            c.execute(ddl)
        copiadas = {t: self._copiar_checks(t) for t in TABELAS}
        # as CHECK que existem hoje em produção (05/10/2026) chegaram à TEMP
        self.assertGreaterEqual(copiadas['fact_estoque_diario'], 2)
        self.assertGreaterEqual(copiadas['fact_vendas_snapshot'], 1)
        self.assertGreaterEqual(copiadas['dim_kit_composicao'], 3)
        self.assertGreaterEqual(copiadas['fact_ads_performance'], 1)
        self.assertGreaterEqual(copiadas['fact_pedidos_itens_marketplace'], 1)
        c.execute("""
            SELECT bool_and(c.relnamespace = pg_my_temp_schema())
            FROM unnest(%s) AS t(nome) JOIN pg_class c ON c.oid = t.nome::regclass
        """, (TABELAS + ['vw_visitas_dia', 'sinal_ciente', 'vw_views_shopee_30d',
                          'sinal_historico'],))
        self.assertTrue(c.fetchone()[0], 'tabela não resolveu para pg_temp')
        self._popular()
        self.c = _Conexao(self.conn)

    def tearDown(self):
        self.cur.close()
        self.conn.rollback()

    def _popular(self):
        c, ML = self.cur, 'MERCADO LIVRE'
        c.execute("INSERT INTO dim_lojas VALUES (%s, %s), (%s, %s), ('SHOPEE', 'Shopee-LPT')",
                  (ML, LPT, ML, NALA))
        c.execute("INSERT INTO dim_metas_loja VALUES (%s, %s, '2026-10', 200000)", (LPT, ML))
        # K-L-0421-A: dois anúncios, dois estoques (números reais de 04/10)
        c.executemany("INSERT INTO dim_estoque_anuncio VALUES (%s,%s,%s,NULL,%s,%s,true,true)", [
            (ML, LPT, 'MLB5183384677', 'MLBU5071722523', 'K-L-0421-A'),
            (ML, LPT, 'MLB4793561323', 'MLBU4131286394', 'K-L-0421-A'),
            (ML, NALA, 'MLB7175749876', 'MLBU4334892068', 'LTT-CT3821')])
        foto = date(2026, 10, 4)
        c.executemany("INSERT INTO fact_estoque_diario VALUES "
                      "(%s,%s,%s,%s,%s,%s,NULL,0,'snapshot',NULL,%s)", [
            (ML, LPT, 'MLBU5071722523', foto, 6, 2, datetime(2026, 10, 5, 4, 8)),
            (ML, LPT, 'MLBU4131286394', foto, 89, 0, datetime(2026, 10, 5, 4, 8)),
            (ML, LPT, 'MLBU5071722523', date(2026, 10, 3), 99, 0, datetime(2026, 10, 4, 4, 8))]
            # Garrafa: Full zerado em 4 dos 7 dias da semana (só 3 com Full)
            + [(ML, NALA, 'MLBU4334892068', date(2026, 9, 26 + i) if i < 5 else date(2026, 10, i - 4),
                10 if i < 3 else 0, 0, datetime(2026, 10, 5, 4, 1)) for i in range(7)]
            + [(ML, NALA, 'MLBU4334892068', foto, 59, 0, datetime(2026, 10, 5, 4, 1))])
        vendas = []
        for i in range(21):                       # 21 un. na semana 26/09–02/10
            vendas.append((ML, LPT, f'P{i}', date(2026, 9, 26 + i % 5), 'K-L-0421-A',
                           'MLB5183384677', 1, 30, 3))
        for i in range(44):                       # +44 no resto dos 30 dias
            vendas.append((ML, LPT, f'Q{i}', date(2026, 9, 3 + i % 20), 'K-L-0421-A',
                           'MLB5183384677', 1, 30, 3))
        vendas.append((ML, LPT, 'H', HOJE, 'K-L-0421-A', 'MLB5183384677', 500, 9999, 0))  # hoje: fora
        vendas.append((ML, LPT, 'D1', date(2026, 10, 4), 'K-L-0421-A', 'MLB5183384677',
                       3, 90, 9))             # D-1: só no resumo
        vendas.append((ML, NALA, 'G1', date(2026, 9, 27), 'LTT-CT3821', 'MLB7175749876', 5, 150, 20))
        vendas.append((ML, NALA, 'G2', date(2026, 9, 10), 'LTT-CT3821', 'MLB7175749876', 100, 3200, 300))
        c.executemany("INSERT INTO fact_vendas_snapshot VALUES "
                      "(%s,%s,%s,%s,%s,%s,%s,%s,%s,'2026-10-05 06:04','API')", vendas)
        c.executemany("INSERT INTO fact_ads_performance VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)", [
            (date(2026, 9, 10), ML, NALA, 'MLB7175749876', 100, 10, 300, 'MLB7175749876 x',
             datetime(2026, 10, 5, 5, 31)),
            (date(2026, 9, 27), ML, NALA, 'MLB7175749876', 5, 0, 10, 'MLB7175749876 x',
             datetime(2026, 10, 5, 5, 31))])
        c.executemany("INSERT INTO fact_ads_campanha_config VALUES (%s,%s,%s,%s,%s,%s,%s)", [
            (ML, LPT, 'MLB5183384677 kit', datetime(2026, 10, 4, 5, 30), 14, 12, 'MLB5183384677 kit'),
            (ML, LPT, 'MLB5183384677 kit', datetime(2026, 10, 5, 5, 30), 10, 12, 'MLB5183384677 kit'),
            (ML, LPT, 'MLB4793561323 kit', datetime(2026, 10, 4, 5, 30), 9, 5, 'MLB4793561323 kit'),
            (ML, LPT, 'MLB4793561323 kit', datetime(2026, 10, 5, 5, 30), 9, 5, 'MLB4793561323 kit')])
        c.executemany("INSERT INTO fact_saude_anuncio VALUES (%s,%s,%s,%s,'experiencia_compra',%s,%s,%s)", [
            (ML, LPT, 'MLB5183384677', HOJE, 65, 'Média', datetime(2026, 10, 5, 5, 6)),
            (ML, LPT, 'MLB5183384677', date(2026, 9, 28), 100, 'Boa', datetime(2026, 9, 28, 5, 6)),
            (ML, LPT, 'MLB4793561323', HOJE, 100, 'Boa', datetime(2026, 10, 5, 5, 6))])

    def test_ponta_a_ponta_pelas_sql(self):
        dados, erros = sd.ler_tudo(self.c, HOJE, [LPT, NALA])
        self.assertEqual(erros, {})
        self.assertEqual(dados['visitas'], [])               # TEMP vazia, nunca a real
        self.assertEqual(sd.ler_cientes(self.c, [LPT, NALA]), [])   # TEMP vazia, nunca a real
        s, erros_bloco = sd.montar_sinais(dados, HOJE)
        self.assertEqual(erros_bloco, {})
        por = {(x['loja'], x['tipo'], x['numero'].split(':')[0].split(' ')[0]): x for x in s}

        full = [x for x in s if x['tipo'] == 'FULL' and x['loja'] == LPT]
        self.assertEqual(len(full), 1)                       # só o de 6 un.; o de 89 não
        self.assertEqual(full[0]['anuncio'], 'MLB5183384677')
        self.assertTrue(full[0]['urgente'])
        self.assertIn('(6 disp. + 2 em transf.', full[0]['numero'])

        self.assertIn('ROAS objetivo 14,0 → 10,0', por[(LPT, 'ADS', 'Mudou')]['numero'])
        self.assertEqual(len([x for x in s if x['numero'].startswith('Mudou')]), 1)
        self.assertIn('100 (Boa) → 65 (Média)', por[(LPT, 'EXPERIÊNCIA', 'Experiência')]['numero'])

        rup = por[(NALA, 'FULL', 'Ruptura')]
        self.assertIn('só 3 de 7 dias da semana, Full em 3', rup['numero'])
        self.assertFalse(any(x['numero'].startswith('Espiral') for x in s))

        r = sd.resumo_lojas(dados['resumo'], dados['metas'], [LPT, NALA], HOJE)
        self.assertAlmostEqual(r[LPT]['ontem'], 90)          # D-1 entra no resumo
        self.assertEqual(r[LPT]['meta'], 200000)
        self.assertIsNone(r[NALA]['meta'])

        fres = {(f, l) for f, l, _u in dados['frescor']}
        self.assertIn(('Estoque', LPT), fres)
        self.assertIn(('Experiência', LPT), fres)

    def _popular_shopee(self):
        c, SH = self.cur, 'SHOPEE'
        vendas = [(SH, S_LPT, f'S{i}', date(2026, 9, 10), 'L-0320', 'L-0320', 1, 200, 5)
                  for i in range(10)]                       # base: R$ 2.000
        vendas.append((SH, S_LPT, 'S99', date(2026, 9, 27), 'L-0320', 'L-0320', 1, 100, 5))
        c.executemany("INSERT INTO fact_vendas_snapshot VALUES "
                      "(%s,%s,%s,%s,%s,%s,%s,%s,%s,'2026-10-05 06:04','API')", vendas)
        c.executemany("INSERT INTO fact_pedidos_itens_marketplace VALUES "
                      "(%s,%s,%s,%s,'2249','0',1)",
                      [(SH, S_LPT, v[2], 'L-0320') for v in vendas])
        c.executemany("INSERT INTO fact_vendas_snapshot VALUES "
                      "(%s,%s,%s,%s,%s,%s,%s,%s,%s,'2026-09-20 10:00','upload Yanni.xlsx')",
                      [(SH, S_YANNI, f'Y{i}', date(2026, 9, 12), 'L-0500', 'L-0500', 1, 200, 5)
                       for i in range(12)])
        c.execute("INSERT INTO fact_ads_performance VALUES "
                  "('2026-09-20', %s, %s, '2249', 50, 5, 100, '106156986', '2026-10-05 05:32')",
                  (SH, S_LPT))
        c.executemany("INSERT INTO fact_ads_campanha_config VALUES (%s,%s,'Kit Escova',%s,%s,10,"
                      "'106156986', '{\"item_id\": \"2249\"}')", [
                          (SH, S_LPT, datetime(2026, 10, 4, 5, 32), 5),
                          (SH, S_LPT, datetime(2026, 10, 5, 5, 32), 0)])
        c.executemany("INSERT INTO vw_views_shopee_30d VALUES (%s,%s,'2249',%s,%s,%s,%s,%s)", [
            (SH, S_LPT, date(2026, 10, 4), 100, 4.5, 20, datetime(2026, 10, 5, 4, 30)),
            (SH, S_LPT, date(2026, 9, 4), 400, 4.8, 10, datetime(2026, 9, 5, 4, 30))])

    def test_shopee_ponta_a_ponta_pelas_sql(self):
        self._popular_shopee()
        lojas = [S_LPT, S_YANNI]
        dados, erros = sd.ler_tudo(self.c, HOJE, lojas, sd.SHOPEE)
        self.assertEqual(erros, {})
        self.assertNotIn('experiencia', dados)               # só do ML
        anuncios = {(r[0], r[1]) for r in dados['vendas']}
        self.assertEqual(anuncios, {(S_LPT, '2249'), (S_YANNI, 'SKU:L-0500')})
        s, erros_bloco = sd.montar_sinais(dados, HOJE, sd.SHOPEE)
        self.assertEqual(erros_bloco, {})
        por = {(x['loja'], x['regra']): x for x in s}
        self.assertEqual(por[(S_LPT, 'vendas_queda')]['anuncio'], '2249')
        self.assertEqual(por[(S_YANNI, 'vendas_queda')]['objeto'], 'SKU:L-0500')
        cfg = por[(S_LPT, 'ads_config')]
        self.assertIn('ROAS objetivo 5,0 → automático', cfg['numero'])
        self.assertEqual(cfg['objeto'], '2249@2026-10-05')
        custo = por[(S_LPT, 'ads_custo')]['numero']
        self.assertIn('Custo por PEDIDO de ads R$ 20,00', custo)   # 100 ÷ 5 pedidos
        self.assertIn('margem por PEDIDO R$ 5,00', custo)          # 55 ÷ 11 pedidos
        views = por[(S_LPT, 'visitas_queda')]
        self.assertIn('Visualizações de página (30d) 100 × 400', views['numero'])
        self.assertEqual(views['periodo'], '05/09–04/10 × 06/08–04/09')
        nota = por[(S_LPT, 'exp_anuncio')]
        self.assertAlmostEqual(nota['medida'], 4.2)                # (90 − 48) ÷ 10 novas
        self.assertFalse([x for x in s if x['regra'] in ('ads_espiral', 'full_ruptura')])
        self.assertEqual(sd.lojas_carga_pendente(dados['frescor'], lojas, HOJE), set())
        self.assertEqual(sd.lojas_upload_atrasado(dados['frescor'], lojas, HOJE),
                         {S_YANNI: date(2026, 9, 12)})

    def test_litstore_relida_ate_o_ultimo_upload_pelas_sql(self):
        c, SH = self.cur, 'SHOPEE'
        c.executemany("INSERT INTO fact_vendas_snapshot VALUES "
                      "(%s,%s,%s,%s,%s,%s,%s,%s,%s,'2026-10-02 10:00','upload Yanni.xlsx')",
                      [(SH, S_YANNI, f'B{i}', date(2026, 9, 5), 'L-0500', 'L-0500', 1, 200, 5)
                       for i in range(12)]
                      + [(SH, S_YANNI, 'U1', date(2026, 9, 30), 'L-0500', 'L-0500', 1, 50, 1)])
        dados, erros = sd.ler_tudo(self.c, HOJE, [S_YANNI], sd.SHOPEE)
        self.assertEqual(erros, {})
        self.assertEqual(sd.lojas_upload_atrasado(dados['frescor'], [S_YANNI], HOJE),
                         {S_YANNI: date(2026, 9, 30)})
        sinais, ate = sd.reler_vendas_upload(self.c, dados, [], HOJE, [S_YANNI], sd.SHOPEE)
        self.assertEqual(ate, {S_YANNI: date(2026, 9, 30)})
        q = [x for x in sinais if x['regra'] == 'vendas_queda']
        self.assertEqual(len(q), 1)
        j = sd.janelas(date(2026, 10, 1))                    # janelas do último upload
        self.assertIn(f"semana {j['sem_ini']:%d/%m}–{j['fim']:%d/%m}", q[0]['numero'])

    def test_backfill_le_a_foto_e_a_config_daquele_dia(self):
        # captura de config em 03/10 (ROAS 20); as de 04 e 05/10 já estão no cenário
        self.cur.execute("INSERT INTO fact_ads_campanha_config VALUES "
                         "('MERCADO LIVRE', %s, 'MLB5183384677 kit', '2026-10-03 05:30', 20, 12, "
                         "'MLB5183384677 kit')", (LPT,))
        dia = date(2026, 10, 4)                              # backfill de 04/10
        dados, erros = sd.ler_tudo(self.c, dia, [LPT])
        self.assertEqual(erros, {})
        foto = {r[1]: (r[2], r[3]) for r in dados['foto']}
        self.assertEqual(foto['MLBU5071722523'], (date(2026, 10, 3), 99))   # a foto de 03/10
        cfg = [r for r in dados['config'] if r[1] == 'MLB5183384677 kit']
        self.assertEqual(len(cfg), 1)                        # 03/10 -> 04/10 (não 05/10)
        self.assertEqual((float(cfg[0][5]), float(cfg[0][6])), (20.0, 14.0))

    def test_historico_grava_com_upsert_e_a_tela_le(self):
        x = sd._sinal(LPT, 'VENDAS', 'MLB1', [], 'x', 's', 100, regra='vendas_queda', medida=-0.5)
        y = sd._sinal(LPT, 'FULL', 'MLB2', [], 'x', 's', 50, regra='full_cobertura',
                      objeto='E2', medida=6)
        for dia in (date(2026, 10, 3), date(2026, 10, 4)):
            sd.gravar_historico(self.c, [x, y], dia, sd.MARKETPLACE)
        sd.gravar_historico(self.c, [dict(x, em_jogo=999)], date(2026, 10, 4), sd.MARKETPLACE)
        self.cur.execute("SELECT data, objeto, em_jogo FROM sinal_historico ORDER BY 1, 2")
        self.assertEqual([(r[0], r[1], float(r[2])) for r in self.cur.fetchall()],
                         [(date(2026, 10, 3), 'E2', 50.0), (date(2026, 10, 3), 'MLB1', 100.0),
                          (date(2026, 10, 4), 'E2', 50.0), (date(2026, 10, 4), 'MLB1', 999.0)])
        dados, erros = sd.ler_tudo(self.c, HOJE, [LPT])
        self.assertEqual(erros, {})
        self.assertEqual(sorted(dados['historico_dias']), [date(2026, 10, 3), date(2026, 10, 4)])
        marcado = sd.aparece_desde([x], dados['historico'], dados['historico_dias'], HOJE)[0]
        self.assertEqual(marcado['desde_txt'], 'há 2 dias (desde 03/10)')

    def test_check_do_historico_iguais_as_de_producao_quando_aplicado(self):
        self.cur.execute("SELECT to_regclass('public.sinal_historico') IS NOT NULL")
        if not self.cur.fetchone()[0]:
            self.skipTest('sql/sinais_historico.sql ainda não aplicado em produção')
        q = ("SELECT conname, pg_get_constraintdef(oid) FROM pg_constraint "
             "WHERE conrelid = %s::regclass AND contype = 'c' ORDER BY conname")
        self.cur.execute(q, ('public.sinal_historico',))
        real = self.cur.fetchall()
        self.cur.execute(q, ('pg_temp.sinal_historico',))
        self.assertEqual(real, self.cur.fetchall())

    def test_lojas_ml_e_restricao_pela_sql(self):
        lojas = [r[0] for r in sd._ler(self.c, sd.SQL_LOJAS_ML, {'marketplace': sd.MARKETPLACE})]
        self.assertEqual(lojas, [LPT, NALA])
        dados, _ = sd.ler_tudo(self.c, HOJE, [NALA])
        self.assertTrue(all(r[0] == NALA for r in dados['vendas']))
        self.assertTrue(all(r[0] == NALA for r in dados['foto']))

    def test_visitas_le_da_view_quando_ela_existe(self):
        self.cur.execute("""CREATE OR REPLACE TEMP VIEW vw_visitas_dia AS
            SELECT 'MERCADO LIVRE'::varchar AS marketplace, 'ML-LPT'::varchar AS loja,
                   'MLB5183384677'::varchar AS codigo_anuncio, d::date AS data,
                   CASE WHEN d >= '2026-09-26' THEN 10 ELSE 100 END AS visitas
              FROM generate_series('2026-08-29'::date, '2026-10-04'::date, '1 day') d""")
        dados, erros = sd.ler_tudo(self.c, HOJE, [LPT])
        self.assertEqual(erros, {})
        self.assertEqual([tuple(r[:2]) for r in dados['visitas']], [(LPT, 'MLB5183384677')])
        self.assertEqual(int(dados['visitas'][0][2]), 70)    # 7 dias × 10
        self.assertEqual(int(dados['visitas'][0][3]), 2800)  # 28 dias × 100


@unittest.skipUnless(DB_URL, 'defina NALA_TEST_DB_URL (usuário de permissão mínima)')
class CienteComBanco(unittest.TestCase):
    """Ciclo do Ciente em TEMP criada com o DDL LIDO de sql/sinais_ciente.sql."""

    @classmethod
    def setUpClass(cls):
        import psycopg2
        cls.conn = psycopg2.connect(DB_URL, connect_timeout=15)
        cls.conn.autocommit = False

    @classmethod
    def tearDownClass(cls):
        cls.conn.rollback()
        cls.conn.close()

    def setUp(self):
        c = self.cur = self.conn.cursor()
        for ddl in _ddl_ciente_em_temp():
            c.execute(ddl)
        c.execute("SELECT 'sinal_ciente'::regclass::oid IN "
                  "(SELECT oid FROM pg_class WHERE relnamespace = pg_my_temp_schema())")
        self.assertTrue(c.fetchone()[0], 'tabela não resolveu para pg_temp')
        c.execute("SELECT count(*) FROM pg_constraint WHERE conrelid = 'sinal_ciente'::regclass "
                  "AND contype = 'c'")
        self.assertEqual(c.fetchone()[0], 7)

    def tearDown(self):
        self.cur.close()
        self.conn.rollback()

    def _s(self, objeto='E1', regra='full_cobertura', loja=LPT, medida=6):
        return sd._sinal(loja, 'FULL', 'MLB1', ['A'], 'texto do sinal', 's', 1,
                         regra=regra, objeto=objeto, medida=medida)

    def _abertos(self, lojas=(LPT, NALA)):
        return sd._ler(_Conexao(self.conn), sd.SQL_CIENTES_ABERTOS,
                       {'marketplace': sd.MARKETPLACE, 'lojas': list(lojas)})

    def test_grava_le_e_silencia(self):
        hoje = datetime.now(sd.BRT).date()
        n = _gravar(_Conexao(self.conn), [self._s(), self._s('E2')], 'outro',
                              'abraçadeira em falta', hoje + timedelta(days=7), 'larissa',
                              hoje, [LPT])
        self.assertEqual(n, 2)
        abertos = self._abertos()
        self.assertEqual(len(abertos), 2)
        c = dict(zip(sd.CIENTE_COLS, abertos[0]))
        self.assertEqual((c['motivo'], c['nota'], c['criado_por']),
                         ('outro', 'abraçadeira em falta', 'larissa'))
        self.assertEqual(c['texto_no_ciente'], 'texto do sinal')
        self.assertEqual(float(c['medida_no_ciente']), 6)
        ativos, sil = sd.aplicar_cientes([self._s(), self._s('E3')], abertos, hoje)
        self.assertEqual([x['objeto'] for x in ativos], ['E3'])
        self.assertEqual([x['objeto'] for x in sil], ['E1'])

    def test_novo_ciente_substitui_o_aberto(self):
        hoje = datetime.now(sd.BRT).date()
        for motivo in ('em_andamento', 'proposital'):
            _gravar(_Conexao(self.conn), [self._s()], motivo, '', hoje, 'larissa',
                              hoje, [LPT])
        self.cur.execute("SELECT motivo, como_encerrou, encerrado_por FROM sinal_ciente ORDER BY id")
        self.assertEqual(self.cur.fetchall(), [('em_andamento', 'substituido', 'larissa'),
                                               ('proposital', None, None)])

    def test_indice_impede_dois_abertos_na_mesma_chave(self):
        import psycopg2
        ins = ("INSERT INTO sinal_ciente (marketplace, loja, regra, objeto, motivo, "
               "silenciar_ate, criado_por) VALUES ('MERCADO LIVRE', 'ML-LPT', "
               "'full_cobertura', 'E1', 'proposital', CURRENT_DATE, 'x')")
        self.cur.execute(ins)
        with self.assertRaises(psycopg2.errors.UniqueViolation):
            self.cur.execute(ins)

    def test_check_recusam_o_que_o_codigo_nunca_deveria_gravar(self):
        import psycopg2
        base = {'marketplace': 'MERCADO LIVRE', 'loja': LPT, 'regra': 'full_cobertura',
                'objeto': 'E9', 'motivo': 'proposital', 'nota': None,
                'silenciar_ate': datetime.now(sd.BRT).date(), 'medida': 1, 'texto': 't',
                'usuario': 'x'}
        ruins = [{'motivo': 'esqueci'}, {'motivo': 'outro', 'nota': '  '},
                 {'regra': 'qualquer'}, {'marketplace': 'Mercado Livre'},
                 {'silenciar_ate': datetime.now(sd.BRT).date() + timedelta(days=61)},
                 {'silenciar_ate': datetime.now(sd.BRT).date() - timedelta(days=1)},
                 {'objeto': ' '}]
        self.cur.execute(sd.SQL_INSERIR_CIENTE, {**base, 'objeto': 'E8'})
        for sem_quem in ("encerrado_por = NULL", "encerrado_por = '  '"):
            self.cur.execute('SAVEPOINT q')
            with self.assertRaises(psycopg2.errors.CheckViolation, msg=sem_quem):
                self.cur.execute(f"UPDATE sinal_ciente SET encerrado_em = now(), {sem_quem}, "
                                 "como_encerrou = 'reativado' WHERE objeto = 'E8'")
            self.cur.execute('ROLLBACK TO SAVEPOINT q')
        for mudanca in ruins:
            self.cur.execute('SAVEPOINT r')
            with self.assertRaises(psycopg2.errors.CheckViolation, msg=str(mudanca)):
                self.cur.execute(sd.SQL_INSERIR_CIENTE, {**base, **mudanca})
            self.cur.execute('ROLLBACK TO SAVEPOINT r')
        self.cur.execute(sd.SQL_INSERIR_CIENTE, base)        # o caso bom passa

    def test_transacao_unica_um_erro_nao_grava_nada(self):
        hoje = datetime.now(sd.BRT).date()
        ruim = self._s('E2', regra='regra_que_nao_existe')
        with self.assertRaises(Exception):
            _gravar(_Conexao(self.conn, savepoint=True), [self._s(), ruim],
                              'proposital', '', hoje, 'larissa', hoje, [LPT])
        self.cur.execute('SELECT count(*) FROM sinal_ciente')
        self.assertEqual(self.cur.fetchone()[0], 0)

    def test_cada_linha_com_seu_motivo_numa_transacao(self):
        hoje = datetime.now(sd.BRT).date()
        n = sd.gravar_cientes(_Conexao(self.conn), [
            (self._s('E1'), 'falta_fornecedor', 'abraçadeira', hoje + timedelta(days=30)),
            (self._s('E2'), 'outro', 'teste A/B', hoje)], 'larissa', hoje, [LPT])
        self.assertEqual(n, 2)
        self.cur.execute("SELECT objeto, motivo, nota, silenciar_ate - CURRENT_DATE "
                         "FROM sinal_ciente ORDER BY objeto")
        self.assertEqual(self.cur.fetchall(), [('E1', 'falta_fornecedor', 'abraçadeira', 30),
                                               ('E2', 'outro', 'teste A/B', 0)])

    def test_reativar_so_nas_lojas_do_usuario(self):
        hoje = datetime.now(sd.BRT).date()
        _gravar(_Conexao(self.conn), [self._s()], 'proposital', '', hoje,
                          'larissa', hoje, [LPT])
        id_ = self._abertos()[0][0]
        self.assertEqual(sd.reativar_ciente(_Conexao(self.conn), id_, 'patricia', [NALA]), 0)
        self.assertEqual(sd.reativar_ciente(_Conexao(self.conn), id_, 'larissa', [LPT]), 1)
        self.assertEqual(self._abertos(), [])
        self.cur.execute("SELECT como_encerrou, encerrado_por FROM sinal_ciente")
        self.assertEqual(self.cur.fetchall(), [('reativado', 'larissa')])

    def test_ler_tudo_le_os_cientes_quando_a_tabela_existe(self):
        hoje = datetime.now(sd.BRT).date()
        _gravar(_Conexao(self.conn), [self._s(), self._s(loja=NALA)], 'proposital',
                          '', hoje, 'larissa', hoje, [LPT, NALA])
        c = _Conexao(self.conn)
        self.assertTrue(sd._ler(c, sd.SQL_EXISTE_CIENTE, {})[0][0])  # ler_tudo usa isto
        p = sd.params(hoje, [LPT])
        self.assertEqual([r[1] for r in sd._ler(c, sd.SQL_CIENTES_ABERTOS, p)], [LPT])

    def test_check_de_producao_iguais_as_do_arquivo_quando_aplicado(self):
        self.cur.execute("SELECT to_regclass('public.sinal_ciente') IS NOT NULL")
        if not self.cur.fetchone()[0]:
            self.skipTest('sql/sinais_ciente.sql ainda não aplicado em produção')
        q = ("SELECT conname, pg_get_constraintdef(oid) FROM pg_constraint "
             "WHERE conrelid = %s::regclass AND contype = 'c' ORDER BY conname")
        self.cur.execute(q, ('public.sinal_ciente',))
        real = self.cur.fetchall()
        self.cur.execute(q, ('pg_temp.sinal_ciente',))
        self.assertEqual(real, self.cur.fetchall())


if __name__ == '__main__':
    unittest.main()
